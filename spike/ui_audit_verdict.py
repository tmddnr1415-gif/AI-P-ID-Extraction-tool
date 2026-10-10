"""hotfix82 — O/X 평가를 띄워서 눌러 확인한다.

    python3 spike/ui_audit_verdict.py <data_dir 사본> <job> out/hotfix82/ui

★ 실 DB 를 열지 않는다 — 넘긴 데이터 폴더(사본)를 PID_DATA_DIR 로 가리킨다.

  ① 목록 O/X 칸의 O 단추 → 서버 O · 단추 켜짐 · 도구줄 수
  ② X 단추 → 서버 X · 행이 출력에서 빠짐(취소선) · 다시 누르면 지워지고 되살아남
  ③ 칸을 고르고 O / X 키 → 적히고 아래 칸으로 · Backspace → 지움
  ④ Shift+↓ 범위 뒤 X → 범위 전체
  ⑤ 도구줄 O/X — 고른 행 · 묶음
  ⑥ Ctrl+Z / Ctrl+Y
  ⑦ VOC 함 — 요청마다 한 건 · 행 전부 · 단일 X 는 조각
  ⑧ 근거 패널 단추 · 열 필터 선택지
"""
from __future__ import annotations
import json, os, socket, subprocess, sys, time, urllib.request
from pathlib import Path
from urllib.parse import quote

ROOT = Path(__file__).resolve().parent.parent
data, JOB, OUT = Path(sys.argv[1]), sys.argv[2], Path(sys.argv[3])
OUT.mkdir(parents=True, exist_ok=True)
s = socket.socket(); s.bind(("127.0.0.1", 0)); port = s.getsockname()[1]; s.close()
env = dict(os.environ, PID_DATA_DIR=str(data), PYTHONPATH=str(ROOT), PID_PAGE_WARM="0", PID_RENDER_POOL="0", PID_ROWS_WARM="0")
env.pop("PID_VOC_DIR", None)
srv = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port)],
                       cwd=str(ROOT), env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
base = f"http://127.0.0.1:{port}"
FOUND, BAD = [], []
def note(x): print(x, flush=True); FOUND.append(x)
def check(ok, what):
    note(("  ✔ " if ok else "  ✘ ") + what)
    if not ok: BAD.append(what)
def srv_row(key):
    rows = json.loads(urllib.request.urlopen(f"{base}/jobs/{JOB}/rows?keys={quote(key)}", timeout=30).read())
    return rows[0]
def srv_vx(key):
    r = srv_row(key); return (r.get("verdict") or {}).get("verdict", ""), bool(r.get("removed"))
inbox = data / "voc" / "inbox"
def settle(pg, ms=250):
    # 평가 한 번이 끝날 때까지 (`NET.vxBusy` — 요청 · 그 행 다시 받기 · 다음 칸 옮기기까지).  첫 쓰기는 행 메모를
    # 다시 만들어 1초 넘게 걸리므로 고정 대기로는 안 된다 (첫 판이 그렇게 틀렸다).
    pg.wait_for_timeout(150)
    # `NET` 는 const 라 window.NET 이 아니다 — typeof 로 본다 (첫 판은 이 때문에 기다리지 않았다)
    pg.wait_for_function("() => typeof NET === 'undefined' || !(NET.vxBusy > 0 || NET.writing > 0)", timeout=60000)
    pg.wait_for_timeout(ms)
def vocs():
    return sorted(p for p in inbox.iterdir() if p.is_dir() and not p.name.startswith(".")) if inbox.exists() else []

try:
    for _ in range(90):
        try: urllib.request.urlopen(base + "/version", timeout=2); break
        except Exception: time.sleep(1)
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        br = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
        ctx = br.new_context(viewport={"width": 1920, "height": 1080})
        ctx.add_init_script("try{localStorage.setItem('pid.side.off','1');localStorage.removeItem('pid.split')}catch(e){}")
        pg = ctx.new_page(); errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)))
        pg.on("console", lambda m: errs.append("console:" + m.text) if m.type == "error" else None)
        pg.goto(f"{base}/?embed=1&user={quote('홍길동')}"); settle(pg)
        pg.evaluate(f"open('{JOB}')")
        pg.wait_for_function("() => typeof S !== 'undefined' && S.job && !S.loading && S.vlist && S.vlist.length > 6", timeout=180000)
        settle(pg)
        if pg.evaluate("() => !!S.side"):
            pg.evaluate("() => toggleSide(false)"); pg.wait_for_timeout(500)
        keys = pg.evaluate("""() => { const ok = r => !r.deleted && !r.removed && !verdictOf(r) && !r.added && r.values.tag_no; const L = S.vlist;
            for (let i = 0; i + 5 < L.length; i++) if ([0,1,2,3,4,5].every(j => ok(L[i+j]))) return [0,1,2,3,4,5].map(j => L[i+j].key);
            return null; }""")
        k0, k1, k2, k3, k4, k5 = keys
        td = lambda k, c: f'#body tr[data-key="{k}"] td[data-col="{c}"]'
        vxb = lambda k, v: f'{td(k, "verdict")} button.vx-b.{v.lower()}'
        cur = lambda: pg.evaluate("() => S.cell ? [S.cell.key, S.cell.col] : null")
        on = lambda k, v: pg.evaluate(f"() => !!document.querySelector('{vxb(k, v)}.on')")
        cnt = lambda: pg.text_content("#vx-count")
        n0 = len(vocs())
        check(pg.evaluate("() => gridCols().map(c => c[0]).indexOf('verdict') >= 0"), "목록에 O/X 열이 있다")
        hd = pg.evaluate("() => gridCols().map(c => c[0])")
        check(hd.index("verdict") + 1 == hd.index("type"), f"O/X 열은 Type 바로 앞 ({hd[:6]})")

        note("① 목록 칸의 O 단추")
        pg.click(vxb(k0, "O")); settle(pg)
        check(srv_vx(k0) == ("O", False), f"서버 O · 행은 그대로 {srv_vx(k0)}")
        check(on(k0, "O") and not on(k0, "X"), "O 단추 켜짐")
        check("O 1" in cnt(), f"도구줄 수 '{cnt()}'")
        check(len(vocs()) == n0 + 1, "VOC 한 건 늘었다")
        pg.screenshot(path=str(OUT / "1_O.png"), clip={"x": 960, "y": 500, "width": 960, "height": 420})

        note("② X 단추 → 출력에서 빠짐 → 다시 누르면 되살아남")
        pg.click(vxb(k1, "X")); settle(pg)
        check(srv_vx(k1) == ("X", True), f"서버 X · removed {srv_vx(k1)}")
        check(pg.evaluate(f"() => document.querySelector('#body tr[data-key=\"{k1}\"]').classList.contains('deleted')"), "행에 취소선(deleted)")
        check(on(k1, "X"), "X 단추 켜짐")
        check(pg.evaluate(f"() => !!document.querySelector('#body tr[data-key=\"{k1}\"] button[data-act=restore]')"), "되돌리기 단추가 섰다")
        pg.screenshot(path=str(OUT / "2_X.png"), clip={"x": 960, "y": 500, "width": 960, "height": 420})
        pg.click(vxb(k1, "X")); settle(pg)
        check(srv_vx(k1) == ("", False), f"켜진 X 를 다시 누르면 평가 지움 · 되살아남 {srv_vx(k1)}")
        check(not on(k1, "X") and not on(k1, "O"), "단추 둘 다 꺼짐")

        note("③ 자판 — 칸 고르고 o · x · Backspace")
        pg.click(td(k2, "verdict")); pg.wait_for_timeout(300)
        check(cur() == [k2, "verdict"], f"O/X 칸이 골라짐 {cur()}")
        pg.keyboard.press("o"); settle(pg)
        check(srv_vx(k2)[0] == "O", "o 키 → O")
        check(cur() == [k3, "verdict"], f"적은 뒤 아래 칸으로 {cur()}")
        pg.keyboard.press("x"); settle(pg)
        check(srv_vx(k3) == ("X", True), f"x 키 → X · 출력에서 빠짐 {srv_vx(k3)}")
        check(cur() == [k4, "verdict"], f"다시 아래 칸 {cur()}")
        pg.keyboard.press("ArrowUp"); pg.wait_for_timeout(300)
        pg.keyboard.press("Backspace"); settle(pg)
        check(srv_vx(k3) == ("", False), f"Backspace → 지움 · 되살아남 {srv_vx(k3)}")
        check(not pg.evaluate("() => !!CELL.editing"), "O/X 칸은 편집이 열리지 않는다")

        note("④ Shift+↓ 범위 뒤 x")
        pg.click(td(k3, "verdict")); pg.wait_for_timeout(300)
        pg.keyboard.press("Shift+ArrowDown"); pg.keyboard.press("Shift+ArrowDown"); pg.wait_for_timeout(300)
        check(pg.evaluate("() => rangeRows().length") == 3, "범위 3행")
        pg.keyboard.press("x"); settle(pg)
        check(all(srv_vx(k)[0] == "X" for k in (k3, k4, k5)), "범위 3행 전부 X")
        check(len(vocs()) == n0 + 5, f"VOC 는 요청마다 한 건 — 지우기는 안 남긴다 (5건 기대 · {len(vocs()) - n0}건)")
        pg.screenshot(path=str(OUT / "3_범위_X.png"), clip={"x": 960, "y": 500, "width": 960, "height": 420})

        note("⑤ 도구줄 O/X — 고른 행 · 묶음")
        pg.keyboard.press("Escape"); pg.wait_for_timeout(200)
        pg.click(td(k1, "page_no")); pg.wait_for_timeout(300)
        pg.click("#vx-o"); settle(pg)
        check(srv_vx(k1)[0] == "O", "도구줄 O → 고른 행 O")
        pg.click("#vx-o"); settle(pg)
        check(srv_vx(k1)[0] == "", "같은 단추 다시 → 지움")
        pg.evaluate(f"() => {{ clearMulti && clearMulti(); select('{k4}', true); toggleMulti('{k5}'); }}"); pg.wait_for_timeout(400)
        check(pg.evaluate("() => S.multi.size") == 2, "묶음 2행")
        check(pg.evaluate("() => !!document.querySelector('#evidence #mv-o')"), "묶음 판에 O/X 줄")
        pg.click("#evidence #mv-o"); settle(pg)
        check(srv_vx(k4)[0] == "O" and srv_vx(k5)[0] == "O", "묶음 판 O → 둘 다 O · 되살아남")
        check(not srv_vx(k4)[1] and not srv_vx(k5)[1], "X 가 뺐던 행이 O 로 되살아났다")

        note("⑥ Ctrl+Z / Ctrl+Y")
        pg.evaluate("() => clearMulti()"); pg.wait_for_timeout(200)
        pg.click(td(k1, "page_no")); pg.keyboard.press("Escape"); pg.wait_for_timeout(200)
        pg.evaluate("() => document.getElementById('gridwrap').focus()")
        pg.keyboard.press("Control+z"); settle(pg)
        check(srv_vx(k4)[0] == "X" and srv_vx(k5)[0] == "X", f"Ctrl+Z → 묶음 O 가 X 로 돌아감 {srv_vx(k4)} {srv_vx(k5)}")
        pg.keyboard.press("Control+y"); settle(pg)
        check(srv_vx(k4)[0] == "O" and srv_vx(k5)[0] == "O", "Ctrl+Y → 다시 O")

        note("⑦ VOC 함")
        vs = [json.loads((p / "voc.json").read_text(encoding="utf-8")) for p in vocs()[n0:]]
        check(all(v["source"] == "VERDICT" and v["category"] == "VERDICT" for v in vs), "전부 VERDICT")
        check(all(len(v.get("verdicts", [])) == len(v.get("rows", [])) >= 1 for v in vs), "verdicts 와 rows 가 행마다")
        three = [v for v in vs if len(v["verdicts"]) == 3]
        check(bool(three) and three[0]["counts"] == {"O": 0, "X": 3, "qty_O": 0, "qty_X": 0}, "범위 3행 X 가 한 건에 셋")
        one_x = [p for p in vocs()[n0:] if (p / "crop.png").exists()]
        check(bool(one_x), f"단일 X 는 도면 조각을 남긴다 ({len(one_x)}건)")
        check(all(v["author"] == "홍길동" for v in vs), "작성자 = 대시보드 로그인 이름 (묻지 않았다)")
        api = json.loads(urllib.request.urlopen(f"{base}/jobs/{JOB}/verdicts", timeout=30).read())
        check(api["counts"]["O"] == 4 and api["counts"]["X"] == 1, f"GET verdicts 집계 (O k0·k2·k4·k5 · X k3) {api['counts']}")
        lst = subprocess.run([sys.executable, "spike/voc.py", "--inbox", str(inbox), "--ledger", str(data / "ledger.json"),
                              "list"], cwd=str(ROOT), capture_output=True, text=True)
        check("O/X 평가" in lst.stdout and "VOC-" not in lst.stdout, f"spike/voc.py list 는 평가를 숨기고 안내만 ({lst.stdout.strip().splitlines()[:2]})")

        note("⑧ 근거 패널 · 열 필터")
        pg.click(td(k0, "page_no")); pg.wait_for_timeout(800)
        check(pg.evaluate("() => !!document.querySelector('#evidence .vxed button.o.on')"), "근거 패널 O 켜짐")
        pg.click("#evidence .vxed button.x"); settle(pg)
        check(srv_vx(k0) == ("X", True), "근거 패널 X → 서버 X")
        vals = pg.evaluate("() => columnValues('verdict').map(x => x[0])")
        check("O" in vals and "X" in vals, f"O/X 열 필터 선택지 {vals}")
        pg.screenshot(path=str(OUT / "4_근거패널.png"), clip={"x": 960, "y": 500, "width": 960, "height": 580})
        note("⑨ 식별 VOC 탭 — 수량 O/X · 비고 · 누락 ＋행 · VOC Excel")
        pg.click('#tabs button[data-tab="VOC"]'); pg.wait_for_timeout(800)
        check(pg.evaluate("() => S.tab") == "VOC", "식별 VOC 탭이 열린다")
        hd2 = pg.evaluate("() => gridCols().map(c => c[0])")
        check(hd2[:4] == ["page_no", "pid_no", "verdict", "qty_verdict"] and "vx_note" in hd2 and "vx_who" in hd2,
              f"VOC 탭 열 순서 {hd2}")
        check(pg.evaluate("() => S.vlist.length") >= 1900, f"VOC 탭은 전체 행을 보인다 ({pg.evaluate('() => S.vlist.length')})")
        check(pg.evaluate("() => !document.getElementById('vx-xlsx').classList.contains('hidden')"), "[VOC Excel] 단추가 보인다")
        badge = pg.text_content('#tabs button[data-tab="VOC"] .n')
        check(badge and int(badge) >= 5, f"탭 배지 = 평가를 적은 행 수 ({badge})")
        # 칸 가운데를 누르면 단추에 맞을 수 있어(좁은 열) 칸은 코드로 고른다 — 자판 흐름을 재는 것이 목적
        pg.evaluate(f"() => {{ setCell('{k0}', 'qty_verdict'); document.getElementById('gridwrap').focus(); }}"); pg.wait_for_timeout(300)
        check(cur() == [k0, "qty_verdict"], f"수량 O/X 칸이 골라짐 {cur()}")
        pg.keyboard.press("x"); settle(pg)
        rr = srv_row(k0)
        check(rr["verdict"]["qty_verdict"] == "X" and rr["verdict"]["verdict"] == "X", f"수량 X 는 식별 X 와 따로 ({rr['verdict']})")
        check(cur() != [k0, "qty_verdict"] and cur()[1] == "qty_verdict", f"수량 칸도 적으면 아래로 {cur()}")
        pg.click(td(k2, "qty_verdict") + " button[data-act=qvx].o"); settle(pg)
        check(srv_row(k2)["verdict"]["qty_verdict"] == "O", "수량 O 단추")
        pg.dblclick(td(k2, "vx_note")); pg.wait_for_timeout(300)
        pg.keyboard.press("Control+A"); pg.keyboard.type("x2 여야 함"); pg.keyboard.press("Enter"); settle(pg)
        check(srv_row(k2)["verdict"]["note"] == "x2 여야 함", "비고 칸 → 서버 note")
        check(pg.text_content(td(k2, "vx_note")).strip() == "x2 여야 함", "비고가 목록에 보인다")
        pg.screenshot(path=str(OUT / "5_식별VOC탭.png"), clip={"x": 960, "y": 440, "width": 960, "height": 480})
        # 누락 ＋행 — 마크업 추가 행이 이 탭에 선다 (추가 대화상자는 마크업 자기검증이 맡는다 · 여기서는 API 로)
        import urllib.request as _u
        body = json.dumps({"page_no": 6, "tab": "FIELD", "values": {"type": "LIT", "qty": 1}, "rect": [50, 50, 80, 90],
                           "author": "홍길동", "reason_class": "MISSING", "note": "누락"}).encode()
        req = _u.Request(f"{base}/jobs/{JOB}/rows", data=body, headers={"Content-Type": "application/json"})
        added = json.loads(_u.urlopen(req, timeout=30).read())["key"]
        pg.evaluate(f"() => refreshRows('{added}', {{keys: ['{added}'], toGrid: true}})"); pg.wait_for_timeout(1500)
        check(pg.evaluate(f"() => !!S.rowByKey['{added}'] && S.vlist.some(r => r.key === '{added}')"), "사람이 더한 행이 VOC 탭 목록에 선다")
        pg.click(vxb(added, "O")); settle(pg)
        check(srv_row(added)["verdict"]["verdict"] == "O", "더한 행에 O")
        xr = _u.urlopen(f"{base}/jobs/{JOB}/verdicts.xlsx", timeout=60)
        xb = xr.read()
        import openpyxl, io
        wb = openpyxl.load_workbook(io.BytesIO(xb))
        ws = wb["식별 VOC"]; hd = [c.value for c in ws[1]]
        got = {r[hd.index("Tag No.")] or r[hd.index("No")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        check(ws.max_row >= 1990 and wb["누락 추가"].max_row >= 2, f"VOC Excel — 식별 VOC {ws.max_row - 1}행 · 누락 추가 {wb['누락 추가'].max_row - 1}행")
        k2tag = srv_row(k2)["values"].get("tag_no")
        r2 = got.get(k2tag)
        check(bool(r2) and r2[hd.index("수량 O/X")] == "O" and r2[hd.index("비고")] == "x2 여야 함", f"Excel 에 수량 O · 비고 ({k2tag})")
        check(not errs, f"페이지 오류 {len(errs)} {errs[:3]}")
        br.close()
finally:
    srv.terminate()
    try: srv.wait(10)
    except Exception: srv.kill()
(OUT / "audit.txt").write_text("\n".join(FOUND) + f"\n\n결함 {len(BAD)}\n", encoding="utf-8")
print("BAD", len(BAD))
sys.exit(1 if BAD else 0)
