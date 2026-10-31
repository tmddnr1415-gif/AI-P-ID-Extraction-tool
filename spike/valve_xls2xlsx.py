"""DEV ONLY — .xls -> .xlsx with values, merges, widths, heights, fonts, fills, borders, wrap (xlrd + openpyxl).

The company server converts with Excel itself (app/valve/tbe.py xls_to_xlsx).  This exists so the
Valve List export can be tested in a checkout that has no Excel.  Needs `pip install xlrd` (not shipped).

    python3 spike/valve_xls2xlsx.py in.xls out.xlsx
"""
import sys

import openpyxl
import xlrd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_BORDER = {0: None, 1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick", 6: "double", 7: "hair"}


def _rgb(book, idx):
    c = book.colour_map.get(idx)
    return "{:02X}{:02X}{:02X}".format(*c) if c else None


def convert(src, dst):
    book = xlrd.open_workbook(src, formatting_info=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31])
        for r in range(sh.nrows):
            if r in sh.rowinfo_map and sh.rowinfo_map[r].height:
                ws.row_dimensions[r + 1].height = sh.rowinfo_map[r].height / 20
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                cell = ws.cell(r + 1, c + 1, v if v != "" else None)
                xf = book.xf_list[sh.cell_xf_index(r, c)]
                f = book.font_list[xf.font_index]
                cell.font = Font(name=f.name, size=f.height / 20, bold=bool(f.bold), italic=bool(f.italic),
                                 color=_rgb(book, f.colour_index))
                bg = xf.background
                if bg.fill_pattern and _rgb(book, bg.pattern_colour_index):
                    cell.fill = PatternFill("solid", fgColor=_rgb(book, bg.pattern_colour_index))
                b = xf.border
                side = lambda s, col: Side(style=_BORDER.get(s), color=_rgb(book, col)) if _BORDER.get(s) else Side()  # noqa: E731
                cell.border = Border(left=side(b.left_line_style, b.left_colour_index),
                                     right=side(b.right_line_style, b.right_colour_index),
                                     top=side(b.top_line_style, b.top_colour_index),
                                     bottom=side(b.bottom_line_style, b.bottom_colour_index))
                al = xf.alignment
                cell.alignment = Alignment(wrap_text=bool(al.text_wrapped),
                                           horizontal={1: "left", 2: "center", 3: "right"}.get(al.hor_align),
                                           vertical={0: "top", 1: "center", 2: "bottom"}.get(al.vert_align))
        for c, info in sh.colinfo_map.items():
            ws.column_dimensions[get_column_letter(c + 1)].width = info.width / 256
        for r0, r1, c0, c1 in sh.merged_cells:
            ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1, end_column=c1)
    wb.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
