"""Excel 출력을 **끝에서 끝까지** 확인한다 (사용자 요구 — 54회차 뒤).

    python3 spike/excel_check.py out/round54/excel

무엇을 하는가 — 사람이 `Excel 출력` 을 누르는 것과 같은 길로 간다:
    저장된 결과로 임시 DB → 마크업 행 하나 추가 → 칸 하나 편집 →
    `POST /jobs/{id}/snapshot` → `GET /revisions/{id}/excel` → zip 을 풀어
    **실제 xlsx 를 열어** 센다.

★ 양식이 없으면 아무것도 지어내지 않는다.  이 환경에는 발주처 양식이 없으므로
  21회차가 **그 양식으로 써 낸 산출물**(`out/handover_r21/*.xlsx`)을 양식 자리에
  둔다 — 발주처의 시트 이름·열·서식을 그대로 담고 있고, `write_deliverable` 은
  쓰기 전에 행을 지우므로 옛 값이 남지 않는다.  PNEUMATIC·MASTER 는 양식이
  없으므로 **건너뛴 사유가 나오는지**를 확인한다 (그것도 요구된 동작이다).

★ 실 DB 를 열지 않는다 (17회차 격리).
"""
from __future__ import annotations
import hashlib, io, json, os, shutil, socket, subprocess, sys, tempfile, time, zipfile
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "out/round54/excel")
OUT.mkdir(parents=True, exist_ok=True)
FOUND: list[str] = []

DOCS = [("TC2", "out/regression_3p/TC2.json", "data/TC2_260821.pdf"),
        ("AL NOUF1", "out/regression_3p/AL_NOUF1.json", "data/pid_total.pdf")]
FORMS = {"FIELD": "out/handover_r21/field_pid_total.xlsx",
         "MOV": "out/handover_r21/mov_pid_total.xlsx",
         "BFV": "out/handover_r21/bfv_pid_total.xlsx"}


def note(s):
    print(s, flush=True); FOUND.append(s)


def free_port():
    s = socket.socket(); s.bind(("127.0.0.1", 0)); p = s.getsockname()[1]; s.close(); return p


def call(method, path, body=None):
    req = urllib.request.Request(f"http://127.0.0.1:{port}{path}", method=method)
    data = None
    if body is not None:
        data = json.dumps(body).encode()
        req.add_header("Content-Type", "application/json")
    with urllib.request.urlopen(req, data, timeout=300) as r:
        raw = r.read()
    ctype = r.headers.get("Content-Type", "")
    return json.loads(raw) if "json" in ctype else raw


data = Path(tempfile.mkdtemp(prefix="xlcheck-"))
os.environ["PID_DATA_DIR"] = str(data)
from app import db, revisions                                    # noqa: E402

# 양식을 사람이 올리는 그 자리에 둔다
(data / "templates").mkdir(parents=True, exist_ok=True)
for kind, p in FORMS.items():
    src = ROOT / p
    if src.exists():
        shutil.copyfile(src, data / "templates" / f"{kind}.xlsx")
    else:
        note(f"⚠ {kind} 양식 없음 — 건너뛴 사유가 나오는지만 본다")

con = db.connect(data / "app.db"); jobs = {}
for i, (name, res_p, pdf_p) in enumerate(DOCS):
    src, pdf = ROOT / res_p, ROOT / pdf_p
    if not src.exists() or not pdf.exists():
        note(f"⚠ {name}: 자료 없음 — 건너뛴다"); continue
    result = json.loads(src.read_text()); result = result.get("result", result)
    job = f"xlcheck{i:04d}"; jobs[name] = job
    revisions.create_project(data, name)
    con.execute("INSERT INTO job (id, pdf_name, pdf_sha256, pdf_path, created_at, status,"
                " progress, message, fingerprint, project, revision)"
                " VALUES (?,?,?,?,?,'done',1.0,'',?,?,?)",
                (job, pdf.name, hashlib.sha256(pdf.read_bytes()).hexdigest(),
                 str(pdf.resolve()), time.time() + i,
                 result.get("fingerprint", ""), name, "A"))
    con.commit(); db.store_result(con, job, result); con.commit()
    revisions.record_revision(data, name, "A", job_id=job, pdf_name=pdf.name,
                              compared_with="",
                              result={"counts": {}, "states": {}, "radii": {},
                                      "deleted_candidates": []})
con.close()

port = free_port()
env = dict(os.environ, PID_DATA_DIR=str(data), PYTHONPATH=str(ROOT))
srv = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app",
                        "--host", "127.0.0.1", "--port", str(port)],
                       cwd=str(ROOT), env=env,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
try:
    for _ in range(90):
        try:
            urllib.request.urlopen(f"http://127.0.0.1:{port}/home", timeout=2); break
        except Exception:
            time.sleep(1)
    import openpyxl
    from app import excel_out
    from app.pipeline import CFG

    for name, job in jobs.items():
        note(f"\n══ {name} ══════════════════════════════════════════")
        rows = call("GET", f"/jobs/{job}/rows")
        rows = rows["rows"] if isinstance(rows, dict) else rows
        note(f"행 {len(rows)}")

        # ① 사람이 만든 값 둘 — 마크업 행 하나 · 칸 편집 하나
        page_no = rows[0]["page_no"]
        added = call("POST", f"/jobs/{job}/rows",
                     {"page_no": page_no, "rect": [300.0, 470.0, 370.0, 525.0],
                      "tab": "FIELD", "drawing_no": rows[0].get("drawing_no") or "",
                      "values": {"type": "PIT", "qty": 3, "scope": "SCT"},
                      "author": "출력확인", "note": "Excel 출력 확인용",
                      "reason_class": "MISSING"})
        note(f"① 마크업 행 추가 — {json.dumps(added, ensure_ascii=False)[:160]}")
        target = next(r for r in rows if r["tab"] == "FIELD")
        edited = call("PATCH", f"/jobs/{job}/rows/{target['key']}",
                      {"field": "description", "value": "사람이 적은 설명",
                       "author": "출력확인"})
        note(f"① 칸 편집 — {target['values'].get('type')} description"
             f" → user {json.dumps(edited.get('user'), ensure_ascii=False)[:100]}")

        # ② 스냅샷 → 출력
        snap = call("POST", f"/jobs/{job}/snapshot", {"label": "출력 확인"})
        rev_id = snap["revision_id"] if isinstance(snap, dict) else snap
        note(f"② 스냅샷 {json.dumps(snap, ensure_ascii=False)[:200]}")
        blob = call("GET", f"/revisions/{rev_id}/excel")
        zpath = OUT / f"{name.replace(' ', '_')}.zip"
        zpath.write_bytes(blob)
        note(f"② zip {len(blob):,} 바이트 → {zpath}")

        # ③ 실제 xlsx 를 열어 센다
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            names = z.namelist()
            note(f"③ zip 안 {names}")
            man = json.loads(z.read("MANIFEST.json"))
            note(f"③ MANIFEST — form_scope {man['form_scope']!r}"
                 f" · held {man['held_rows']} · out_of_scope {man['out_of_scope_rows']}"
                 f" · legacy {man['legacy_no_scope_rows']}")
            for w in man["written"]:
                note(f"   쓴 것 {w['kind']} — 시트 {w['sheet']!r} · 행 {w['rows']}"
                     f" · 양식 행 {w['template_rows']} · 미대응 편집 {w['unmapped_values']}")
            for s in man["skipped"]:
                note(f"   건너뜀 {s['kind']} — {s['rows']}행 · {s['reason'][:70]}")
            out_dir = OUT / name.replace(" ", "_")
            out_dir.mkdir(exist_ok=True)
            for n in names:
                if n.endswith(".xlsx"):
                    (out_dir / n).write_bytes(z.read(n))

        for x in sorted(out_dir.glob("*.xlsx")):
            wb = openpyxl.load_workbook(x)
            kind = x.name.split("_")[0].upper()
            cfgsheet = wb.sheetnames[0]
            ws = wb[cfgsheet]
            # 머리글 줄과 첫 데이터 줄은 **설정이 답한다** — 짐작하지 않는다
            # (이 양식은 머리글이 두 줄이라 7 로 두면 둘째 머리글을 데이터로 읽는다).
            spec = excel_out.DELIVERABLES[kind]
            _sheet, first, _cols = excel_out._cfg_columns(CFG, spec["config"])
            hdr_row = first - 2
            headers = [ws.cell(hdr_row, c).value for c in range(1, ws.max_column + 1)]
            # NO 열이 1..N 로 이어지는가 · 마지막 데이터 행
            nos, last = [], first - 1
            for r in range(first, ws.max_row + 1):
                v = ws.cell(r, 1).value
                if isinstance(v, int):
                    nos.append(v); last = r
            gaps = [i for i, v in enumerate(nos, 1) if v != i]
            note(f"③ {x.name} — 시트 {cfgsheet!r} · 데이터 {len(nos)}행"
                 f" (r{first}~r{last}) · NO 이어짐 {'예' if not gaps else f'★ 아니오 {gaps[:5]}'}")
            note(f"   머리글 r{hdr_row} {[h for h in headers if h][:9]}")
            # 첫 행과 마지막 행이 실제로 값을 갖는가
            def cells(r):
                return {headers[c - 1]: ws.cell(r, c).value
                        for c in range(1, ws.max_column + 1)
                        if headers[c - 1] and ws.cell(r, c).value not in (None, "")}
            note(f"   첫 행 {json.dumps(cells(first), ensure_ascii=False, default=str)[:260]}")
            note(f"   끝 행 {json.dumps(cells(last), ensure_ascii=False, default=str)[:260]}")
            # 사람이 만든 값이 실제로 들어갔나
            hits = {"사용자 추가": 0, "사람이 적은 설명": 0}
            for r in range(first, last + 1):
                for c in range(1, ws.max_column + 1):
                    v = ws.cell(r, c).value
                    if not isinstance(v, str):
                        continue
                    for k in hits:
                        if k in v:
                            hits[k] += 1
            note(f"   사람이 만든 값 — {hits}")
            # 양식 뒤쪽이 깨끗한가 (옛 데이터가 남지 않았는가)
            tail = [r for r in range(last + 1, min(ws.max_row, last + 40) + 1)
                    if any(ws.cell(r, c).value not in (None, "")
                           for c in range(1, 10))]
            note(f"   데이터 뒤 남은 줄 {len(tail)}{'' if not tail else f' — r{tail[:5]}'}")

        # ④ 두 번 써서 같은가 (결정성) · ⑤ SCT 만 담는 모드
        con2 = db.connect(data / "app.db")
        snapshot = db.get_revision(con2, rev_id)
        con2.close()
        a_dir, b_dir = OUT / "det_a", OUT / "det_b"
        tpl = {k: (data / "templates" / f"{k}.xlsx") for k in FORMS}
        r1 = excel_out.write_all(snapshot, tpl, a_dir, CFG)
        r2 = excel_out.write_all(snapshot, tpl, b_dir, CFG)
        same = True
        for w1, w2 in zip(r1["written"], r2["written"]):
            s1, s2 = (openpyxl.load_workbook(w["path"])[w["sheet"]] for w in (w1, w2))
            for r in range(1, max(s1.max_row, s2.max_row) + 1):
                for c in range(1, max(s1.max_column, s2.max_column) + 1):
                    if s1.cell(r, c).value != s2.cell(r, c).value:
                        same = False
                        note(f"④ ★ 두 번이 다르다 — {w1['kind']} r{r}c{c}")
                        break
                if not same:
                    break
            if not same:
                break
        note(f"④ 두 번 써서 칸이 같은가 — {'같다' if same else '★ 다르다'}")

        # `form_scope_mode` 는 `cfg.data["client_form"]["scope_filter"]` 하나를
        # 본다 (판정하는 곳이 거기 하나다).  그 칸만 바꾼 대역을 만든다.
        class _SctCfg:
            def __init__(self, base):
                self._b = base
                self.data = dict(base.data)
                self.data["client_form"] = dict(self.data.get("client_form") or {},
                                                scope_filter=excel_out.FORM_SCOPE_SCT)

            def __getattr__(self, n):
                return getattr(self._b, n)

        r3 = excel_out.write_all(snapshot, tpl, OUT / "sct", _SctCfg(CFG))
        note("⑤ SCT 만 담기 — "
             + " · ".join(f"{w['kind']} {w['rows']}" for w in r3["written"])
             + f" · 빠진 행 {r3['held_rows']} (전량 모드에서는 "
             + " · ".join(f"{w['kind']} {w['rows']}" for w in r1["written"]) + ")")
finally:
    srv.terminate()
(OUT / "확인결과.txt").write_text("\n".join(FOUND) + "\n", encoding="utf8")
print("\n=== 정리 ===\n" + "\n".join(FOUND))
