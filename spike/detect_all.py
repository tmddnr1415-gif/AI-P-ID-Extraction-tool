#!/usr/bin/env python3
"""Phase 0 spike 2b — apply the page-6 rules to every P&ID page and see where
they break.

No new detection rules.  This runs exactly the ruleset validated on page 6
(spike/detect_symbols.py) across every in-scope P&ID page and measures the
damage, so that rule work in the next phase is aimed at real failures rather
than guessed ones.

Two things are measured rather than assumed, because page 6 is not
representative of the set:

* **Bubble size.**  Page 6 uses a 68.0 x 22.6 pt stadium, but the document has
  at least two families.  The detector derives bubbles from geometry (two arc
  caps joined by straight sides) with no size constant, and this script reports
  the size distribution it actually found per page.

* **Vendor mark meaning.**  The mark dictionary is page-scoped
  (docs/design.md §10.1).  Where a page's NOTES do not define a mark that
  appears on a bubble, the exclusion is *held*, not guessed, and the case is
  reported as VENDOR_MARK_UNDEFINED.

Comparison
----------
Joined to the Field Instrument Excel on `P&ID No.`.  Three drawing numbers
appear on two pages each, so the join is 1:N and everything is aggregated **by
drawing number**, never by page.  Only row counts are compared: `Q'ty` is a
note multiplier rather than a symbol count (docs/design.md §8) and belongs to
spike 3.

Usage
-----
    python spike/detect_all.py --pdf data/pid_total.pdf \
           --compare data/CZE_Field_Instrument.xlsx \
           --titleblocks out/titleblocks.csv \
           --json out/detect_all.json --failures out/failure_report.md
"""

from __future__ import annotations

import argparse
import collections
import csv
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from detect_symbols import (  # noqa: E402
    FIELD_TYPE_MAP,
    NOT_FIELD_INSTRUMENT,
    bubble_sizes,
    detect,
    find_bubbles,
    find_marks,
)

from pidcache import load_pages  # noqa: E402

# A vendor mark written as literal text, e.g. "(*)" or "(**)".  Page 6 draws its
# marks as vector glyphs; other sheets type them instead, and the page-6 rule
# only knows how to see the drawn form.
MARK_TEXT_RE = re.compile(r"^\({1,2}\*{1,3}\)?$")


def mark_notation(pc, mark_dict):
    """How this page expresses vendor marks — drawn glyphs, typed text, or none.

    Reported because the exclusion rule validated on page 6 can only see the
    drawn form, so a page that types its marks is silently unfiltered.
    """
    notes_text = sum(1 for r, t in pc.words if MARK_TEXT_RE.match(t) and r.x0 > 1960)
    draw_text = sum(1 for r, t in pc.words if MARK_TEXT_RE.match(t) and r.x0 < 1960)
    draw_glyph = sum(1 for m in find_marks(pc) if m.x < 1960)
    return {"glyph_definitions": len(mark_dict),
            "text_definitions_in_notes": notes_text,
            "text_marks_in_drawing": draw_text,
            "glyph_marks_in_drawing": draw_glyph}

# The two rules whose contribution is being measured.  The other two
# (NOT_FIELD_INSTRUMENT, VALVE) are type filters, not scope rules, and are
# always on.
SCOPE_RULES = ("VENDOR_MARK", "SCT_SUPPLIER_SCOPE")

COMBOS = {
    "none": frozenset(),
    "vendor_only": frozenset({"VENDOR_MARK"}),
    "sct_only": frozenset({"SCT_SUPPLIER_SCOPE"}),
    "both": frozenset(SCOPE_RULES),
}

FAILURE_TYPES = [
    "NO_ANCHOR",
    "AMBIGUOUS_GEOMETRY",
    "SCOPE_CONFLICT",
    "VENDOR_MARK_UNDEFINED",
    "TYPE_MAPPING_MISS",
    "OVER_DETECT",
]


def included_under(det, active: frozenset) -> bool:
    """Whether a detection survives a given rule combination.

    Every rule that fired is recorded on the detection regardless of whether it
    was allowed to exclude, so all four combinations come from a single
    detection pass instead of four.
    """
    hits = det.rules_hit
    if "NOT_FIELD_INSTRUMENT" in hits or "VALVE" in hits:
        return False
    return not any(r in active and r in hits for r in SCOPE_RULES)


def system_of(drawing_no: str) -> str:
    """KKS system code — 'D00P-10LBA10-M05-0001' -> 'LBA'."""
    parts = drawing_no.split("-")
    return parts[1][2:5] if len(parts) > 1 and len(parts[1]) >= 5 else "???"


def load_excel(path: Path) -> tuple[dict, int, int]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["2.0_Instrument List"]
    by_drawing: dict[str, list] = collections.defaultdict(list)
    total = spare = 0
    for i in range(8, ws.max_row + 1):
        no = ws.cell(i, 1).value
        if not isinstance(no, (int, float)):
            continue
        total += 1
        pid = ws.cell(i, 7).value
        pid = str(pid).strip() if pid else ""
        if not pid:
            spare += 1
            continue
        by_drawing[pid].append({
            "row": i, "no": no,
            "type": str(ws.cell(i, 8).value).strip() if ws.cell(i, 8).value else "",
            "system": ws.cell(i, 6).value,
            "description": ws.cell(i, 10).value,
        })
    return by_drawing, total, spare


def main() -> int:
    ap = argparse.ArgumentParser(description="Phase 0 spike 2b — all pages")
    ap.add_argument("--pdf", default="data/pid_total.pdf")
    ap.add_argument("--compare", default="data/CZE_Field_Instrument.xlsx")
    ap.add_argument("--titleblocks", default="out/titleblocks.csv")
    ap.add_argument("--json", default="out/detect_all.json")
    ap.add_argument("--failures", default="out/failure_report.md")
    args = ap.parse_args()

    tb = {int(r["page_no"]): r for r in
          csv.DictReader(open(args.titleblocks, encoding="utf-8-sig"))}
    excel, excel_total, excel_spare = load_excel(Path(args.compare))

    doc, pages = load_pages(args.pdf)
    targets = [
        pc for pc in pages
        if tb[pc.page_no]["page_kind"] == "PID"
        and tb[pc.page_no]["analysis_scope"] == "True"
    ]
    print(f"pages in scope: {len(targets)}   "
          f"excel rows: {excel_total} ({excel_spare} SPARE without P&ID No.)")

    # ---- detect ------------------------------------------------------
    per_page = {}
    size_dist_global = collections.Counter()
    for i, pc in enumerate(targets, 1):
        dets, scopes, mark_dict, unverified, unmapped = detect(pc)
        sizes = bubble_sizes(find_bubbles(pc))
        notation = mark_notation(pc, mark_dict)
        size_dist_global.update(sizes)
        per_page[pc.page_no] = {
            "drawing_no": tb[pc.page_no]["drawing_no"],
            "title": tb[pc.page_no]["drawing_title"],
            "detections": dets,
            "scopes": scopes,
            "mark_dict": mark_dict,
            "unverified": unverified,
            "unmapped": unmapped,
            "sizes": sizes,
            "notation": notation,
        }
        print(f"  [{i:>2}/{len(targets)}] p{pc.page_no:<3} {tb[pc.page_no]['drawing_no']}"
              f"  anchors={len(dets):<4} unverified={len(unverified):<3}"
              f" marks_defined={len(mark_dict)}", flush=True)

    # ---- aggregate by drawing number (the join is 1:N) ---------------
    by_drawing: dict[str, dict] = collections.defaultdict(
        lambda: {"pages": [], "detections": [], "unverified": [], "unmapped": [],
                 "mark_dicts": [], "scopes": []})
    for page_no, info in per_page.items():
        d = by_drawing[info["drawing_no"]]
        d["pages"].append(page_no)
        d["detections"].extend((page_no, x) for x in info["detections"])
        d["unverified"].extend((page_no, x) for x in info["unverified"])
        d["unmapped"].extend((page_no, x) for x in info["unmapped"])
        d["mark_dicts"].append((page_no, info["mark_dict"]))
        d["scopes"].extend((page_no, s) for s in info["scopes"])

    all_drawings = sorted(set(by_drawing) | set(excel))

    # ---- metrics per rule combination --------------------------------
    def score(active: frozenset):
        tp = det_total = exp_total = 0
        for dn in all_drawings:
            det_counts = collections.Counter(
                x.excel_type for _, x in by_drawing.get(dn, {}).get("detections", [])
                if included_under(x, active)
            )
            exp_counts = collections.Counter(r["type"] for r in excel.get(dn, []))
            for t in set(det_counts) | set(exp_counts):
                tp += min(det_counts.get(t, 0), exp_counts.get(t, 0))
            det_total += sum(det_counts.values())
            exp_total += sum(exp_counts.values())
        return {
            "tp": tp, "detected": det_total, "expected": exp_total,
            "recall": tp / exp_total if exp_total else 0.0,
            "precision": tp / det_total if det_total else 0.0,
        }

    combo_scores = {name: score(active) for name, active in COMBOS.items()}

    # ---- per-drawing detail for the shipped combination --------------
    ACTIVE = COMBOS["both"]
    drawing_rows = []
    for dn in all_drawings:
        info = by_drawing.get(dn)
        det_counts = collections.Counter(
            x.excel_type for _, x in (info["detections"] if info else [])
            if included_under(x, ACTIVE)
        )
        exp_counts = collections.Counter(r["type"] for r in excel.get(dn, []))
        tp = sum(min(det_counts.get(t, 0), exp_counts.get(t, 0))
                 for t in set(det_counts) | set(exp_counts))
        drawing_rows.append({
            "drawing_no": dn,
            "system": system_of(dn),
            "pages": info["pages"] if info else [],
            "title": (per_page[info["pages"][0]]["title"] if info else ""),
            "detected": sum(det_counts.values()),
            "expected": sum(exp_counts.values()),
            "tp": tp,
            "det_counts": dict(det_counts),
            "exp_counts": dict(exp_counts),
            "has_page": bool(info),
            "in_excel": dn in excel,
        })

    # ---- failure classification --------------------------------------
    failures = classify_failures(drawing_rows, by_drawing, excel, ACTIVE)

    # ---- report -------------------------------------------------------
    print_report(combo_scores, drawing_rows, failures, size_dist_global, per_page)

    write_failure_report(Path(args.failures), failures, drawing_rows,
                         combo_scores, size_dist_global, per_page,
                         excel_total, excel_spare, excel)

    payload = {
        "pages_in_scope": len(targets),
        "excel_rows_total": excel_total,
        "excel_rows_spare": excel_spare,
        "combo_scores": combo_scores,
        "bubble_size_distribution": {f"{k[0]}x{k[1]}": v
                                     for k, v in size_dist_global.most_common()},
        "drawings": drawing_rows,
        "failures": {k: [f for f in failures if f["type"] == k] for k in FAILURE_TYPES},
    }
    out = Path(args.json)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nwrote {out} and {args.failures}")
    return 0


def classify_failures(drawing_rows, by_drawing, excel, active):
    """Explain every count mismatch, using evidence from the page itself."""
    failures = []
    for row in drawing_rows:
        dn = row["drawing_no"]
        info = by_drawing.get(dn)
        det_counts, exp_counts = row["det_counts"], row["exp_counts"]

        for t in sorted(set(det_counts) | set(exp_counts)):
            got, want = det_counts.get(t, 0), exp_counts.get(t, 0)
            if got == want:
                continue

            if got > want:
                extras = [(p, x) for p, x in (info["detections"] if info else [])
                          if x.excel_type == t and included_under(x, active)]
                failures.append(_mk(dn, row, "OVER_DETECT", t, got - want,
                                    extras[:3],
                                    f"detected {got} vs excel {want}"))
                continue

            deficit = want - got
            if not row["has_page"]:
                failures.append(_mk(dn, row, "NO_ANCHOR", t, deficit, [],
                                    "no page in the PDF carries this drawing number"))
                continue

            # Which evidence on the page explains the shortfall?
            held = [(p, x) for p, x in info["detections"]
                    if x.excel_type == t and "VENDOR_MARK_UNDEFINED" in x.rules_hit]
            scoped = [(p, x) for p, x in info["detections"]
                      if x.excel_type == t and not included_under(x, active)
                      and any(r in x.rules_hit for r in SCOPE_RULES)]
            mistyped = [(p, x) for p, x in info["detections"]
                        if x.anchor == t and "NOT_FIELD_INSTRUMENT" in x.rules_hit]
            unmapped = [(p, u) for p, u in info["unmapped"]]
            ambiguous = [(p, u) for p, u in info["unverified"]]

            if mistyped:
                failures.append(_mk(dn, row, "TYPE_MAPPING_MISS", t, deficit,
                                    mistyped[:3],
                                    f"'{t}' is on the drawing but the ruleset treats it "
                                    f"as NOT_FIELD_INSTRUMENT"))
            elif held:
                failures.append(_mk(dn, row, "VENDOR_MARK_UNDEFINED", t, deficit,
                                    held[:3],
                                    "bubble carries a mark this page's NOTES do not define; "
                                    "exclusion held"))
            elif scoped:
                failures.append(_mk(dn, row, "SCOPE_CONFLICT", t, deficit,
                                    scoped[:3],
                                    "excluded by a scope rule but the Excel keeps the row"))
            elif unmapped:
                failures.append(_mk(dn, row, "TYPE_MAPPING_MISS", t, deficit,
                                    unmapped[:3],
                                    "bubble holds a tag the anchor dictionary does not cover: "
                                    + ", ".join(sorted({u['token'] for _, u in unmapped})[:6])))
            elif ambiguous:
                failures.append(_mk(dn, row, "AMBIGUOUS_GEOMETRY", t, deficit,
                                    ambiguous[:3],
                                    "anchors found but no single enclosing bubble"))
            else:
                failures.append(_mk(dn, row, "NO_ANCHOR", t, deficit, [],
                                    f"no '{t}' anchor found in any bubble on this drawing"))
    return failures


def _mk(dn, row, kind, type_, count, samples, cause):
    out_samples = []
    for page_no, obj in samples:
        if hasattr(obj, "center"):
            out_samples.append({"page": page_no, "anchor": obj.anchor,
                                "center": [round(obj.center[0], 1), round(obj.center[1], 1)],
                                "rules_hit": list(obj.rules_hit)})
        else:
            out_samples.append({"page": page_no,
                                "anchor": obj.get("anchor") or obj.get("token"),
                                "center": obj.get("center")})
    return {"type": kind, "drawing_no": dn, "system": row["system"],
            "title": row["title"], "excel_type": type_, "count": count,
            "cause": cause, "samples": out_samples}


def system_table(drawing_rows):
    agg = collections.defaultdict(lambda: {"tp": 0, "det": 0, "exp": 0,
                                           "drawings": 0, "titles": []})
    for r in drawing_rows:
        a = agg[r["system"]]
        a["tp"] += r["tp"]
        a["det"] += r["detected"]
        a["exp"] += r["expected"]
        a["drawings"] += 1
        if r["title"]:
            a["titles"].append(r["title"])
    rows = []
    for sysname, a in agg.items():
        if a["exp"] == 0 and a["det"] == 0:
            continue
        rows.append({
            "system": sysname,
            "drawings": a["drawings"],
            "detected": a["det"],
            "expected": a["exp"],
            "tp": a["tp"],
            "recall": a["tp"] / a["exp"] if a["exp"] else float("nan"),
            "precision": a["tp"] / a["det"] if a["det"] else float("nan"),
            "title": a["titles"][0] if a["titles"] else "",
        })
    rows.sort(key=lambda r: (r["recall"] if r["expected"] else 9, -r["expected"]))
    return rows


def print_report(combo_scores, drawing_rows, failures, size_dist, per_page):
    print("\n" + "=" * 78)
    print("BUBBLE GEOMETRY — measured per page, not assumed")
    print("=" * 78)
    print(f"{'stadium (long x short)':<26} {'count':>7}   pages using it")
    for (lo, sh), n in size_dist.most_common(8):
        users = sum(1 for i in per_page.values() if (lo, sh) in i["sizes"])
        print(f"  {lo:>7.1f} x {sh:<12.1f} {n:>7}   {users}")
    p6 = (68.0, 22.6)
    other = sum(n for k, n in size_dist.items() if k != p6)
    print(f"\n  page-6 size {p6[0]}x{p6[1]} covers {size_dist.get(p6,0)} bubbles; "
          f"{other} bubbles use a different size")

    print("\n" + "=" * 78)
    print("VENDOR MARK NOTATION — where the page-6 rule can even apply")
    print("=" * 78)
    glyph_pages = [p for p, i in per_page.items() if i["notation"]["glyph_definitions"]]
    text_pages = [p for p, i in per_page.items()
                  if not i["notation"]["glyph_definitions"]
                  and i["notation"]["text_marks_in_drawing"]]
    none_pages = [p for p, i in per_page.items()
                  if not i["notation"]["glyph_definitions"]
                  and not i["notation"]["text_marks_in_drawing"]]
    print(f"  marks drawn as glyphs, legend readable : {len(glyph_pages):>3} pages  (rule works)")
    print(f"  marks typed as text '(*)' / '(**)'     : {len(text_pages):>3} pages  "
          f"(rule blind)  {sorted(text_pages)}")
    print(f"  no marks found at all                  : {len(none_pages):>3} pages")

    print("\n" + "=" * 78)
    print("RULE CONTRIBUTION — the core question")
    print("=" * 78)
    print(f"{'active rules':<16} {'detected':>9} {'excel':>7} {'TP':>6} {'recall':>9} {'precision':>10}")
    for name in ("none", "sct_only", "vendor_only", "both"):
        s = combo_scores[name]
        print(f"{name:<16} {s['detected']:>9} {s['expected']:>7} {s['tp']:>6} "
              f"{s['recall']:>8.1%} {s['precision']:>10.1%}")

    print("\n" + "=" * 78)
    print("PER-SYSTEM ACCURACY  (recall ascending)")
    print("=" * 78)
    rows = system_table(drawing_rows)
    print(f"{'sys':<5} {'dwg':>4} {'det':>5} {'excel':>6} {'TP':>5} {'recall':>8} {'prec':>8}  title")
    for r in rows:
        rc = f"{r['recall']:.1%}" if r["expected"] else "   n/a"
        pr = f"{r['precision']:.1%}" if r["detected"] else "   n/a"
        print(f"{r['system']:<5} {r['drawings']:>4} {r['detected']:>5} {r['expected']:>6} "
              f"{r['tp']:>5} {rc:>8} {pr:>8}  {r['title'][:44]}")

    print("\n" + "=" * 78)
    print("FAILURE TYPES")
    print("=" * 78)
    counts = collections.Counter(f["type"] for f in failures)
    rowsum = collections.Counter()
    for f in failures:
        rowsum[f["type"]] += f["count"]
    for t in FAILURE_TYPES:
        print(f"  {t:<24} cases={counts.get(t,0):>4}   rows={rowsum.get(t,0):>4}")


def write_failure_report(path, failures, drawing_rows, combo_scores,
                         size_dist, per_page, excel_total, excel_spare, excel):
    counts = collections.Counter(f["type"] for f in failures)
    rowsum = collections.Counter()
    for f in failures:
        rowsum[f["type"]] += f["count"]

    L = []
    L.append("# 전수 적용 실패 리포트 — Phase 0 spike 2b\n")
    L.append("page 6 에서 검증한 규칙을 **그대로** 전 페이지에 적용한 결과입니다. "
             "새 규칙은 추가하지 않았습니다.\n")

    s = combo_scores["both"]
    L.append(f"- Excel 총 {excel_total}행 중 `P&ID No.` 가 있는 "
             f"{excel_total - excel_spare}행이 조인 대상 (나머지 {excel_spare}행은 SPARE)\n")
    L.append(f"- 검출 {s['detected']}행 / 정답 {s['expected']}행 / 일치 {s['tp']}행\n")
    L.append(f"- **재현율 {s['recall']:.1%}, 정밀도 {s['precision']:.1%}**\n")

    # The join itself is a confound: some Excel drawings have no page, some
    # pages have no Excel rows, and three numbers are shared by two drawings.
    orphan_excel = [r for r in drawing_rows if not r["has_page"] and r["in_excel"]]
    orphan_page = [r for r in drawing_rows if r["has_page"] and not r["in_excel"]]
    shared = [r for r in drawing_rows if len(r["pages"]) > 1]
    L.append("\n## 조인 무결성 — 지표를 왜곡하는 요인\n")
    L.append(f"- Excel 에는 있으나 **PDF 에 해당 도면이 없는** 도면번호 "
             f"{len(orphan_excel)}건 / {sum(r['expected'] for r in orphan_excel)}행 "
             f"→ 검출 불가, 전부 재현율 손실로 계상됨")
    for r in orphan_excel:
        L.append(f"  - `{r['drawing_no']}` ({r['expected']}행)")
    L.append(f"- PDF 에는 있으나 **Excel 에 행이 없는** 도면 {len(orphan_page)}건 "
             f"/ 검출 {sum(r['detected'] for r in orphan_page)}행 → 전부 정밀도 손실로 계상됨")
    L.append(f"- **도면번호를 공유하는 페이지** {len(shared)}건 — 1:N 조인이라 "
             f"서로 다른 도면의 검출이 한 번호로 합산됩니다")
    for r in shared:
        systems = sorted({str(x["system"]) for x in excel.get(r["drawing_no"], [])})
        titles = [per_page[p]["title"] for p in r["pages"]]
        L.append(f"  - `{r['drawing_no']}` pages {r['pages']} 검출 {r['detected']} / "
                 f"Excel {r['expected']}")
        for pno, t in zip(r["pages"], titles):
            L.append(f"    - p{pno}: {t}")
        L.append(f"    - Excel SYSTEM: {', '.join(systems) if systems else '(행 없음)'}")

    L.append("\n## 규칙 조합별 성능\n")
    L.append("| 활성 규칙 | 검출 | 정답 | 일치 | 재현율 | 정밀도 |")
    L.append("|---|---|---|---|---|---|")
    label = {"none": "둘 다 끔", "sct_only": "SCT 만",
             "vendor_only": "벤더마크 만", "both": "둘 다"}
    for name in ("none", "sct_only", "vendor_only", "both"):
        c = combo_scores[name]
        L.append(f"| {label[name]} | {c['detected']} | {c['expected']} | {c['tp']} "
                 f"| {c['recall']:.1%} | {c['precision']:.1%} |")

    L.append("\n## 버블 치수 분포 (페이지별 실측)\n")
    L.append("| 스타디움 (긴변 x 짧은변) | 개수 | 사용 페이지 수 |")
    L.append("|---|---|---|")
    for (lo, sh), n in size_dist.most_common(8):
        users = sum(1 for i in per_page.values() if (lo, sh) in i["sizes"])
        L.append(f"| {lo} x {sh} | {n} | {users} |")

    glyph_pages = [p for p, i in per_page.items() if i["notation"]["glyph_definitions"]]
    text_pages = [p for p, i in per_page.items()
                  if not i["notation"]["glyph_definitions"]
                  and i["notation"]["text_marks_in_drawing"]]
    none_pages = [p for p, i in per_page.items()
                  if not i["notation"]["glyph_definitions"]
                  and not i["notation"]["text_marks_in_drawing"]]
    L.append("\n## 벤더 마크 표기 방식 (규칙이 적용 가능한 범위)\n")
    L.append("| 표기 방식 | 페이지 수 | 규칙 적용 | 페이지 |")
    L.append("|---|---|---|---|")
    L.append(f"| 벡터 글리프 (page 6 방식) | {len(glyph_pages)} | 가능 | {sorted(glyph_pages)} |")
    L.append(f"| 텍스트 `(*)` / `(**)` | {len(text_pages)} | **불가 — 규칙이 못 봄** | {sorted(text_pages)} |")
    L.append(f"| 마크 없음 | {len(none_pages)} | 해당 없음 | {sorted(none_pages)} |")

    L.append("\n## 실패 유형별 집계\n")
    L.append("| 유형 | 케이스 | 행 수 |")
    L.append("|---|---|---|")
    for t in FAILURE_TYPES:
        L.append(f"| `{t}` | {counts.get(t,0)} | {rowsum.get(t,0)} |")

    L.append("\n## 유형별 대표 사례\n")
    for t in FAILURE_TYPES:
        items = sorted([f for f in failures if f["type"] == t],
                       key=lambda f: -f["count"])
        L.append(f"\n### `{t}` — {counts.get(t,0)} 케이스 / {rowsum.get(t,0)} 행\n")
        if not items:
            L.append("해당 없음.\n")
            continue
        for f in items[:3]:
            L.append(f"**{f['drawing_no']}** (계통 `{f['system']}`, TYPE `{f['excel_type']}`, "
                     f"{f['count']}행)  ")
            L.append(f"{f['title']}  ")
            L.append(f"추정 원인: {f['cause']}  ")
            if f["samples"]:
                for smp in f["samples"]:
                    rules = ("/".join(smp["rules_hit"]) if smp.get("rules_hit") else "-")
                    L.append(f"- p{smp['page']} `{smp['anchor']}` @ {smp['center']}  규칙: {rules}")
            L.append("")

    L.append("\n## 계통별 정확도 (재현율 오름차순)\n")
    L.append("| 계통 | 도면 | 검출 | Excel | 일치 | 재현율 | 정밀도 | 대표 도면명 |")
    L.append("|---|---|---|---|---|---|---|---|")
    for r in system_table(drawing_rows):
        rc = f"{r['recall']:.1%}" if r["expected"] else "n/a"
        pr = f"{r['precision']:.1%}" if r["detected"] else "n/a"
        L.append(f"| `{r['system']}` | {r['drawings']} | {r['detected']} | {r['expected']} "
                 f"| {r['tp']} | {rc} | {pr} | {r['title'][:40]} |")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(L) + "\n", encoding="utf-8")


if __name__ == "__main__":
    raise SystemExit(main())
