"""Excel 출력이 **지금의 데이터**를 담는지 화면에서 끝까지 확인한다 (hotfix15).

    python3 spike/excel_reflect_check.py out/regression_3p/UAD-DXF.json out/round56/excel_reflect [앱뿌리]

순서 — 실제 사람이 하는 그대로:
    ① 검토 완료 체크 (스냅샷)
    ② 그 **뒤에** 편집: SCOPE · Description · Q'ty 각 한 행, 한 행 삭제, 마크업 한 행 추가
    ③ Excel 단추 → 내려받은 zip 을 풀어 FIELD 파일을 연다
    ④ 편집 값 셋이 칸에 있는가 · 지운 행이 없는가 · 추가 행이 있는가

양식은 이 저장소에 없다(발주처 파일) — 설정(`excel.columns`)이 말하는 자리만 가진
**시험용 빈 양식**을 만들어 쓴다.  실 DB 를 열지 않는다 (끝에서 sha256 대조).
세 번째 인자로 옛 코드 뿌리를 주면 같은 절차로 고치기 전을 잰다.
"""
from __future__ import annotations
import hashlib, io, json, os, socket, subprocess, sys, tempfile, time, zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
src = Path(sys.argv[1] if len(sys.argv) > 1 else "out/regression_3p/UAD-DXF.json")
OUT = Path(sys.argv[2] if len(sys.argv) > 2 else "out/round56/excel_reflect")
APP = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else ROOT
OUT.mkdir(parents=True, exist_ok=True)
FOUND: list[str] = []
REAL = ROOT / "app" / "_data" / "app.db"
real_before = hashlib.sha256(REAL.read_bytes()).hexdigest() if REAL.exists() else ""


def note(s):
    print(s, flush=True); FOUND.append(s)


def free_port():
    s = socket.socket(); s.bind(("127.0.0.1", 0)); p = s.getsockname()[1]; s.close(); return p


data = Path(tempfile.mkdtemp(prefix="xlsxui-"))
os.environ["PID_DATA_DIR"] = str(data)
from app import db, revisions                                    # noqa: E402
import openpyxl                                                  # noqa: E402
import yaml                                                      # noqa: E402

cfg = yaml.safe_load((ROOT / "config" / "project_alnouf1.yaml").read_text(encoding="utf-8"))
ex = cfg["excel"]
(data / "templates").mkdir(parents=True, exist_ok=True)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = ex["sheet"]
for name, col in ex["columns"].items():
    ws.cell(ex["first_data_row"] - 1, int(col)).value = name.upper()
ws.cell(ex["first_data_row"], 41).value = None
wb.save(data / "templates" / "FIELD.xlsx")
COLS = {k: int(v) for k, v in ex["columns"].items()}

blob = json.loads(src.read_text()); result = blob.get("result", blob)
con = db.connect(data / "app.db")
job = "xlsxrefl0001"
pack = ROOT / "data" / "uad_dxf.zip"
sha = hashlib.sha256(pack.read_bytes()).hexdigest()
revisions.create_project(data, "UAD-DXF")
con.execute("INSERT INTO job (id, pdf_name, pdf_sha256, pdf_path, created_at, status,"
            " progress, message, fingerprint, project, revision, input_kind)"
            " VALUES (?,?,?,?,?,'done',1.0,'',?,?,?,'DXF')",
            (job, pack.name, sha, str(pack.resolve()), time.time(),
             result.get("fingerprint", ""), "UAD-DXF", "A"))
con.commit(); db.store_result(con, job, result); con.commit(); con.close()
revisions.record_revision(data, "UAD-DXF", "A", job_id=job, pdf_name=pack.name,
                          compared_with="", result={"counts": {}, "states": {}, "radii": {},
                                                    "deleted_candidates": []})

port = free_port()
env = dict(os.environ, PID_DATA_DIR=str(data), PYTHONPATH=str(APP))
srv = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app",
                        "--host", "127.0.0.1", "--port", str(port)],
                       cwd=str(APP), env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
try:
    import urllib.request
    for _ in range(90):
        try:
            urllib.request.urlopen(f"http://127.0.0.1:{port}/home", timeout=2); break
        except Exception:
            time.sleep(1)
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        exe = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
        br = pw.chromium.launch(executable_path=exe if Path(exe).exists() else None)
        ctx = br.new_context(viewport={"width": 1700, "height": 1000}, accept_downloads=True)
        pg = ctx.new_page()
        errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)))
        pg.on("dialog", lambda d: d.accept(""))
        pg.goto(f"http://127.0.0.1:{port}/")
        pg.wait_for_timeout(2500)
        pg.evaluate("() => document.querySelectorAll('details.pjt').forEach(d => d.open = true)")
        pg.wait_for_timeout(400)
        pg.query_selector("a.revrow").click()
        pg.wait_for_timeout(8000)
        # ① 검토 완료 체크
        pg.check("#gate-check")
        pg.wait_for_function("() => window.__rev", timeout=60000)
        rev1 = pg.evaluate("() => window.__rev")
        # ② 체크 뒤의 편집 — 화면이 쓰는 fetch 로 (래퍼를 지나게)
        keys = pg.evaluate("""() => S.rows.filter(r => r.tab === 'FIELD' && !r.removed && r.values.tag_no)
                                     .slice(0, 4).map(r => ({key: r.key, tag: r.values.tag_no}))""")
        edits = [("scope", "VENDOR(TEST SUPPLIER)"), ("description", "EDITED BY REVIEWER"),
                 ("qty", 7)]
        for (field, value), k in zip(edits, keys[:3]):
            pg.evaluate("""async ([key, field, value]) => {
                const r = await fetch(`/jobs/${S.job.id}/rows/${key}`, {method: 'PATCH',
                  headers: {'Content-Type': 'application/json'},
                  body: JSON.stringify({field, value, author: 'tester'})});
                return r.ok; }""", [k["key"], field, value])
        gone = keys[3]
        pg.evaluate("""async (key) => { const r = await fetch(
            `/jobs/${S.job.id}/rows/${key}?reason_class=FALSE_POSITIVE&author=tester&exclude=true`,
            {method: 'DELETE'}); return r.ok; }""", gone["key"])
        pg.evaluate("""async () => { const r = await fetch(`/jobs/${S.job.id}/rows`, {method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({page_no: 12, tab: 'FIELD', drawing_no: '1A5J-00EKD00-M05-0002',
              rect: [471, 323, 499, 335], author: 'tester', reason_class: 'MISSING',
              values: {type: 'PDIT', tag_no: 'MARKUP-TAG-01', system: 'P&ID FOR FUEL OIL',
                       qty: 1, scope: 'SCT'}})}); return r.ok; }""")
        pg.wait_for_timeout(800)
        label = pg.evaluate("() => document.querySelector('#excel').textContent")
        note(f"① 검토 완료 rev {rev1} · ② 편집 3 · 삭제 1 · 마크업 추가 1 뒤 단추: {label!r}")
        # ③ Excel
        with pg.expect_download(timeout=120000) as dl:
            pg.click("#excel")
        path = OUT / "download.zip"
        dl.value.save_as(str(path))
        rev2 = pg.evaluate("() => window.__rev")
        note(f"③ 내려받음 — 스냅샷 rev {rev1} → {rev2}")
        note("콘솔 오류: " + (" | ".join(errs[:4]) if errs else "없음"))
        br.close()
finally:
    srv.terminate(); srv.wait(timeout=20)

with zipfile.ZipFile(path) as z:
    name = next(n for n in z.namelist() if n.startswith("field_"))
    wb = openpyxl.load_workbook(io.BytesIO(z.read(name)))
ws = wb[ex["sheet"]]
by_tag = {}
for r in range(ex["first_data_row"], ws.max_row + 1):
    t = ws.cell(r, COLS["tag_no"]).value
    if t:
        by_tag[str(t)] = r
checks = []
for (field, value), k in zip(edits, keys[:3]):
    r = by_tag.get(k["tag"])
    got = ws.cell(r, COLS[field]).value if r else None
    checks.append((f"{field} 편집 ({k['tag']})", got == value, got))
checks.append((f"삭제한 행 ({gone['tag']}) 이 없다", gone["tag"] not in by_tag, gone["tag"] in by_tag))
r = by_tag.get("MARKUP-TAG-01")
checks.append(("마크업 추가 행이 있다", r is not None,
               [ws.cell(r, COLS[c]).value for c in ("type", "system", "qty", "scope")] if r else None))
for what, ok, got in checks:
    note(f"④ {'통과' if ok else '실패'} — {what}: {got!r}")
note(f"   FIELD 파일 행 {len(by_tag)} (태그 있는 행)")
real_after = hashlib.sha256(REAL.read_bytes()).hexdigest() if REAL.exists() else ""
note(f"실 DB sha256 — 같은가: {real_before == real_after}")
(OUT / "README.md").write_text("# hotfix15 — Excel 이 지금의 데이터를 담는가\n\n"
                               + "\n".join("- " + s for s in FOUND) + "\n", encoding="utf-8")
print("→", OUT)
sys.exit(0 if all(ok for _w, ok, _g in checks) else 1)
