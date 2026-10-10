"""hotfix74 — 같은 스냅샷의 [Excel 출력] 을 여럿이 같은 순간에 누른다 (8동시 × 3회).

    python3 spike/excel_race.py <data_dir 사본>

사본 데이터에 시험용 양식(설정의 시트 이름·열)을 올리고 스냅샷을 찍은 뒤, 내려받은 zip 의 xlsx 가 전부 열리는지 본다.
잠금 전: 24회 중 17~22회 500 (BadZipFile — 한 요청이 다른 요청의 반쯤 쓴 파일을 읽음) · 잠금 뒤: 0.
"""
import os, sys, io, zipfile, json, threading, sqlite3
os.environ.update(PID_DATA_DIR=sys.argv[1], PID_PAGE_WARM="0", PID_RENDER_POOL="0", PID_ANALYSIS_INPROCESS="1")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from fastapi.testclient import TestClient
from app import main
import openpyxl
c = TestClient(main.app, raise_server_exceptions=False)
print("revisions", [dict(r) for r in main.CON.execute("select id, job_id from revision order by id desc limit 3")])
# make templates for all deliverables (empty workbooks)
from app import excel_out
for kind in excel_out.DELIVERABLES:
    sheet, first_row, cols = excel_out._cfg_columns(main.CFG, excel_out.DELIVERABLES[kind]["config"])
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = sheet
    for name, col in (cols.items() if isinstance(cols, dict) else []):
        try: ws.cell(max(1, first_row - 1), int(col), name)
        except Exception: pass
    b = io.BytesIO(); wb.save(b)
    r = c.post("/templates", data={"kind": kind}, files={"file": (f"{kind}.xlsx", b.getvalue())})
    print(kind, r.status_code)
rev = main.CON.execute("select id from revision order by id desc limit 1").fetchone()
if rev is None:
    job = main.CON.execute("select id from job where status='done' order by created_at desc").fetchone()[0]
    r = c.post(f"/jobs/{job}/snapshot", json={"label": "race"}); print("snap", r.status_code, r.text[:200])
    rev = main.CON.execute("select id from revision order by id desc limit 1").fetchone()
RID = rev[0]
print("rid", RID)
res = []
def go(i):
    r = c.get(f"/revisions/{RID}/excel")
    if r.status_code != 200:
        res.append((i, r.status_code, r.text[:200])); return
    z = zipfile.ZipFile(io.BytesIO(r.content)); bad = []
    for n in z.namelist():
        if n.endswith(".xlsx"):
            try: openpyxl.load_workbook(io.BytesIO(z.read(n)))
            except Exception as e: bad.append(f"{n}: {type(e).__name__}")
    res.append((i, 200, bad))
for rnd in range(3):
    ts = [threading.Thread(target=go, args=(i,)) for i in range(8)]
    [t.start() for t in ts]; [t.join() for t in ts]
for x in res: print(x)
print("broken", sum(1 for x in res if x[1] != 200 or x[2]))
