#!/usr/bin/env python3
"""판독 결과(extraction.json)를 지정된 Excel 포맷으로 저장한다.

원본 템플릿 파일을 복사해서 데이터 행만 교체하는 방식이라 헤더 3줄(프로젝트/RFQ/
DESCRIPTION), 병합된 2단 헤더, 열 너비, 셀 서식이 템플릿 그대로 유지된다.

컬럼은 excel_format.json의 source에 따라 채워진다:
  drawing      판독 결과에서
  typical      INST. TYPICAL TYPE으로 typicals.json에서 조회
  auto_index   1부터 증가
  design_table / blank / review   비워 둠 (사람 또는 후속 조인)

리뷰 UI가 읽을 result.json도 함께 만든다(값 + 확신도 + 근거 + 도면 페이지).

사용:
  python3 scripts/build_excel.py outputs/run_20260815_2230/extraction.json
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("extraction", type=Path, help="extract_instruments.py가 만든 extraction.json")
    ap.add_argument("--template", type=Path,
                    help="템플릿 xlsx (기본: inputs/ 안의 excel_format.json이 기록한 파일)")
    ap.add_argument("--out", type=Path, help="출력 xlsx 경로 (기본: extraction.json 옆)")
    args = ap.parse_args()

    fmt = json.loads((ROOT / "config/excel_format.json").read_text(encoding="utf-8"))
    typicals = json.loads((ROOT / "config/typicals.json").read_text(encoding="utf-8"))
    ext = json.loads(args.extraction.read_text(encoding="utf-8"))

    template = args.template or (ROOT / "inputs" / fmt["template_file"])
    if not template.exists():
        print(f"[!] 템플릿을 찾을 수 없습니다: {template}\n"
              f"    --template 으로 경로를 지정하세요.", file=sys.stderr)
        return 1

    out_xlsx = args.out or args.extraction.parent / "instrument_list.xlsx"
    shutil.copy2(template, out_xlsx)

    wb = load_workbook(out_xlsx)
    ws = wb[fmt["sheet"]]
    start = fmt["data_start_row"]

    # 템플릿 첫 데이터 행의 서식을 원본으로 확보한 뒤 기존 데이터를 지운다.
    styles = {c["col"]: copy(ws[f"{c['col']}{start}"]._style) for c in fmt["columns"]}
    last = ws.max_row
    if last >= start:
        ws.delete_rows(start, last - start + 1)

    rows = ext["instruments"]
    unknown_typicals: set[str] = set()

    for i, rec in enumerate(rows):
        r = start + i
        x = (rec.get("inst_typical_type") or "").strip()
        typ = typicals.get(x, {})
        if x and x not in typicals:
            unknown_typicals.add(x)

        for c in fmt["columns"]:
            cell = ws[f"{c['col']}{r}"]
            cell._style = copy(styles[c["col"]])
            src, key = c["source"], c["key"]
            if src == "auto_index":
                cell.value = i + 1
            elif src == "drawing":
                cell.value = coerce(rec.get(key, ""), key)
            elif src == "typical":
                cell.value = typ.get(key, "")
            else:
                cell.value = None       # design_table / blank / review

    wb.save(out_xlsx)

    # ── 리뷰 UI용 result.json ────────────────────────────────────
    page_of = {p["page"]: p for p in ext.get("pages", [])}
    result = {
        "run_id": ext["run_id"],
        "created_at": ext["created_at"],
        "model": ext["model"],
        "sheet": fmt["sheet"],
        "xlsx": out_xlsx.name,
        "columns": [{k: c[k] for k in ("col", "key", "label", "source")} for c in fmt["columns"]],
        "instrument_types": fmt["instrument_types"],
        "systems": fmt["systems"],
        "type_to_typical_types": fmt["type_to_typical_types"],
        "pages": ext.get("pages", []),
        "excluded": ext.get("excluded", []),
        "review_findings": ext.get("review_findings", []),
        "rows": [],
    }
    for i, rec in enumerate(rows):
        x = (rec.get("inst_typical_type") or "").strip()
        typ = typicals.get(x, {})
        values = {}
        for c in fmt["columns"]:
            if c["source"] == "auto_index":
                values[c["key"]] = i + 1
            elif c["source"] == "drawing":
                values[c["key"]] = rec.get(c["key"], "")
            elif c["source"] == "typical":
                values[c["key"]] = typ.get(c["key"], "")
            else:
                values[c["key"]] = ""
        result["rows"].append({
            "index": i,
            "values": values,
            "confidence": rec.get("confidence", ""),
            "source_tokens": rec.get("source_tokens", ""),
            "page": rec.get("_page"),
            "drawing_title": (page_of.get(rec.get("_page")) or {}).get("title", ""),
        })

    out_json = out_xlsx.with_name("result.json")
    out_json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"[✓] {out_xlsx}   ({len(rows)}행)")
    print(f"[✓] {out_json}   (리뷰 UI 입력)")
    filled = [c["key"] for c in fmt["columns"] if c["source"] in ("drawing", "typical", "auto_index")]
    blank = [c["label"] or c["key"] for c in fmt["columns"] if c["source"] == "design_table"]
    print(f"[i] 채운 컬럼: {filled}")
    print(f"[i] 비워 둔 컬럼(Design Table 조인 필요, {len(blank)}개): {blank[:6]} …")
    if unknown_typicals:
        print(f"[!] typicals.json에 없는 INST. TYPICAL TYPE: {sorted(unknown_typicals)} → 사양 컬럼이 비었습니다")
    low = [r for r in result["rows"] if r["confidence"] == "low"]
    if low:
        print(f"[!] 확신도 low 행 {len(low)}개 — 리뷰 UI에서 우선 확인하세요")
    print(f"\n다음: review-ui/index.html 을 열고 {out_json.name} 을 불러와 검토하세요")
    return 0


def coerce(value, key: str):
    """Q'ty처럼 숫자여야 하는 값은 숫자로 넣는다."""
    s = "" if value is None else str(value).strip()
    if key == "qty" and s.isdigit():
        return int(s)
    return s


if __name__ == "__main__":
    raise SystemExit(main())
