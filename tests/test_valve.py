"""Valve Datasheet Review (app/valve) — extraction, grouping, revision compare, Valve List fill, API.

The client's datasheets / TBE are not in the repository.  These tests draw a small ruled datasheet with
PyMuPDF in the same shape (row numbers · label cells · value cells · title block · NOTES on sheet 2) and
a matching appendix workbook.  With VALVE_SAMPLES=<folder with Datasheet.pdf> the real file is read too.
"""
from __future__ import annotations

import io
import os
import time
from pathlib import Path

import openpyxl
import pymupdf
import pytest

from app.valve import api, extract, schema, store, tbe

ROWS = [  # (no, [labels], [values]) — consecutive like the real form; fields are found by label, not number  — two labels pair with left / right value halves
    (1, ["Tag No.", "Item No."], ["00XV 921", "VTA"]),
    (2, ["Manufacturer", "Model"], ["VTA", "VTA"]),
    (3, ["Service"], ["AGAS FR N2 BOOT LVL VLV"]),
    (4, ["Piping class"], ["1C0K8"]),
    (5, ["Operating Pressure", "Design Pressure"], ["", ""]),
    (6, ["Body Size", "Body Type"], ['2"', "Ball (Full Bore) - Floating"]),
    (7, ["Connections", "Rating"], ["Flanged RF (ASME B16.5)", "150#"]),
    (8, ["Body material"], ["316/316L SS A351-CF8M/CF3M"]),
    (9, ["Seat Material", "Seal Material"], ["RTFE", "Graphite (VTA)"]),
    (10, ["Bolting / Gasket"], ["Note-1"]),
    (11, ["Packing Type"], ["General note - 8"]),
    (12, ["Manufacturer", "Model No."], ["Bettis", "G-01"]),
]
X_NUM, X_L1, X_L2, X_V, X_MID, X_R = 122, 136, 222, 313, 438, 572


def _sheet(doc, rev, sheet, of, tag_rows, notes=None, values=None):
    pg = doc.new_page(width=595, height=842)
    pg.insert_text((400, 110), "ON / OFF VALVE", fontsize=14)
    pg.insert_text((420, 125), "00-09-IN-DTS-00003", fontsize=8)
    pg.insert_text((410, 140), f"SHEET {sheet} OF {of}", fontsize=8)
    pg.insert_text((410, 150), "REV.", fontsize=8)
    pg.insert_text((470, 150), rev, fontsize=8)
    pg.insert_text((410, 160), "DATE", fontsize=8)
    pg.insert_text((470, 160), "19/9/2026", fontsize=8)
    y = 200
    caption_top = y
    for no, labels, vals in tag_rows:
        if values and no in values:
            vals = values[no]
        pg.draw_rect(pymupdf.Rect(43, y, X_R, y + 12), width=0.5)
        xs = [X_NUM, X_L1, X_V, X_R] + ([X_L2, X_MID] if len(labels) == 2 else [])
        for x in xs:
            pg.draw_line((x, y), (x, y + 12), width=0.5)
        pg.insert_text((X_NUM + 4, y + 9), str(no), fontsize=7)
        lx = [X_L1 + 3, X_L2 + 3]
        for i, lab in enumerate(labels):
            pg.insert_text((lx[i], y + 9), lab, fontsize=7)
        vx = [X_V + 3, X_MID + 3] if len(labels) == 2 else [X_V + 3]
        for i, v in enumerate(vals):
            if v:
                pg.insert_text((vx[i], y + 9), v, fontsize=7)
        y += 12
    # caption bands left of the numbers: the last row sits in its own ACTUATOR band, as on the real form
    split = y - 12 if len(tag_rows) > 4 else y
    for a in (caption_top, split, y):
        pg.draw_line((43, a), (X_NUM, a), width=0.5)
    pg.insert_text((50, (caption_top + split) / 2), "GENERAL" if sheet == 1 else "ACCESSORIES", fontsize=7)
    if split < y:
        pg.insert_text((50, split + 9), "ACTUATOR" if sheet == 1 else "ACCESSORIES", fontsize=7)
    if notes:
        pg.insert_text((57, y + 30), "NOTES :", fontsize=8)
        ny = y + 45
        for i, t in enumerate(notes, 1):
            pg.insert_text((49, ny), f"{i}. {t}", fontsize=7)
            ny += 12
    return pg


def make_pdf(rev="A", values=None) -> bytes:
    doc = pymupdf.open()
    _sheet(doc, rev, 1, 2, ROWS, values=values)
    acc = [(52 + i, ["Filter Regulator" if i == 0 else f"Tubing {i}"], ["NA"]) for i in range(9)]
    _sheet(doc, rev, 2, 2, acc, notes=[
        "Gasket material: FG-316 SS PERF TANG INSERT, CL150. Stud Bolt material: Continuous thread, CR-MO A193-B7.",
        "Position Transmitter with 24V DC."])
    return doc.tobytes()


@pytest.fixture()
def pdf_path(tmp_path):
    p = tmp_path / "ds.pdf"
    p.write_bytes(make_pdf())
    return p


def test_tag_normalisation_and_type():
    assert schema.norm_tag("00XV 921") == schema.norm_tag("00XV-921") == "00XV921"
    assert schema.valve_type("98ESDV -004") == "ESDV"
    assert schema.valve_type("315MOV980004") == "MOV"


def test_revision_order():
    assert sorted(["1", "A", "0", "B", "2"], key=store.rev_key) == ["A", "B", "0", "1", "2"]


def test_extract_fields_positions_notes(pdf_path):
    res = extract.read_pdf(str(pdf_path))
    dss = extract.group_datasheets(res)
    assert len(dss) == 1
    d = dss[0]
    assert d["tag_key"] == "00XV921" and d["rev"] == "A" and d["pages"] == [1, 2]
    f = d["fields"]
    assert f["body_size"]["value"] == '2"' and f["body_size"]["page"] == 1 and f["body_size"]["bbox"]
    assert f["rating"]["value"] == "150#"
    assert f["seat_material"]["value"] == "RTFE" and f["seal_material"]["value"] == "Graphite (VTA)"
    assert f["manufacturer"]["status"] == "placeholder"          # VTA is kept, flagged, not guessed
    assert f["op_pressure"]["status"] == "blank"                 # empty cell -> Not Found, never invented
    assert f["act_manufacturer"]["value"] == "Bettis"            # same label in a later group, no clash
    assert [n["no"] for n in d["notes"]] == [1, 2]
    assert f["bolting_gasket"]["ref"]["kind"] == "note" and "FG-316" in f["bolting_gasket"]["ref"]["text"]


def test_compare_revisions():
    def full(rev, values):
        p = io.BytesIO(make_pdf(rev, values))
        tmp = Path(os.environ.get("TMPDIR", "/tmp")) / f"valve_cmp_{rev}.pdf"
        tmp.write_bytes(p.getvalue())
        d = extract.group_datasheets(extract.read_pdf(str(tmp)))[0]
        return {"fields": [dict(v, key=k, edited_value=None) for k, v in d["fields"].items()], "notes": d["notes"]}
    a = full("A", None)
    b = full("B", {6: ['3"', "Ball (Full Bore) - Floating"], 9: ["PEEK", "Graphite (VTA)"]})
    rows = {r["key"]: r for r in api.compare(a, b)}
    assert rows["body_size"]["change"] == "Modified" and rows["body_size"]["new"] == '3"'
    assert rows["seat_material"]["change"] == "Modified"
    assert rows["rating"]["change"] == "Unchanged"
    assert rows["body_size"]["new_bbox"]                         # the screen can point at it


def _template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TBE Appendix-XV"
    ws["A1"] = "TBE APPENDIX FOR ON-OFF BALL VALVE (XVs - Pneumatic Actuators)"
    hdr = ["S. No.", "Tag Number", "Vendor", "Valve Body", None, None, None]
    sub = [None, None, "Supplier", "Mfr. / Model", "Piping Class", "Body Size", "Seat Material"]
    for c, v in enumerate(hdr, 1):
        ws.cell(2, c, v)
    for c, v in enumerate(sub, 1):
        ws.cell(4, c, v)
    ws.merge_cells("D2:G3")
    rows = [(1, "00XV921", "Requirement", "VTC", "1C0K8", '2"', "PTFE"), (None, None, "AMPO", None, None, None, None)]
    for r, vals in enumerate(rows, 5):
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_valve_list_fill_keeps_format_and_marks_changes(pdf_path):
    d = extract.group_datasheets(extract.read_pdf(str(pdf_path)))[0]
    d["document_id"] = 1
    new = dict(d, tag="19XV 122", tag_key="19XV122")
    out, rep = tbe.build(_template(), [d, new], {1: "ds.pdf"})
    wb = openpyxl.load_workbook(io.BytesIO(out))
    ws = wb["TBE Appendix-XV"]
    assert rep["updated"] == ["00XV921"] and rep["new"] == ["19XV122"]
    seat = [c for c in rep["changes"] if c["column"] == "seat material"]
    assert seat and seat[0]["previous"] == "PTFE" and seat[0]["new"] == "RTFE"
    assert ws.cell(5, 7).value == "RTFE" and ws.cell(5, 7).comment and "Previous: PTFE" in ws.cell(5, 7).comment.text
    assert ws.cell(5, 4).value == "VTC"                         # 'VTA / VTA' says less than the list: kept
    assert "D2:G3" in [str(m) for m in ws.merged_cells.ranges]  # template structure untouched
    tags = [ws.cell(r, 2).value for r in range(5, ws.max_row + 1)]
    assert "19XV122" in tags                                     # new tag appended to the XV sheet
    assert "Datasheet Notes" in wb.sheetnames


def test_api_upload_review_and_export(tmp_path, monkeypatch, pdf_path):
    monkeypatch.setenv("PID_DATA_DIR", str(tmp_path / "data"))
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    app = FastAPI()
    app.include_router(api.router)
    c = TestClient(app)
    H = {"X-Author": "%ED%85%8C%EC%8A%A4%ED%8A%B8"}
    pid = c.post("/api/valve/projects", json={"name": "CO2 PJT"}).json()["id"]
    j = c.post(f"/api/valve/projects/{pid}/upload", data={"kind": "datasheet"},
               files=[("files", ("ds.pdf", pdf_path.read_bytes(), "application/pdf"))], headers=H).json()
    for _ in range(100):
        if c.get(f"/api/valve/jobs/{j['id']}").json()["status"] == "done":
            break
        time.sleep(0.1)
    s = c.get(f"/api/valve/projects/{pid}/summary").json()
    assert s["kpi"]["tags"] == 1
    ds = c.get(f"/api/valve/datasheets/{s['tags'][0]['id']}").json()
    f = next(x for x in ds["fields"] if x["key"] == "body_size")
    ds2 = c.put(f"/api/valve/datasheets/{ds['id']}/fields/{f['id']}", json={"value": '3"'}, headers=H).json()
    g = next(x for x in ds2["fields"] if x["key"] == "body_size")
    assert g["value"] == '2"' and g["edited_value"] == '3"' and g["edited_by"] == "테스트"   # read value kept
    assert ds2["history"][0]["action"] == "edit"
    r = c.post(f"/api/valve/datasheets/{ds['id']}/comments", json={"text": "check", "page": 1, "bbox": [300, 300, 400, 320]},
               headers=H).json()
    assert r["comments"][0]["bbox"] == [300, 300, 400, 320]
    m = c.post("/api/valve/markup.pdf", json={"document_id": ds["document_id"], "pages": ds["pages"],
               "items": [{"page": 1, "bbox": [300, 300, 400, 320], "category": "comment", "text": "check"}]})
    assert m.status_code == 200 and len(list(pymupdf.open(stream=m.content, filetype="pdf")[0].annots())) == 1
    assert c.get(f"/api/valve/documents/{ds['document_id']}/page/1.png?zoom=1").status_code == 200
    x = c.get(f"/api/valve/projects/{pid}/valvelist.xlsx")
    assert x.status_code == 200 and "Valve List (tool format)" in openpyxl.load_workbook(io.BytesIO(x.content)).sheetnames


@pytest.mark.skipif(not os.environ.get("VALVE_SAMPLES"), reason="client sample files are not in the repository")
def test_real_datasheet_sample():
    """VALVE_SAMPLES=<folder>/Datasheet.pdf (CO2 PJT, 456 pages): 164 datasheets, notes on every one."""
    res = extract.read_pdf(str(Path(os.environ["VALVE_SAMPLES"]) / "Datasheet.pdf"))
    dss = extract.group_datasheets(res)
    assert len(dss) == 164
    assert all(d["notes"] for d in dss)
    review = sum(1 for d in dss for f in d["fields"].values() if f["status"] == "review")
    assert review <= 80


# ── material equivalence · PMS · issues · contractor exports (round 2) ─────────────────────────

from app.valve import exports, issues, material, pms  # noqa: E402


@pytest.mark.parametrize("req, ds, want", [
    ("ASTM A351 CF8M", "316/316L SS A351-CF8M/CF3M", material.COMPLIANT),     # user report 2026-10-31
    ("316 SS / HF", "SS 316/HF", material.COMPLIANT),                          # user report 2026-10-31
    ("SS 316/HF", "316 SS / HF", material.COMPLIANT),
    ("CS A216-WCB", "Carbon Steel ASTM A216 Grade WCB", material.COMPLIANT),
    ("RTFE", "RTFE ST", material.COMPLIANT),
    ("ASTM A351 CF8M", "Carbon Steel ASTM A216 Grade WCB", material.NON),
    ("RTFE", "PEEK", material.NON),
    ("316/316L SS A351-CF8M/CF3M", "ASTM A351 CF8M", material.COMPLIANT),     # the other direction, same material
    ("ASTM A351 CF8M", "A182 F316", material.REVIEW),                          # same family: engineer decides
    ("316 SS/HF", "SS 316", material.REVIEW),                                  # hard-facing not stated
    ("316 SS / HF", "VTA", material.REVIEW),
])
def test_material_compare(req, ds, want):
    assert material.compare(req, ds)[0] == want


def _pms_pdf(path):
    doc = pymupdf.open()
    pg = doc.new_page(width=595, height=842)
    pg.insert_text((10, 40), "PIPING MATERIAL SPECIFICATION - QN AND W15 AREA", fontsize=8)
    pg.insert_textbox(pymupdf.Rect(10, 60, 400, 200), "Class: 1C0K8\nClass 150, 316/316L, SS, RF, NACE\n"
                      "NACE Requirement: MR0175/ISO 15156\nFlange Class: CLASS 150\nMaterial: STAINLESS STEEL\n"
                      "Valve Trim: TRIM 12 (316 SS/HF)", fontsize=7)
    pg.insert_textbox(pymupdf.Rect(10, 300, 470, 360), "PVBSED15RF02TD08U703\n0.5\n-\n2\nBALL VALVE, 316/316L SS "
                      "A351-CF8M/CF3M, CL150, RF, B16.10, FLOATING, FULL PORT, 316 SS BALL/STEM, RTFE ST, FIRE TESTED TO "
                      "API 607, NACE MR0175/ISO 15156", fontsize=6)
    doc.save(str(path))


def test_pms_read_and_check(tmp_path):
    p = tmp_path / "pms.pdf"
    _pms_pdf(p)
    P = pms.read_pms(str(p))
    c = P["areas"]["QN"]["1C0K8"]
    assert c["rating"] == 150 and "316 SS/HF" in c["trim"] and c["items"][0]["to"] == 2.0
    F = {k: {"value": v, "page": 1, "bbox": [0, 0, 1, 1]} for k, v in {
        "piping_class": "1C0K8", "body_size": '2"', "body_type": "Ball (Full Bore) - Floating", "rating": "150#",
        "nace": "Yes", "body_material": "ASTM A351 CF8M", "ball_material": "SS 316/HF", "stem_material": "316 SS / HF",
        "seat_material": "PEEK"}.items()}
    rows = {r["item"]: r for r in pms.check(F, P, "QN")}
    assert rows["Pressure Class"]["result"] == material.COMPLIANT
    assert rows["Body Material"]["result"] == material.COMPLIANT          # CF8M inside CF8M/CF3M
    assert rows["Trim (Ball) Material"]["result"] == material.COMPLIANT   # notation only
    assert rows["Seat Material"]["result"] == material.NON and rows["Seat Material"]["pms_page"] == 1
    F["rating"]["value"] = "300#"
    assert {r["item"]: r for r in pms.check(F, P, "QN")}["Pressure Class"]["result"] == material.NON


def _ds_for_issues():
    fields = [
        {"key": "tag", "value": "00XV 921", "page": 1, "bbox": [1, 1, 2, 2], "row": 1, "display": "Tag"},
        {"key": "line_no", "value": 'PW2409631-2"(1C0K8)-N', "page": 1, "bbox": [1, 3, 2, 4], "row": 5, "display": "Line No."},
        {"key": "piping_class", "value": "1C3AS1", "page": 1, "bbox": [1, 5, 2, 6], "row": 9, "display": "Piping class"},
        {"key": "body_size", "value": "", "page": 1, "bbox": [1, 7, 2, 8], "row": 21, "display": "Valve Size"},
        {"key": "rating", "value": "HOLD", "page": 1, "bbox": [1, 9, 2, 10], "row": 22, "display": "Rating"},
        {"key": "elec_connection", "value": "M20 (General Note-46)", "page": 1, "bbox": [1, 11, 2, 12], "row": 49,
         "display": "Electrical Connection"},
        {"key": "raw.13", "value": "Liquid", "page": 1, "bbox": [1, 13, 2, 14], "row": 13, "status": "review", "display": "Row 13"},
    ]
    return {"fields": [dict(f, edited_value=None, status=f.get("status", "ok")) for f in fields], "notes": [],
            "pages": [1, 2], "sheets": [[1, 1, 2], [2, 2, 2]], "doc_no": "00-09-IN-DTS-00003", "comments": []}


def test_issues_rules():
    got = issues.compute(_ds_for_issues(), {str(n): "x" for n in range(1, 20)})
    t = {(i["type"], i["key"]) for i in got}
    assert ("Data Missing", "body_size") in t
    assert ("Open Item", "rating") in t
    assert ("Data Error", "piping_class") in t                           # (1C0K8) in Line No. vs 1C3AS1
    assert ("Data Error", "elec_connection") in t                        # General Note 46 not in the document
    assert ("Extraction Review", "raw.13") in t
    ext = next(i for i in got if i["type"] == "Extraction Review")
    assert ext["no"] == len(got)                                         # internal checks numbered last
    assert all(i["sheet"] == 1 and i["of"] == 2 for i in got)


def test_contractor_exports(tmp_path):
    src = tmp_path / "ds.pdf"
    src.write_bytes(make_pdf())
    rows = [dict(i, ds_id=7, tag="00XV921", doc_no="00-09-IN-DTS-00003", rev="A", document_id=3, filename="ds.pdf",
                 valve_type="XV") for i in issues.compute(_ds_for_issues(), {str(n): "x" for n in range(1, 20)})
            if i["type"] != "Extraction Review"]
    x = exports.issues_xlsx(rows, "http://pc:8000", "CO2 PJT", "tester", ["ds.pdf"])
    ws = openpyxl.load_workbook(io.BytesIO(x))["Issue List"]
    head = [c.value for c in ws[1]]
    for h in ("Tag Number", "Doc. No.", "Rev.", "Sheet", "PDF Page", "Issue Type", "Description", "Required / Expected"):
        assert h in head
    assert ws.cell(2, head.index("Sheet") + 1).value == "1 of 2"
    assert "ds=7" in ws.cell(2, head.index("Open in Tool") + 1).hyperlink.target
    pdf = pymupdf.open(stream=exports.markup_pdf([{"path": str(src), "pages": [1, 2], "tag": "00XV921", "doc_no": "D",
                                                   "rev": "A", "filename": "ds.pdf", "issues": rows}], "CO2 PJT"), filetype="pdf")
    flags = [a.flags for p in pdf for a in p.annots()]
    assert flags and all(f & pymupdf.PDF_ANNOT_IS_PRINT for f in flags)   # mark-ups print with the page
    assert pdf.page_count == 3                                           # 2 datasheet pages + Issue Summary
