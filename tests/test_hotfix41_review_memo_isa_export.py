"""hotfix41 — 스스로 찾은 보완 셋 + 비교 창 보완.

① QFE 범례가 글자를 세 번 겹쳐 인쇄해 FIRST LETTER 의 뜻이 `PRESSURE PRESSURE PRESSURE` 가 되고
   Description 1124/1991행에 변수어가 세 번 들어갔다 → ISA 표 안에서 겹친 낱말을 하나로.
② 검토자가 올린 FreeText 메모(`ESDV OUTLET TAG 중복 확인.`)가 도면 글자와 같은 얼굴로 들어와
   기기 라벨이 됐다 → 사람 이름이 적힌 FreeText 안의 낱말은 뺀다 (SHX 주석은 그대로).
③ 개정 변경 내역 Excel (추가 · 수정 · 삭제 · 장 · 요약).
④ 비교 창에서 오른쪽 장을 사람이 고른다 (직전에만 있는 장) · Alt+←/→.
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pymupdf

from app import revision_export
from app.engine import isa_table, pidcache

ROOT = Path(__file__).resolve().parent.parent
JS = (ROOT / "app/static/app.js").read_text(encoding="utf-8")
HTML = (ROOT / "app/static/index.html").read_text(encoding="utf-8")
MAIN = (ROOT / "app/main.py").read_text(encoding="utf-8")


def _r(x0, y0, x1, y1):
    return pymupdf.Rect(x0, y0, x1, y1)


def test_overprinted_words_collapse_to_one_but_real_repeats_stay():
    words = [(_r(10, 10, 60, 18), "PRESSURE")] * 3 + [(_r(10, 30, 60, 38), "PRESSURE")]
    out = isa_table.dedupe_overprint(words)
    assert [t for _r_, t in out] == ["PRESSURE", "PRESSURE"]       # 같은 자리 셋 → 하나 · 다른 줄은 남는다
    # FIRST LETTER 열의 본문도 같은 함수를 지난다
    src = (ROOT / "app/engine/isa_table.py").read_text(encoding="utf-8")
    body = src[src.index("def derive("):src.index("def _letter_column")] if "def _letter_column" in src else src[src.index("def derive("):]
    assert "body = dedupe_overprint(" in body


def test_memo_text_is_dropped_but_drawing_text_under_the_memo_and_shx_notes_stay(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=400, height=300)
    page.insert_text((20, 40), "PIT 11LBB50CP001", fontsize=8)
    page.insert_text((210, 125), "FE", fontsize=8)          # 메모 상자 **밑**에 깔린 도면 글자
    a = page.add_freetext_annot(_r(200, 100, 360, 140), "ESDV OUTLET TAG 중복 확인.", fontsize=8)
    a.set_info(title="sungho1.choi")
    a.update()
    b = page.add_rect_annot(_r(20, 200, 120, 220))
    b.set_info(title=pidcache.SHX_AUTHOR, content="VALVES ACTUATORS")
    b.update()
    pdf = tmp_path / "memo.pdf"
    doc.save(pdf)
    raw = [w[4] for w in pymupdf.open(pdf)[0].get_text("words")]
    assert "ESDV" in raw and "FE" in raw                      # 메모 글자가 도면 글자와 같은 얼굴로 들어온다
    _doc, pages = pidcache.load_pages(pdf)
    pc = pages[0]
    texts = [t for _r_, t in pc.words]
    assert "ESDV" not in texts and not any("중복" in t for t in texts)
    assert "FE" in texts and "PIT" in texts                   # 상자 밑의 도면 글자는 산다
    assert any("VALVES" in t for t in texts)                  # SHX 글자는 들어온다
    assert pc.memo_words_dropped == 5
    # 숨김은 메모리에만 — 파일은 그대로다
    assert "ESDV" in [w[4] for w in pymupdf.open(pdf)[0].get_text("words")]


def test_memo_rule_is_by_author_not_by_annotation_kind():
    src = (ROOT / "app/engine/pidcache.py").read_text(encoding="utf-8")
    body = src[src.index("def _hide_memos"):src.index("def _shx_words")]
    assert 'a.type[1] != "FreeText"' in body and "SHX_AUTHOR" in body and "PDF_ANNOT_IS_HIDDEN" in body
    # 사각형 안 낱말을 빼는 방식은 없다 — 상자 밑의 도면 글자를 삼킨다 (AL NOUF1 1137 → 1102 의 원인)
    assert "_without_memo_words" not in src and "_memo_rects" not in src
    lp = src[src.index("def load_pages"):]
    assert lp.index("_hide_memos(page)") < lp.index("words = [(R.rect(w[:4]), w[4])")


def _row(key, page, state, **v):
    return {"key": key, "page_no": page, "tab": "FIELD", "values": dict(v),
            "rev": {"id": f"ID-{key}", "state": state, "basis": "TAG", "moved_pt": 1.5,
                    "changed": [{"field": "description", "was": "A", "now": "B"}] if state == "MODIFIED" else []}}


def test_change_report_workbook_has_the_five_sheets_and_copies_judgements():
    rows = [_row("a", 6, "ADDED", type="PIT", tag_no="11LBB50CP001", qty=1, scope="SCT", description="X"),
            _row("m", 6, "MODIFIED", type="TI", tag_no="11LBB50CT001", qty=2, scope="SCT", description="B"),
            _row("u", 7, "UNCHANGED", type="LI", tag_no="", qty=1, scope="SCT", description="")]
    deleted = [{"id": "00PAB10-032", "drawing_no": "D-1", "page_no": 11, "tab": "FIELD", "type": "RO",
                "tag_no": "50PAB11BP601", "description": "ORIFICE", "anchor": [207.8, 699.0],
                "nearest_distance": None, "radius": 54.4, "confirmed": False, "tag_elsewhere": []}]
    sheets = {"renumbered": [{"before": "D-3", "now": "D-203", "shared": 13, "before_tags": 20, "now_tags": 22}],
              "only_before": ["D-4"], "only_now": ["D-204"]}
    data = revision_export.build({"pdf_name": "x.pdf", "revision": "Rev.B", "compared_with": "Rev.A"},
                                 rows, {}, deleted, {6: "D-6", 7: "D-7"}, sheets, {"TAG": 2, "GEOMETRY": 0})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert [ws.title for ws in wb.worksheets] == ["요약", "추가", "수정", "삭제", "장"]
    assert wb["추가"].max_row == 2 and wb["추가"]["E2"].value == "PIT" and wb["추가"]["B2"].value == "D-6"
    mod = [c.value for c in wb["수정"][2]]
    assert "description: A → B" in mod and 1.5 in mod
    dele = [c.value for c in wb["삭제"][2]]
    assert "삭제 후보" in dele and "50PAB11BP601" in dele and any("어디에도 없습니다" in str(v) for v in dele)
    sheet = [[c.value for c in r] for r in wb["장"].iter_rows(min_row=2)]
    assert sheet[0][:4] == ["도면번호 바뀜", "D-3", "D-203", 13]
    norm = [[v or "" for v in r] for r in sheet]          # openpyxl 은 빈 칸을 None 으로 돌려준다
    assert ["직전에만 있음", "D-4", "", "", "", ""] in norm and ["이번에만 있음", "", "D-204", "", "", ""] in norm
    summary = {r[0].value: r[1].value for r in wb["요약"].iter_rows(max_col=2) if r[0].value}
    assert summary["추가"] == 1 and summary["수정"] == 1 and summary["변경 없음"] == 1 and summary["삭제 후보"] == 1


def test_export_endpoint_and_button_exist_and_judge_nothing():
    body = MAIN[MAIN.index('"/jobs/{job_id}/revision/changes.xlsx"'):MAIN.index('@app.post("/jobs/{job_id}/deleted/{stable_id}/confirm")')]
    assert "revision_export.build(" in body and "db.revision_states(" in body and "job_revision(job_id)" in body
    assert "match_radius" not in body and "compare(" not in body       # 저장된 판정을 옮겨 적을 뿐
    assert 'id="rev-export"' in HTML
    assert "/revision/changes.xlsx" in JS


def test_compare_pane_lets_the_person_pick_the_previous_sheet_and_step_with_keys():
    body = JS[JS.index("async function cmpShow"):JS.index("function cmpApplyZoom")]
    assert 'id="cmp-pick"' in body and "S.cmpPick[want]" in body and "only_before" in body
    # 자동(같은 도면번호 · 바뀐 번호)이 기본이고 고른 것은 그 왼쪽 장에만 기억된다
    assert "|| auto" in body
    kd = JS[JS.index('document.addEventListener("keydown"'):]
    kd = kd[:kd.index("\n});\n") + 5]
    assert "ArrowLeft" in kd and "ArrowRight" in kd and "focusChange(" in kd and "ev.altKey" in kd
