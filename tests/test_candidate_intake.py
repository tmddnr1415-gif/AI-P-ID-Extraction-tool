"""후보 병렬 워크북 읽어들이기 (10회차) — 빠른 시험.  실제 워크북을 쓴다."""
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "spike"))
from app import axis_overrides as ov          # noqa: E402
import read_candidates as rc                  # noqa: E402

BOOK = ROOT / "out" / "candidates_nouf1.xlsx"


def _sheet(wb):
    return wb[wb.sheetnames[1]] if wb.sheetnames[0].startswith("0.") else wb.active


@pytest.mark.skipif(not BOOK.exists(), reason="후보 워크북이 없는 기계")
def test_empty_workbook_applies_nothing(tmp_path):
    """빈 워크북(예시 3행만 채워진 상태)은 0행을 반영한다."""
    got = rc.apply(BOOK, tmp_path, "AXIS10", {}, origin_job="j1")
    assert got == {"selected": 0, "written": 0, "no_stable_id": 0}
    assert not (tmp_path / "projects").exists() or \
        ov.load(ov.path_for(tmp_path, "AXIS10")) == {}


@pytest.mark.skipif(not BOOK.exists(), reason="후보 워크북이 없는 기계")
def test_example_rows_are_skipped(tmp_path):
    """'작성 예시' 행은 선택이 채워져 있어도 건너뛴다."""
    ws = _sheet(openpyxl.load_workbook(BOOK))
    marked = [i for i in range(2, ws.max_row + 1)
              if rc.EXAMPLE_NOTE in str(ws.cell(i, rc.COL["memo"]).value or "")]
    assert len(marked) == 3, "예시는 3행이어야 한다"
    for i in marked:                       # 예시 행에는 선택이 실제로 들어 있다
        assert str(ws.cell(i, rc.COL["pick"]).value or "").strip()
    assert rc.selections(ws) == []         # 그런데 읽기에는 잡히지 않는다


@pytest.mark.skipif(not BOOK.exists(), reason="후보 워크북이 없는 기계")
def test_a_filled_row_reaches_the_ledger(tmp_path):
    """사람이 채운 한 행이 장부에 `후보선택(워크북)` 출처로 적힌다."""
    wb = openpyxl.load_workbook(BOOK)
    ws = _sheet(wb)
    row = next(i for i in range(2, ws.max_row + 1)
               if not str(ws.cell(i, rc.COL["memo"]).value or "")
               and ws.cell(i, rc.COL["c1"] + 3).value)      # 후보4 가 있는 행
    key = str(ws.cell(i := row, rc.COL["key"]).value)
    want = ws.cell(row, rc.COL["c1"] + 3).value
    ws.cell(row, rc.COL["pick"]).value = "4"
    ws.cell(row, rc.COL["scope"]).value = "전 페이지 규칙 후보"
    ws.cell(row, rc.COL["reason"]).value = "기기가 없고 목적지가 정체성이다"
    edited = tmp_path / "filled.xlsx"
    wb.save(edited)

    got = rc.apply(edited, tmp_path, "AXIS10", {key: "10LBA10-999"},
                   origin_job="j1")
    assert got["selected"] == 1 and got["written"] == 1
    entry = ov.load(ov.path_for(tmp_path, "AXIS10"))["10LBA10-999"]
    assert entry["sentence"] == want
    assert entry["source_from"] == ov.SOURCE_WORKBOOK
    assert entry["candidate"] == "4"
    assert entry["applies_to"] == "전 페이지 규칙 후보"
    assert entry["reason"]
    # 안정 ID 가 안 물리면 적지 않고 센다 — 지어내지 않는다
    again = rc.apply(edited, tmp_path, "AXIS10", {}, origin_job="j1")
    assert again == {"selected": 1, "written": 0, "no_stable_id": 1}
