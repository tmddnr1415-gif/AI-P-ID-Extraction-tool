"""What a shipped build has to be able to do, checked without shipping one.

Three separable things are pinned here:

  * an error report carries the evidence *automatically*.  The whole design
    depends on a person typing two fields and no more, so a change that quietly
    starts asking them for a coordinate should fail a test rather than a user.
  * the diagnostic export names what it leaves out.  "The drawings do not leave
    the PC" is a promise, and a promise that is only in a comment is not kept.
  * the packaging separates what is read from what is written, and does not
    bundle the client's documents.

None of this needs an analysis: the fixtures build a job by hand.
"""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app import db, paths, version    # noqa: E402


@pytest.fixture()
def con(tmp_path):
    return db.connect(tmp_path / "app.db")


@pytest.fixture()
def job(con, tmp_path):
    """One done job with one row, one edit and one review state."""
    pdf = tmp_path / "sheet.pdf"
    pdf.write_bytes(b"%PDF-1.4\n")
    db.create_job(con, "j1", "sheet.pdf", "abc123", pdf)
    db.store_result(con, "j1", {
        "fingerprint": "ff00ff00",
        "pages": [{"page_no": 6, "drawing_no": "D00P-10LBA10-M05-0001",
                   "title": "P&ID FOR HP STEAM", "page_kind": "PID",
                   "in_scope": True, "scope_reason": "", "width": 2384.0,
                   "height": 1684.0}],
        "layers": {"6": {}},
        "rows": [{
            "key": "row1", "tab": "FIELD", "page_no": 6,
            "drawing_no": "D00P-10LBA10-M05-0001", "origin": "DRAWING",
            "rect": [10.0, 20.0, 30.0, 40.0],
            "type": "PIT", "qty": 2, "system": "HP Steam", "valve_type": "",
            "vendor_supply": "", "scope": "", "tag_no": "",
            "description": "UNIT #10 HP STEAM PRESSURE",
            "description_grade": "CONFIRMED", "remark": "",
            "needs_review": "", "annotation": "",
            "evidence": {"anchor": "PT", "rules_hit": ["VENDOR_MARK_GLYPH"],
                         "qty_basis": "sheet multiplier 2"},
        }],
        "applied_rules": {"instrument_ruleset": "v3-anchor-audit",
                          "layout": {"items": [
                              {"key": "sheet.width_pt", "value": 2384.0,
                               "source": "DERIVED", "evidence": "page rectangle"},
                              {"key": "broken_line.brk_max_mark", "value": 0,
                               "source": "UNAVAILABLE", "evidence": "not isolated"}],
                              "moved": [["sheet.width_pt", 1.0, 2384.0]]}},
        "multipliers": {}, "legend": {}, "glyphs": {}, "job_review": [],
        "timings": {}, "pipe_trace": {}, "description_scope": {},
        "description_build": {}, "description_grades": {},
    })
    db.set_user_value(con, "j1", "row1", "description", "UNIT #10 HP STEAM PRESSURE A")
    db.set_review_state(con, "j1", "row1", "DESC_GRADE_PARTIAL", "EDITED", "")
    return "j1"


# --------------------------------------------------------------------------
# reports
# --------------------------------------------------------------------------

def test_a_report_accepts_two_answers_and_refuses_a_third_kind(con, job):
    rid = db.add_report(con, job, "ROW", "DESCRIPTION", "발주처는 SUCTION 이라고 씁니다",
                        row_key="row1", page_no=6)
    assert db.report_count(con, job) == 1
    with pytest.raises(ValueError):
        db.add_report(con, job, "ROW", "COLOUR", "…")
    with pytest.raises(ValueError):
        db.add_report(con, job, "GUESS", "TYPE", "…")
    # editing changes what a person wrote and nothing else
    before = db.get_report(con, rid)
    after = db.update_report(con, rid, detail="고쳤습니다")
    assert after["detail"] == "고쳤습니다"
    assert after["capture"] == before["capture"]
    assert after["created_at"] == before["created_at"]
    db.delete_report(con, rid)
    assert db.report_count(con, job) == 0
    with pytest.raises(KeyError):
        db.delete_report(con, rid)


def test_the_report_captures_the_evidence_the_reporter_never_types(con, job, monkeypatch):
    """The two typed fields are the *only* two.  Everything else is taken."""
    import app.main as m
    monkeypatch.setattr(m, "CON", con)
    cap = m._capture_row(job, "row1")
    ident = cap["identity"]
    assert ident["drawing_no"] == "D00P-10LBA10-M05-0001"
    assert (ident["page_no"], ident["type"], ident["rect"]) == (6, "PIT", [10.0, 20.0, 30.0, 40.0])
    assert cap["applied_rules"] == ["VENDOR_MARK_GLYPH"]
    assert cap["evidence"]["qty_basis"] == "sheet multiplier 2"
    # both sides of every field the engine and the reviewer disagree on
    assert cap["ai_values"]["description"] == "UNIT #10 HP STEAM PRESSURE"
    assert cap["user_values"]["description"] == "UNIT #10 HP STEAM PRESSURE A"
    assert cap["review_state"]["DESC_GRADE_PARTIAL"]["state"] == "EDITED"
    ctx = m._capture_context(job)
    assert ctx["fingerprint"] == "ff00ff00"
    assert ctx["pdf_sha256"] == "abc123"
    assert ctx["build"]["version"] == version.VERSION
    assert ctx["layout"]["moved"] == [["sheet.width_pt", 1.0, 2384.0]]


# --------------------------------------------------------------------------
# diagnostic export
# --------------------------------------------------------------------------

def test_the_export_carries_the_decisions_and_not_the_documents(
        con, job, tmp_path, monkeypatch):
    import app.main as m
    monkeypatch.setattr(m, "CON", con)
    monkeypatch.setattr(m, "DIAGNOSTICS", tmp_path / "diagnostics")
    db.add_report(con, job, "ROW", "QTY", "승수가 4 여야 합니다", row_key="row1", page_no=6)
    out = m._diagnostic_zip(job)

    info = version.info()
    assert out["filename"].startswith(f"pid_diag_v{info['version']}_")
    assert job in out["filename"]
    assert out["bytes"] > 0

    z = zipfile.ZipFile(tmp_path / "diagnostics" / out["filename"])
    names = set(z.namelist())
    assert {"MANIFEST.json", "reports.json", "edits.json", "feedback.json",
            "analysis/rows.json", "analysis/derived_layout.json",
            "analysis/applied_rules.json", "logs/server.log"} <= names
    # nothing that is the client's, by extension or by content
    assert not [n for n in names if n.lower().endswith((".pdf", ".xlsx"))]

    man = json.loads(z.read("MANIFEST.json"))
    assert man["job"]["fingerprint"] == "ff00ff00"
    assert man["job"]["pdf_sha256"] == "abc123"      # the hash, not the file
    assert man["build"]["version"] == info["version"]
    assert man["build"]["built_at"]
    assert man["counts"]["reports"] == 1
    assert man["counts"]["edits"] == 1
    assert man["excluded"] == m.DIAGNOSTIC_EXCLUDED and man["excluded"]

    # the derived layout is in, because it is the only record of what *that* PC's
    # PDF stated - the whole point of collecting a report from another site
    lay = json.loads(z.read("analysis/derived_layout.json"))
    assert [i["key"] for i in lay["items"]] == ["sheet.width_pt",
                                                "broken_line.brk_max_mark"]
    assert man["counts"]["derived_layout_measured"] == 1
    assert man["counts"]["derived_layout_applied"] == 1
    assert json.loads(z.read("reports.json"))["reports"][0]["what"] == "QTY"


# --------------------------------------------------------------------------
# packaging
# --------------------------------------------------------------------------

def test_what_is_read_and_what_is_written_are_different_roots():
    """A one-file exe unpacks its bundle to a temp dir that is deleted on exit,
    so anything written there is lost.  The two roots must not be the same."""
    assert paths.resource_root() != paths.data_dir()
    assert paths.resource("config", "project_alnouf1.yaml").exists()
    assert not paths.frozen()          # the test suite runs from the checkout


def test_the_recipe_ships_the_code_and_none_of_the_clients_documents():
    spec = (ROOT / "pid_extract.spec").read_text(encoding="utf-8")
    assert '("app/static", "app/static")' in spec
    assert '("config", "config")' in spec
    # the drawings and the workbooks are the client's; `data/` is gitignored for
    # the same reason and must never be a datas entry
    assert '"data"' not in spec.split("datas += [")[1].split("]")[0]
    # the engine's bare-name modules have to be in the bundle under those names
    assert "ENGINE_MODULES" in spec and 'pathex=[".", "app/engine"]' in spec
    build = (ROOT / "build.bat").read_text(encoding="utf-8")
    assert "_build.json" in build and "pytest" in build


def test_the_build_says_which_build_it_is():
    i = version.info()
    assert i["version"] == version.VERSION
    assert i["kind"] in ("exe", "source")
    assert i["built_at"]
    assert version.label().startswith(f"v{i['version']}")


# --------------------------------------------------------------------------
# the deliverable carries our rows and nothing of the template's
# --------------------------------------------------------------------------

def _form(rows=3, cols=12, data_first=8):
    """A miniature of the client's form: title rows, header, data, note block."""
    import openpyxl
    from openpyxl.styles import Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2.0_Instrument List"
    ws["A1"] = "INSTRUMENT LIST"
    ws["A2"] = "Project : TEST"
    ws["A6"] = "NO"
    ws.cell(6, 6).value = "SYSTEM"
    ws.cell(6, 7).value = "P&ID No."
    ws.cell(6, 8).value = "TYPE"
    ws.cell(6, 9).value = "Q'ty"
    ws.cell(6, 10).value = "DESCRIPTION"
    ws.cell(6, 11).value = "Design table No."
    ws.cell(6, 12).value = "Operating condition"
    thin = Side(style="thin")
    for i in range(rows):
        r = data_first + i
        ws.cell(r, 1).value = i + 1
        ws.cell(r, 7).value = f"OLD-DRAWING-{i}"
        ws.cell(r, 8).value = "PIT"
        ws.cell(r, 11).value = f"A{i}"          # unmapped: the design data
        ws.cell(r, 12).value = 187.3 + i        # unmapped: the design data
        for c in range(1, cols + 1):
            ws.cell(r, c).border = Border(left=thin, right=thin, top=thin,
                                          bottom=thin)
        ws.cell(r, 11).fill = PatternFill("solid", fgColor="FF92D050")
        ws.cell(r, 7).font = Font(strike=True)
    ws.cell(data_first + rows + 1, 1).value = "Note)"
    ws.cell(data_first + rows + 2, 1).value = "1. REFER TO THE SPECIFICATION."
    ws.merged_cells.ranges.add("A1:E1")
    ws.column_dimensions["G"].width = 24.5
    return wb


def test_a_deliverable_carries_no_value_from_the_template(tmp_path):
    """The defect this pins: 14,121 cells of the client's own design data went
    out under our rows because only the mapped columns were cleared.

    Row 8 of the client's list was an HP steam PIT; row 8 of ours is an ACW pump
    PDIT.  Everything the config does not map - the design table, the operating
    condition, the pipe material - stayed at row 8 and described the wrong
    instrument.
    """
    import openpyxl
    from app import excel_out

    src = tmp_path / "form.xlsx"
    _form().save(src)

    class Cfg:
        def get(self, key):
            assert key == "excel"
            return {"sheet": "2.0_Instrument List", "first_data_row": 8,
                    "columns": {"no": 1, "system": 6, "pid_no": 7, "type": 8,
                                "qty": 9, "description": 10}}

    rows = [{"key": "k1", "tab": "FIELD", "page_no": 1,
             "drawing_no": "NEW-DRAWING", "values": {"type": "PDIT", "qty": 2,
                                                     "system": "ACW",
                                                     "description": "ACW PUMP"},
             "user": {}, "evidence": {}, "review_codes": []}]
    out = tmp_path / "out.xlsx"
    excel_out.write_deliverable(src, out, rows, Cfg(), "FIELD")

    ws = openpyxl.load_workbook(out)["2.0_Instrument List"]
    assert ws.cell(8, 7).value == "NEW-DRAWING"
    assert ws.cell(8, 8).value == "PDIT"
    # the unmapped columns carry nothing at all
    assert ws.cell(8, 11).value is None, "the template's Design table No. survived"
    assert ws.cell(8, 12).value is None, "the template's operating condition survived"
    # and the rows the new data did not reach are empty across every column
    for r in (9, 10):
        for c in range(1, 13):
            assert ws.cell(r, c).value is None, f"row {r} col {c} not cleared"
    # the form itself is untouched
    assert ws["A1"].value == "INSTRUMENT LIST"
    assert ws.cell(6, 11).value == "Design table No."
    assert ws.cell(12, 1).value == "Note)"
    assert round(ws.column_dimensions["G"].width, 1) == 24.5
    # and the data row kept its borders
    assert ws.cell(8, 1).border.left.style == "thin"


def test_the_blank_form_keeps_the_form_and_drops_the_working_marks(tmp_path):
    """A blank form is the client's form with the data region emptied.

    The shading question was settled by measurement rather than by the brief:
    AL NOUF1's three workbooks carry no grey attribute shading at all.  What is
    there is green FF92D050 and yellow FFFF00 on scattered cells and whole rows,
    not aligned with the strikethrough rows - the client's own working marks.
    They go, because a Rev.A deliverable has to open with no shading on it.
    """
    import openpyxl
    from app import excel_out

    src = tmp_path / "form.xlsx"
    _form().save(src)
    out = tmp_path / "blank.xlsx"
    info = excel_out.blank_form(src, out, "2.0_Instrument List")

    assert info["data_rows"] == 3 and info["values_cleared"] > 0
    assert info["fills_cleared"] == 3 and info["strikes_cleared"] == 3

    ws = openpyxl.load_workbook(out)["2.0_Instrument List"]
    for r in (8, 9, 10):
        for c in range(1, 13):
            cell = ws.cell(r, c)
            assert cell.value is None
            assert cell.fill.patternType is None, "shading survived into Rev.A"
            assert not cell.font.strike, "strikethrough survived into Rev.A"
            assert cell.border.left.style == "thin", "the border was taken too"
    # rows 1-7 and the note block are the form and stay
    assert ws["A1"].value == "INSTRUMENT LIST"
    assert ws.cell(6, 12).value == "Operating condition"
    assert ws.cell(12, 1).value == "Note)"
    assert ws.cell(13, 1).value == "1. REFER TO THE SPECIFICATION."
    assert round(ws.column_dimensions["G"].width, 1) == 24.5
