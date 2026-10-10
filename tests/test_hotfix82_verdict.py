"""hotfix82 — O/X 평가: 행이 식별되어야 하는가(O) · 아닌가(X).  결과는 VOC 함에 쌓여 학습 자료가 된다."""
from __future__ import annotations

import json
import os
import pathlib
import sys
import time

import pytest

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


def _app(tmp_path):
    os.environ["PID_DATA_DIR"] = str(tmp_path)
    os.environ.pop("PID_VOC_DIR", None)
    for mod in [m for m in list(sys.modules) if m.startswith("app")]:
        del sys.modules[mod]
    from fastapi.testclient import TestClient
    from app import main
    return TestClient(main.app), main


def _job(main, job="jx", n=3):
    from app import db
    import pymupdf
    main.UPLOADS.mkdir(parents=True, exist_ok=True)
    pdf = main.UPLOADS / f"{job}_x.pdf"
    doc = pymupdf.open()
    pg = doc.new_page(width=300, height=200)
    pg.insert_text((40, 60), "PIT")
    doc.save(str(pdf)); doc.close()
    main.CON.execute(
        "INSERT INTO job (id,pdf_name,pdf_sha256,pdf_path,created_at,status,"
        "progress,message,fingerprint,project,revision)"
        " VALUES (?,?,?,?,?,'done',1.0,'','fp','QFE','B')",
        (job, "x.pdf", job, str(pdf), time.time()))
    keys = []
    for i in range(n):
        keys.append(db.add_row(main.CON, job, page_no=1, tab="FIELD", drawing_no="D-1",
                               values={"type": "PIT", "qty": 1, "tag_no": f"T{i}"},
                               rect=[35 + 40 * i, 45, 70 + 40 * i, 65]))
    # 검출 행처럼 보이게 — add_row 는 added=1 이라 X 가 지우지 않는다
    main.CON.execute("UPDATE item SET added=0 WHERE job_id=?", (job,))
    main.CON.commit()
    return job, keys


def _rows(c, job):
    return {r["key"]: r for r in c.get(f"/jobs/{job}/rows?tab=ALL").json()}


def test_o_records_and_x_removes_then_o_restores(tmp_path):
    c, main = _app(tmp_path)
    job, (a, b, k3) = _job(main)
    r = c.post(f"/jobs/{job}/verdicts",
               json={"items": [{"key": a, "verdict": "O"}, {"key": b, "verdict": "X"}], "author": "홍길동"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["counts"] == {"O": 1, "X": 1, "open": 1, "rows": 3, "qty_O": 0, "qty_X": 0}
    rows = _rows(c, job)
    assert rows[a]["verdict"]["verdict"] == "O" and rows[a]["verdict"]["author"] == "홍길동"
    assert rows[b]["verdict"]["verdict"] == "X" and rows[b]["removed"]       # X 는 출력에서 뺀다
    assert not rows[a]["removed"] and rows[k3]["verdict"] == {}
    # X → O : X 가 지운 행은 되살아난다
    r = c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": b, "verdict": "O"}], "author": "홍길동"})
    assert r.status_code == 200 and not _rows(c, job)[b]["removed"]
    # 비우면 평가가 사라진다 (행은 그대로)
    r = c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": a, "verdict": ""}], "author": "홍길동"})
    assert r.json()["counts"] == {"O": 1, "X": 0, "open": 2, "rows": 3, "qty_O": 0, "qty_X": 0}
    assert _rows(c, job)[a]["verdict"] == {}


def test_x_does_not_touch_rows_a_person_removed_for_another_reason(tmp_path):
    c, main = _app(tmp_path)
    job, (a, _, _) = _job(main)
    assert c.request("DELETE", f"/jobs/{job}/rows/{a}",
                     json={"klass": "FALSE_POSITIVE", "reason": "배관 주석", "author": "김", "exclude": True}
                     ).status_code == 200
    c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": a, "verdict": "O"}], "author": "홍"})
    assert _rows(c, job)[a]["removed"]            # 사람이 따로 지운 행은 O 가 되살리지 않는다


def test_one_request_is_one_voc_with_every_row(tmp_path):
    c, main = _app(tmp_path)
    job, keys = _job(main)
    items = [{"key": k, "verdict": "O" if i else "X"} for i, k in enumerate(keys)]
    r = c.post(f"/jobs/{job}/verdicts", json={"items": items, "author": "홍길동", "note": "p1 전체"})
    vid = r.json()["voc"]["id"]
    inbox = tmp_path / "voc" / "inbox"
    assert [p.name for p in inbox.iterdir()] == [vid]                     # 행마다가 아니라 요청마다 한 건
    rec = json.loads((inbox / vid / "voc.json").read_text(encoding="utf-8"))
    assert rec["source"] == "VERDICT" and rec["category"] == "VERDICT"
    assert rec["counts"] == {"O": 2, "X": 1, "qty_O": 0, "qty_X": 0} and "p1 전체" in rec["reason"]
    assert sorted(v["key"] for v in rec["verdicts"]) == sorted(keys)
    assert len(rec["rows"]) == 3 and {x["verdict"] for x in rec["rows"]} == {"O", "X"}
    assert rec["rows"][0]["identity"]["drawing_no"] == "D-1"
    assert rec["has_crop"] is False                                       # 여럿이면 조각 없음
    # 한 행 X 는 조각을 남긴다
    r = c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": keys[1], "verdict": "X"}], "author": "홍"})
    v2 = r.json()["voc"]["id"]
    assert (inbox / v2 / "crop.png").is_file()
    # 비우기만 한 요청은 VOC 를 만들지 않는다
    r = c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": keys[1], "verdict": ""}], "author": "홍"})
    assert r.json()["voc"] is None and len(list(inbox.iterdir())) == 2
    # 화면 목록은 접수됨으로 보이고 개발 쪽 목록(list)에서는 기본으로 숨긴다 — spike/voc.py 가 가른다
    assert c.get(f"/jobs/{job}/verdicts").json()["counts"]["X"] == 1


def test_bad_payloads_are_400_and_unknown_keys_are_reported(tmp_path):
    c, main = _app(tmp_path)
    job, (a, _, _) = _job(main)
    for bad in ({}, {"items": []}, {"items": "x"}, {"items": [{"verdict": "O"}]},
                {"items": [{"key": a, "verdict": "Y"}]}):
        assert c.post(f"/jobs/{job}/verdicts", json=bad).status_code == 400, bad
    r = c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": a, "verdict": "O"}, {"key": "nope", "verdict": "X"}]})
    assert r.status_code == 200 and r.json()["missing"] == ["nope"] and r.json()["rows"][0]["key"] == a
    assert c.post("/jobs/none/verdicts", json={"items": [{"key": a, "verdict": "O"}]}).status_code == 404


def test_voc_amend_keeps_the_head_and_adds_rows(tmp_path):
    from app import voc
    rec = voc.write({"category": "VERDICT", "source": "VERDICT", "reason": "O 1"}, root=tmp_path)
    out = voc.amend(rec["id"], {"verdicts": [{"key": "k", "verdict": "O"}], "id": "hack",
                                "category": "UI"}, root=tmp_path)
    back = json.loads((tmp_path / "inbox" / rec["id"] / "voc.json").read_text(encoding="utf-8"))
    assert back["id"] == rec["id"] and back["category"] == "VERDICT" and back["verdicts"][0]["key"] == "k"
    assert out == back
    with pytest.raises(FileNotFoundError):
        voc.amend("VOC-20260101-000000-abcdef", {}, root=tmp_path)


def test_qty_verdict_and_note_are_separate_cells_and_x_only_on_identification(tmp_path):
    c, main = _app(tmp_path)
    job, (a, b, _) = _job(main)
    r = c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": a, "qty_verdict": "X", "note": "x2 여야 함"}], "author": "홍"})
    assert r.status_code == 200, r.text
    row = _rows(c, job)[a]
    assert row["verdict"]["qty_verdict"] == "X" and row["verdict"]["verdict"] == "" and row["verdict"]["note"] == "x2 여야 함"
    assert not row["removed"]                                  # 수량 X 는 행을 빼지 않는다
    assert r.json()["counts"]["qty_X"] == 1 and r.json()["counts"]["open"] == 3
    # 식별 O 를 적어도 수량·비고는 그대로
    c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": a, "verdict": "O"}], "author": "홍"})
    row = _rows(c, job)[a]
    assert (row["verdict"]["verdict"], row["verdict"]["qty_verdict"], row["verdict"]["note"]) == ("O", "X", "x2 여야 함")
    # 비고만 고치기
    c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": a, "note": "확인 완료"}], "author": "김"})
    row = _rows(c, job)[a]
    assert row["verdict"]["note"] == "확인 완료" and row["verdict"]["author"] == "김" and row["verdict"]["verdict"] == "O"
    # 셋 다 비우면 미평가
    c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": a, "verdict": "", "qty_verdict": "", "note": ""}], "author": "홍"})
    assert _rows(c, job)[a]["verdict"] == {}
    # 칸 이름이 하나도 없는 항목은 400
    assert c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": b}]}).status_code == 400
    # VOC 기록에 수량 · 비고가 함께
    r = c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": b, "verdict": "O", "qty_verdict": "O", "note": "ok"}], "author": "홍"})
    inbox = tmp_path / "voc" / "inbox"
    rec = json.loads((inbox / r.json()["voc"]["id"] / "voc.json").read_text(encoding="utf-8"))
    assert rec["verdicts"][0] == {"key": b, "verdict": "O", "qty_verdict": "O", "note": "ok"}
    assert rec["counts"] == {"O": 1, "X": 0, "qty_O": 1, "qty_X": 0} and "수량 O" in rec["reason"]


def test_verdict_excel_lists_every_live_row_with_both_verdicts(tmp_path):
    import io
    import openpyxl
    c, main = _app(tmp_path)
    job, (a, b, k3) = _job(main)
    c.post(f"/jobs/{job}/verdicts", json={"items": [{"key": a, "verdict": "O", "qty_verdict": "X", "note": "x2"},
                                                     {"key": b, "verdict": "X"}], "author": "홍길동"})
    # 누락 — 사람이 더한 행
    from app import db
    added = db.add_row(main.CON, job, page_no=1, tab="FIELD", drawing_no="D-1", values={"type": "LIT", "qty": 1}, rect=[1, 1, 9, 9])
    main.CON.commit()
    r = c.get(f"/jobs/{job}/verdicts.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["요약", "식별 VOC", "누락 추가", "출력 제외"]
    ws = wb["식별 VOC"]
    head = [c.value for c in ws[1]]
    assert head[:6] == ["No", "P&ID No.", "쪽", "탭", "Type", "Tag No."] and "식별 O/X" in head and "수량 O/X" in head
    rows = {r[head.index("Tag No.")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert len(rows) == 4                                      # 검출 3 + 사람 추가 1 (지운 행 0)
    ra = rows["T0"]
    assert (ra[head.index("식별 O/X")], ra[head.index("수량 O/X")], ra[head.index("비고")], ra[head.index("평가자")]) == ("O", "X", "x2", "홍길동")
    rb = rows["T1"]
    assert rb[head.index("식별 O/X")] == "X" and rb[head.index("출력")].startswith("제외")
    assert wb["누락 추가"].max_row == 2 and wb["누락 추가"].cell(2, head.index("출처") + 1).value.startswith("사람 추가")
    assert wb["출력 제외"].max_row == 2
    assert c.get("/jobs/none/verdicts.xlsx").status_code == 404
