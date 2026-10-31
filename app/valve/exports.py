"""Deliverables for the design contractor: Issue List (xlsx) and the marked-up datasheet PDF.

Issue List  one row per issue: Tag · Doc. No. · Rev · Sheet x of y · file · PDF page · row · field · value ·
            issue type · description · required/expected · severity · status · reviewer note, plus two links —
            'Open in tool' (deep link to the page and point) and 'PDF page' (the original at that page).
Mark-up PDF the datasheet pages that have issues, copied from the originals (originals untouched), each issue
            boxed in its colour with a printed number label; after each datasheet an 'Issue Summary' page that
            lists the numbers; a cover page with counts per tag; a bookmark per tag.  Annotations carry the
            print flag, so 'print with mark-ups' shows them on paper.

Colours: red = data error / PMS non-compliance · orange = data missing / open item · yellow = revision change ·
blue = reviewer correction / comment · purple = extraction check (internal).
"""
from __future__ import annotations

import datetime as dt
import html
import io
from urllib.parse import quote

import openpyxl
import pymupdf
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RGB = {"error": (0.86, 0.15, 0.15), "missing": (1.0, 0.55, 0.0), "revision": (0.95, 0.76, 0.0),
       "comment": (0.18, 0.40, 0.90), "extract": (0.48, 0.30, 0.86)}
HEXFILL = {"error": "FBE3E1", "missing": "FFE8CC", "revision": "FFF4C2", "comment": "E3ECFD", "extract": "EFE6FB"}


def deep_link(base: str, project: str, i: dict) -> str:
    return (f"{base.rstrip('/')}/valve/?project={quote(project)}&ds={i['ds_id']}&key={quote(i['key'])}"
            f"&page={i.get('page') or ''}#review")


def pdf_link(base: str, i: dict) -> str:
    return f"{base.rstrip('/')}/api/valve/documents/{i['document_id']}/file#page={i.get('page') or 1}"


# ── Issue List (xlsx) ──────────────────────────────────────────────────────────

COLS = [("No.", 6), ("Tag Number", 16), ("Valve Type", 8), ("Doc. No.", 20), ("Rev.", 6), ("Sheet", 9),
        ("PDF File", 22), ("PDF Page", 8), ("Row", 6), ("Field", 26), ("Datasheet Value", 30), ("Issue Type", 18),
        ("Description", 60), ("Required / Expected", 34), ("Severity", 9), ("Status", 11), ("Reviewer Note", 30),
        ("Open in Tool", 13), ("PDF Page Link", 13)]


def issues_xlsx(rows: list[dict], base: str, project: str, by: str = "", sources: list[str] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Issue List"
    thin = Side(style="thin", color="C9CFDB")
    head_fill, head_font = PatternFill("solid", fgColor="1F2D4D"), Font(bold=True, color="FFFFFF")
    ws.append([c for c, _ in COLS])
    for n, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(n)].width = w
        c = ws.cell(1, n)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for n, i in enumerate(rows, 1):
        sheet = f"{i['sheet']} of {i['of']}" if i.get("sheet") else ""
        val = (i.get("value") or "") + (f" {i['unit']}" if i.get("unit") and i.get("value") else "")
        ws.append([n, i["tag"], i.get("valve_type", ""), i["doc_no"], i["rev"], sheet, i.get("filename", ""), i.get("page"),
                   i.get("row"), i["field"], val or "(blank)", i["type"], i["description"], i.get("expected", ""),
                   i["severity"], i["status"], i.get("note", ""), "Open", "PDF p." + str(i.get("page") or "")])
        r = ws.max_row
        fill = PatternFill("solid", fgColor=HEXFILL.get(i["color"], "FFFFFF"))
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(r, c)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 12).fill = fill
        if base:
            for col, url in ((18, deep_link(base, project, i)), (19, pdf_link(base, i))):
                ws.cell(r, col).hyperlink = url
                ws.cell(r, col).font = Font(color="2F5BD8", underline="single")
            ws.cell(r, 2).hyperlink = deep_link(base, project, i)       # the tag itself also jumps to the point
            ws.cell(r, 2).font = Font(color="2F5BD8", underline="single", bold=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(1, ws.max_row)}"

    sm = wb.create_sheet("Summary by Tag")
    types = [t for t in ["Data Error", "Data Missing", "Open Item", "PMS Non-Compliance", "PMS Review", "Revision Change",
                         "Reviewer Correction", "Reviewer Comment", "Extraction Review"] if any(i["type"] == t for i in rows)]
    sm.append(["Tag Number", "Doc. No.", "Rev.", "Total"] + types)
    for c in range(1, len(types) + 5):
        sm.cell(1, c).fill, sm.cell(1, c).font = head_fill, head_font
        sm.column_dimensions[get_column_letter(c)].width = 14 if c > 3 else 20
    tags = {}
    for i in rows:
        t = tags.setdefault(i["tag"], {"doc": i["doc_no"], "rev": i["rev"], "n": {}})
        t["n"][i["type"]] = t["n"].get(i["type"], 0) + 1
    for tag, t in tags.items():
        sm.append([tag, t["doc"], t["rev"], sum(t["n"].values())] + [t["n"].get(x, 0) or None for x in types])
    sm.append(["Total", "", "", len(rows)] + [sum(1 for i in rows if i["type"] == x) for x in types])
    sm.freeze_panes = "B2"
    sm.auto_filter.ref = f"A1:{get_column_letter(len(types) + 4)}{max(1, sm.max_row - 1)}"

    info = wb.create_sheet("Info")
    for row in [["Valve Datasheet Review — Issue List"], ["Project", project],
                ["Generated", dt.datetime.now().strftime("%Y-%m-%d %H:%M")], ["Generated by", by or "-"],
                ["Source files", ", ".join(sources or [])], [],
                ["Issue Type", "Meaning"],
                ["Data Error", "The datasheet contradicts itself or refers to a note that does not exist"],
                ["Data Missing", "A required field is on the form but has no value"],
                ["Open Item", "A required field is still HOLD / TBC"],
                ["PMS Non-Compliance", "Differs from the Piping Material Specification (material compared by grade, not spelling)"],
                ["PMS Review", "Could not be confirmed against the PMS — engineer to confirm"],
                ["Revision Change", "Changed from the previous revision of the same tag"],
                ["Reviewer Correction / Comment", "Entered by the reviewer"],
                [], ["Every row is linked: 'Open' opens the review tool at the field, 'PDF p.' opens the original page."]]:
        info.append(row)
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 100
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ── marked-up PDF ──────────────────────────────────────────────────────────────

def _ascii(s: str) -> str:
    return (s or "").replace("→", "->").replace("≤", "<=").replace("—", "-").replace("·", "-")


def _story_pages(html_body: str) -> pymupdf.Document:
    css = ("body{font-family:sans-serif;font-size:8pt;} h1{font-size:12pt;margin:0 0 4pt 0;} p{margin:0 0 6pt 0;}"
           "table{border-collapse:collapse;width:100%;} th{background-color:#1F2D4D;color:#FFFFFF;text-align:left;}"
           "th,td{border:0.5pt solid #9AA3B2;padding:2pt 3pt;vertical-align:top;}")
    buf = io.BytesIO()
    w = pymupdf.DocumentWriter(buf)
    story = pymupdf.Story(html=html_body, user_css=css)
    box = pymupdf.paper_rect("a4-l")
    more = 1
    while more:
        dev = w.begin_page(box)
        more, _ = story.place(box + (30, 30, -30, -30))
        story.draw(dev)
        w.end_page()
    w.close()
    return pymupdf.open("pdf", buf.getvalue())


def _swatch(color):
    r, g, b = RGB.get(color, (0.5, 0.5, 0.5))
    return f"#{int(r * 255):02X}{int(g * 255):02X}{int(b * 255):02X}"


def _annotate(page, i):
    if not i.get("bbox"):
        return
    r = pymupdf.Rect(*i["bbox"]) + (-1.5, -1.5, 1.5, 1.5)
    if r.width < 6:
        r.x1 = r.x0 + 6
    if r.height < 6:
        r.y1 = r.y0 + 6
    col = RGB.get(i["color"], RGB["comment"])
    if i["color"] == "revision":
        a = page.add_highlight_annot(r)
        a.set_colors(stroke=col)
        a.set_opacity(0.5)
    else:
        a = page.add_rect_annot(r)
        fill = {"missing": (1, 0.92, 0.80), "error": None}.get(i["color"])
        a.set_colors(stroke=col, fill=fill)
        a.set_border(width=1.4, dashes=[3, 2] if i["color"] in ("comment", "extract") else None)
        if fill:
            a.set_opacity(0.55)
    a.set_info(title=f"#{i['no']} {i['type']} ({i['severity']})",
               content=f"{i['field']}: {i.get('value') or '(blank)'}\n{i['description']}"
                       + (f"\nRequired / expected: {i['expected']}" if i.get("expected") else ""))
    a.set_flags(pymupdf.PDF_ANNOT_IS_PRINT)
    a.update()
    lab = pymupdf.Rect(r.x0 - 13, r.y0 - 1, r.x0 - 1, r.y0 + 8)
    if lab.x0 < 2:
        lab = pymupdf.Rect(r.x0, r.y0 - 10, r.x0 + 12, r.y0 - 1)
    t = page.add_freetext_annot(lab, str(i["no"]), fontsize=6.5, text_color=(1, 1, 1), fill_color=col, align=1)
    t.set_flags(pymupdf.PDF_ANNOT_IS_PRINT)
    t.update(text_color=(1, 1, 1), fill_color=col)


def markup_pdf(entries: list[dict], project: str, title: str = "Valve Datasheet Review — Mark-up",
               summary_pages: bool = True) -> bytes:
    """entries: [{path, pages, tag, doc_no, rev, filename, issues}] in output order."""
    out = pymupdf.open()
    toc, srcs = [], {}
    if len(entries) > 1:                                   # cover: counts per tag
        rows = "".join(
            f"<tr><td>{html.escape(e['tag'])}</td><td>{html.escape(e['doc_no'])}</td><td>{html.escape(e['rev'])}</td>"
            f"<td>{len(e['issues'])}</td><td>{sum(1 for i in e['issues'] if i['color'] == 'error')}</td>"
            f"<td>{sum(1 for i in e['issues'] if i['color'] == 'missing')}</td>"
            f"<td>{sum(1 for i in e['issues'] if i['color'] == 'revision')}</td>"
            f"<td>{sum(1 for i in e['issues'] if i['color'] in ('comment', 'extract'))}</td></tr>" for e in entries)
        cover = _story_pages(
            f"<h1>{html.escape(_ascii(title))}</h1><p>Project: {html.escape(project)} - generated "
            f"{dt.datetime.now():%Y-%m-%d %H:%M} - {sum(len(e['issues']) for e in entries)} issues on {len(entries)} datasheets.</p>"
            "<p>Box colours: <b>red</b> data error / PMS non-compliance, <b>orange</b> data missing / open item, "
            "<b>yellow</b> revision change, <b>blue</b> reviewer correction / comment. Each box has a number; the "
            "Issue Summary page after each datasheet explains it.</p>"
            "<table><tr><th>Tag</th><th>Doc. No.</th><th>Rev.</th><th>Issues</th><th>Error/PMS</th><th>Missing/Open</th>"
            f"<th>Revision</th><th>Reviewer</th></tr>{rows}</table>")
        out.insert_pdf(cover)
        toc.append([1, "Cover - issue counts", 1])
    for e in entries:
        src = srcs.get(e["path"]) or srcs.setdefault(e["path"], pymupdf.open(e["path"]))
        start = out.page_count
        pmap = {}
        for p in e["pages"]:
            if 1 <= p <= src.page_count:
                out.insert_pdf(src, from_page=p - 1, to_page=p - 1)
                pmap[p] = out.page_count - 1
        if not pmap:
            continue
        toc.append([1, f"{e['tag']}  Rev.{e['rev']}  ({len(e['issues'])})", start + 1])
        for i in e["issues"]:
            if i.get("page") in pmap:
                _annotate(out[pmap[i["page"]]], i)
        if summary_pages and e["issues"]:
            trs = "".join(
                f"<tr><td><span style='background-color:{_swatch(i['color'])};color:#FFFFFF'>&nbsp;{i['no']}&nbsp;</span></td>"
                f"<td>{i['sheet'] or ''} of {i['of'] or ''} (p.{i.get('page') or ''})</td><td>{i.get('row') or ''}</td>"
                f"<td>{html.escape(_ascii(i['field']))}</td><td>{html.escape(_ascii(i.get('value') or '(blank)'))}</td>"
                f"<td>{html.escape(i['type'])}</td><td>{html.escape(_ascii(i['description']))}</td>"
                f"<td>{html.escape(_ascii(i.get('expected') or ''))}</td><td>{html.escape(i['status'])}</td></tr>"
                for i in e["issues"])
            s = _story_pages(
                f"<h1>Issue Summary - {html.escape(e['tag'])}</h1><p>Doc. No. {html.escape(e['doc_no'])} - Rev. "
                f"{html.escape(e['rev'])} - file {html.escape(e.get('filename') or '')} - PDF pages "
                f"{', '.join(map(str, e['pages']))}</p><table><tr><th>No.</th><th>Sheet</th><th>Row</th><th>Field</th>"
                f"<th>Datasheet Value</th><th>Issue</th><th>Description</th><th>Required / Expected</th><th>Status</th></tr>"
                f"{trs}</table>")
            out.insert_pdf(s)
    for d in srcs.values():
        d.close()
    if out.page_count == 0:
        out.new_page()
        out[0].insert_text((72, 72), "No issues selected.", fontsize=12)
    out.set_toc(toc)
    return out.tobytes(garbage=3, deflate=True)
