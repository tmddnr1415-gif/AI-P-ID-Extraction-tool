"""L78 — 시스템 설명서 · 시운전 단계.

동작 학습(behavior.json — 명령마다 무엇이 켜고 · 막고 · 유지하고 · 끄는가)을 **시운전 절차**로 바꾼다.
  계통(System)  = 시트 이름의 계통 접두 (HH10_G111 → HH10) · 이름은 그 계통 G 시트 제목들이 함께 쓰는 낱말
  장비(Equipment) = sequence.commands 의 장비 (그룹 시험과 같은 단위)
  순서          = 계통은 도면 순서 · 계통 안 장비는 '선행' 관계(그 장비의 조건 신호가 다른 장비 명령 행에서 오는가) 위상 정렬
단계 (명령마다):
  base   아무 신호도 주지 않음 → 명령 꺼짐 (늘 켜짐이면 켜짐 — 확인)
  perm   꼭 필요한 신호를 하나씩 더한다 → 마지막(보통 누름단추)에서 명령 켜짐
  alt    다른 길
  hold   누름단추 놓음 → 자기유지면 유지 · 아니면 꺼짐 · 펄스면 n초 뒤 꺼짐
  stop   유지된 명령에 끄는 신호 → 꺼짐
  block  인터록 신호 → 막힘
  pair   반대 명령(START·STOP · OPEN·CLOSE) 조건을 함께 → 하나만
판정 0줄 — 기대값은 학습이 배운 것이고, '돌려 보기' 는 같은 묶음 시뮬레이터(grouptest._run)가 낸다.
"""
from __future__ import annotations

import io
import re
import time
from collections import Counter, defaultdict

from . import grouptest, refs, sequence

ACT_ORDER = {"START": 0, "OPEN": 0, "RUN": 1, "STOP": 2, "CLOSE": 2, "TRIP": 3}
SYS_RE = re.compile(r"^([A-Z]{2}\d{2}[A-Z]?)_")
TAG_RE = re.compile(r"\b\d{2}[A-Z]{3}\d{2}[A-Z]{2}\d{3}\b")


TYPICAL = "TYPICAL"


def _code(sheet_no):
    c = system_of(sheet_no)
    if c:
        return c
    return TYPICAL if re.fullmatch(r"\d{3,4}", sheet_no or "") else "기타"


def system_of(sheet_no: str) -> str:
    m = SYS_RE.match(sheet_no or "")
    return m.group(1) if m else ""


def _common_name(titles):
    """계통 이름 — G 시트 제목들의 절반 이상이 쓰는 낱말을 첫 제목 순서대로 (KKS 태그 · 'AUTO CONDITION' 은 뺀다)."""
    ts = [TAG_RE.sub("", t or "").replace("AUTO CONDITION", " ").split() for t in titles if t]
    if not ts:
        return ""
    cnt = Counter(w for t in ts for w in set(t))
    need = max(1, (len(ts) + 1) // 2)
    words = [w for w in ts[0] if cnt[w] >= need]
    seen, out = set(), []
    for w in words:
        if w not in seen:
            out.append(w)
            seen.add(w)
    return " ".join(out) or " ".join(ts[0])


def _row_fm(by_page, page, n):
    s = by_page.get(page) or {}
    for r in s.get("rows") or []:
        if r.get("side") == "L" and r.get("n") == n:
            return r.get("fm") or ""
    return ""


def structure(res: dict, learned: dict) -> dict:
    """계통 → 장비(순서) · 장비 사이 선행 관계."""
    sheets = res.get("sheets") or []
    by_page = {s["page"]: s for s in sheets}
    known = [s.get("sheet_no") for s in sheets if s.get("sheet_no")]
    items = [x for x in (learned or {}).get("items") or [] if x.get("equip")]
    # 명령 행 → 장비
    cmd_at = {}
    for x in items:
        cmd_at[(x.get("sheet"), x.get("row"))] = x["equip"]
    systems = {}
    order = []
    for s in sorted(sheets, key=lambda s: s["page"]):
        code = _code(s.get("sheet_no"))
        if code not in systems:
            systems[code] = {"code": code, "sheets": [], "equips": [], "first_page": s["page"]}
            order.append(code)
        systems[code]["sheets"].append({"page": s["page"], "sheet": s.get("sheet_no"), "title": s.get("title") or ""})
    eq = {}
    for x in items:
        e = eq.setdefault(x["equip"], {"equip": x["equip"], "items": [], "sheets": [], "pages": []})
        e["items"].append(x)
        if x.get("sheet") and x["sheet"] not in e["sheets"]:
            e["sheets"].append(x["sheet"])
            e["pages"].append(x["page"])
    # 선행 — 그 장비 명령의 조건 신호(입력 행)가 FM 으로 다른 장비의 명령 행에서 온다
    for e in eq.values():
        deps, links = set(), []
        for x in e["items"]:
            for l in x.get("labels") or []:
                if not l.get("row"):
                    continue
                fm = _row_fm(by_page, l["page"], l["row"])
                ref = refs.full(fm, known) if fm else None
                if not ref:
                    continue
                other = cmd_at.get((ref[0], ref[1]))
                if other and other != e["equip"]:
                    deps.add(other)
                    links.append({"from": other, "sheet": ref[0], "row": ref[1], "text": l.get("text"), "action": x.get("action")})
        e["deps"] = sorted(deps)
        e["dep_links"] = links[:20]
        e["system"] = Counter(_code(sh) for sh in e["sheets"]).most_common(1)[0][0] if e["sheets"] else "기타"
    for e in eq.values():
        systems.setdefault(e["system"], {"code": e["system"], "sheets": [], "equips": [], "first_page": 10 ** 9})
        if e["system"] not in order:
            order.append(e["system"])
    for code in order:
        S = systems[code]
        es = [e for e in eq.values() if e["system"] == code]
        S["equips"] = [e["equip"] for e in _topo(es)]
        S["name"] = "공통 매크로 (Typical) — 장비 구동 블록의 안쪽 로직" if code == TYPICAL else \
            (_common_name([t["title"] for t in S["sheets"] if "_G" in (t["sheet"] or "")]) or code)
    order.sort(key=lambda c: (c == TYPICAL, c == "기타", systems[c]["first_page"]))   # Typical 은 맨 뒤 (각 장비가 쓰는 공통 블록)
    out = [systems[c] for c in order if systems[c]["equips"] or systems[c]["sheets"]]
    return {"systems": out, "equips": eq}


def resolver(res, learned, sheets=None, cache=None, budget=0.0):
    """입력 행 신호 → 그 행의 FROM 이 가리키는 다른 시트의 명령.  시운전에서 '앞 시트의 진짜 조건' 으로 펼친다.
    배운 것(behavior.json — 장비 · 동작마다 대표 행 하나)에 없으면 budget 초 안에서 그 행을 그 자리에서 배운다
    (behavior.learn_one — 같은 시뮬레이터 · 판정 0줄).  배운 것은 이 resolver 안에 기억한다."""
    from . import behavior
    rs = res.get("sheets") or []
    by_page = {s["page"]: s for s in rs}
    by_name = {s.get("sheet_no"): s for s in rs if s.get("sheet_no")}
    known = list(by_name)
    at = {}
    for x in (learned or {}).get("items") or []:
        at.setdefault((x.get("sheet"), x.get("row")), x)
    t0 = [time.time()]
    memo = {}

    def get(l):
        if not l.get("row") or l.get("value") != 1:
            return None
        fm = _row_fm(by_page, l["page"], l["row"])
        ref = refs.full(fm, known) if fm else None
        if not ref:
            return None
        hit = at.get((ref[0], ref[1]))
        if hit is not None:
            return hit
        if ref in memo:
            return memo[ref]
        s = by_name.get(ref[0])
        if s is None or sheets is None or time.time() - t0[0] > budget:
            return None
        if not any(r.get("side") == "R" and r.get("n") == ref[1] for r in s.get("rows") or []):
            memo[ref] = None
            return None
        try:
            u = behavior.learn_one(sheets, s["page"], ref[1], cache=cache)
        except Exception:
            u = None
        if u is not None:
            r = next((r for r in s.get("rows") or [] if r.get("side") == "R" and r.get("n") == ref[1]), {})
            u = dict(u, sheet=ref[0], desc=r.get("desc") or "", action="")
        memo[ref] = u
        return u
    return get


def _topo(es):
    """선행 장비가 먼저 · 같은 깊이는 시트 순서 (고리면 시트 순서 그대로)."""
    names = {e["equip"] for e in es}
    by = {e["equip"]: e for e in es}
    depth = {}

    def d(n, stack=()):
        if n in depth:
            return depth[n]
        if n in stack:
            return 0
        v = 0
        for p in by[n]["deps"]:
            if p in names:
                v = max(v, d(p, stack + (n,)) + 1)
        depth[n] = v
        return v
    for e in es:
        d(e["equip"])
    return sorted(es, key=lambda e: (depth[e["equip"]], min(e["pages"] or [10 ** 9]), e["equip"]))


# ───────────────────────────── 단계 ─────────────────────────────

def _F(forces):
    return {int(p): dict(v) for p, v in (forces or {}).items()}


def _merge(*Fs):
    G = {}
    for F in Fs:
        for p, d in (F or {}).items():
            G.setdefault(int(p), {}).update(d)
    return G


def _lab(l):
    t = l.get("text") or l.get("key")
    at = f"{l.get('sheet') or ''}-{l['row']}" if l.get("row") else (l.get("sheet") or "")
    return t, at


def _sig(l, value=None):
    t, at = _lab(l)
    return {"text": t, "at": at, "page": l.get("page"), "row": l.get("row"), "key": l.get("key"),
            "value": l.get("value") if value is None else value, "assumed": bool(l.get("assumed")), "via": l.get("via") or []}


def _step(kind, x, title, forces, expect, mode, signals=(), chain=(), note="", added=None):
    return {"kind": kind, "action": x.get("action"), "title": title, "forces": {str(p): v for p, v in forces.items()},
            "chain": [{str(p): v for p, v in c.items()} for c in chain], "expect": expect, "mode": mode,
            "signals": list(signals), "added": added, "note": note,
            "cmd": {"page": x.get("page"), "row": x.get("row"), "sheet": x.get("sheet"), "net": x.get("net"), "desc": x.get("desc")}}


def _expand(x, resolve, depth=0, seen=None):
    """명령의 조건 신호를 펼친다 — 다른 시트 명령에서 오는 입력 행은 그 명령의 조건으로 (최대 4단 · 고리 없음)."""
    seen = set(seen or ()) | {(x.get("sheet"), x.get("row"))}
    out = []
    for l in x.get("labels") or []:
        u = resolve(l) if resolve and depth < 4 else None
        if u is not None and not u.get("reached"):
            q = dict(l)
            q["upstream_gap"] = f"{u.get('sheet')}-{u.get('row')} {u.get('desc') or ''} — 앞 시트에서 시뮬레이션으로 세우지 못함 ({u.get('why') or '안 섬'})"
            out.append(q)
            continue
        if u is not None and (u.get("sheet"), u.get("row")) not in seen:
            via = f"{u.get('sheet')}-{u.get('row')} {u.get('desc') or u.get('action')}"
            sub = [q for q in _expand(u, resolve, depth + 1, seen) if not q.get("assumed")]
            if not sub:                                   # 앞 명령이 신호 없이 서거나 고리 — 그 입력 행을 그대로 준다
                out.append(dict(l))
                continue
            for q in sub:
                q = dict(q)
                q["via"] = [via] + list(q.get("via") or [])
                out.append(q)
            continue
        out.append(dict(l))
    uniq, keys = [], set()
    for l in out:
        k = (l.get("page"), l.get("key"))
        if k not in keys:
            keys.add(k)
            uniq.append(l)
    return uniq


def steps_of(items, opposites=None, resolve=None):
    """장비 하나 — 시운전 단계 (기대값은 학습이 배운 것).  resolve 가 있으면 앞 시트 명령의 조건까지 펼친다."""
    out = []
    items = sorted(items, key=lambda x: (ACT_ORDER.get(x.get("action"), 5), x.get("page") or 0, x.get("row") or 0))
    for x in items:
        act = x.get("action")
        where = f"{x.get('sheet')}-{x.get('row')}"
        if not x.get("reached") or x.get("net") is None:
            b = x.get("blocked_at") or {}
            st = _step("gap", x, f"{act} — 시뮬레이션으로 세우지 못함", {}, None, "none",
                       note=(x.get("why") or "안 섬") + (f" · 막힌 곳 {b.get('label')}" if b.get("label") else "")
                       + " — 현장에서 도면 확인")
            out.append(st)
            continue
        labels = _expand(x, resolve)
        F = {}
        for l in labels:
            F.setdefault(int(l["page"]), {})[l["key"]] = l["value"]
        assumed = [l for l in labels if l.get("assumed")]
        real = [l for l in labels if not l.get("assumed")]
        A = {}
        for l in assumed:
            A.setdefault(int(l["page"]), {})[l["key"]] = l["value"]
        out.append(_step("base", x, f"{act} 출력 확인 전 — 아무 신호 없음 ({where})", A,
                         1 if x.get("always_on") else 0, "last",
                         signals=[_sig(l) for l in assumed],
                         note="신호 없이 켜짐 — 안전 쪽 설계인지 확인" if x.get("always_on") else ""))
        zeros = [l for l in real if l.get("value") == 0]
        perm = [l for l in real if l.get("value") != 0 and not str(l.get("key", "")).startswith("block:")]
        btn = [l for l in real if l.get("value") != 0 and str(l.get("key", "")).startswith("block:")]
        seq = zeros + perm + btn
        acc = {p: dict(d) for p, d in A.items()}
        for i, l in enumerate(seq):
            acc.setdefault(int(l["page"]), {})[l["key"]] = l["value"]
            last = i == len(seq) - 1
            kind_ko = "누름단추" if str(l.get("key", "")).startswith("block:") else ("없어야 하는 신호" if l.get("value") == 0 else "허가 조건")
            t, at = _lab(l)
            if l.get("via"):
                at += " → " + " → ".join(reversed(l["via"]))
            out.append(_step("perm", x, f"{kind_ko} {i + 1}/{len(seq)}: {t} = {l.get('value')}" + (f" ({at})" if at else ""),
                             {p: dict(d) for p, d in acc.items()}, 1 if last else 0, "any" if last else "last",
                             signals=[_sig(q) for q in seq[:i + 1]], added=_sig(l),
                             note=(f"{act} 명령 켜짐 ({where})" if last else f"아직 {act} 꺼져 있어야 함 — 남은 조건 {len(seq) - i - 1}")
                             + (f" · 앞 단계: {l['upstream_gap']}" if l.get("upstream_gap") else "")))
        if not seq:
            out[-1]["note"] = (out[-1]["note"] + " · " if out[-1]["note"] else "") + "꼭 필요한 신호가 없음 — 늘 켜짐"
        for i, a in enumerate(x.get("alts") or []):
            G = _F(a.get("forces"))
            ls = [l for l in a.get("labels") or [] if not l.get("assumed")]
            out.append(_step("alt", x, f"다른 길 {i + 1}: " + " + ".join(_lab(l)[0] for l in ls)[:200], G, 1, "any",
                             signals=[_sig(l) for l in ls], note=f"{act} 명령 켜짐"))
        latch = x.get("latch")
        rest = {p: {k: v for k, v in d.items() if not (x.get("latch_released") == "all" or k.startswith("block:"))}
                for p, d in F.items()}
        rest = {p: d for p, d in rest.items() if d}
        what = "누름단추 놓음" if x.get("latch_released") == "buttons" else "조건 신호 거둠"
        if latch is True:
            out.append(_step("hold", x, f"{what} → {act} 자기유지", rest, 1, "last", chain=[F],
                             note="자기유지(래치) — 신호를 거둬도 명령이 남음"))
            for l in x.get("stoppers") or []:
                G = _merge(rest, {int(l["page"]): {l["key"]: 1}})
                t, at = _lab(l)
                out.append(_step("stop", x, f"정지 신호: {t} = 1" + (f" ({at})" if at else "") + f" → {act} 꺼짐", G, 0, "last",
                                 chain=[F, rest], signals=[_sig(l, 1)], added=_sig(l, 1)))
        elif latch is False:
            out.append(_step("hold", x, f"{what} → {act} 꺼짐", rest, 0, "last", chain=[F],
                             note="누르는 동안만 — 자기유지 없음"))
        elif x.get("pulse_s") is not None:
            out.append(_step("hold", x, f"{act} 펄스 — 약 {x['pulse_s']}초 뒤 꺼짐", F, 0, "last",
                             note=f"펄스 출력 (약 {x['pulse_s']}초)"))
        for l in x.get("interlocks") or []:
            G = _merge(F, {int(l["page"]): {l["key"]: 1}})
            t, at = _lab(l)
            out.append(_step("block", x, f"인터록: {t} = 1" + (f" ({at})" if at else "") + f" → {act} 막힘", G, 0, "never",
                             signals=[_sig(l, 1)], added=_sig(l, 1), note="인터록 · 트립 — 이 신호가 있으면 명령이 나가지 않음"))
    by = {x["action"]: x for x in items if x.get("reached") and x.get("net") is not None}
    for a, b in opposites if opposites is not None else (sequence.vocab().get("opposites") or []):
        if a in by and b in by:
            A, B = by[a], by[b]
            G = _merge(_F(A["forces"]), _F(B["forces"]))
            st = _step("pair", A, f"{a} · {b} 조건을 함께 → 둘 중 하나만 나가야 함", G, "one", "pair",
                       note="반대 명령 동시 금지 확인")
            st["pair"] = [{"action": a, "page": A["page"], "net": A["net"], "row": A["row"], "sheet": A.get("sheet")},
                          {"action": b, "page": B["page"], "net": B["net"], "row": B["row"], "sheet": B.get("sheet")}]
            out.append(st)
    for i, s in enumerate(out):
        s["no"] = i + 1
    return out


def _say(v):
    return "1" if v == 1 else "0" if v == 0 else "모름"


def run_steps(sheets, steps, cache, timers=None, upto=None):
    """단계를 돌린다 — 앞 단계(chain)를 같은 상태로 먼저 돌린 뒤 그 단계의 신호로.  결과는 단계에 got · status 로 붙여 돌려준다."""
    out = []
    memo = {}
    for i, s in enumerate(steps):
        if upto is not None and i >= upto:
            break
        r = dict(s)
        if s["mode"] == "none":
            r.update(got="안 섬", status="gap")
            out.append(r)
            continue
        try:
            if s["mode"] == "pair":
                F = _F(s["forces"])
                vals = []
                for p in s["pair"]:
                    v, _ = grouptest._run(sheets, F, None, p["page"], p["net"], cache, timers)
                    vals.append(bool(v) and v[-1] == 1)
                both = all(vals)
                r.update(got="둘 다 1" if both else "하나만", status="warn" if both else "pass",
                         why="서로 막는 신호 없음 — 현장 확인" if both else "")
                out.append(r)
                continue
            c = s["cmd"]
            st = None
            key = ""
            for ch in s.get("chain") or []:
                key += repr(sorted((p, sorted(d.items())) for p, d in ch.items()))
                if key in memo:
                    st = memo[key]
                    continue
                _, st = grouptest._run(sheets, _F(ch), st, c["page"], c["net"], cache, timers)
                memo[key] = st
            v, _ = grouptest._run(sheets, _F(s["forces"]), st, c["page"], c["net"], cache, timers)
            last = v[-1] if v else None
            if s["mode"] == "any":
                ok = 1 in v
                got = 1 if ok else last
            elif s["mode"] == "never":
                ok = 1 not in v
                got = 1 if not ok else 0
            else:
                got = last
                ok = (last == 1) == (s["expect"] == 1)
            r["got"] = _say(got)
            r["pulse"] = s["mode"] == "any" and ok and last != 1
            if ok:
                r["status"] = "warn" if s["kind"] == "base" and s["expect"] == 1 else "pass"
            else:
                early = s["kind"] == "perm" and s["expect"] == 0 and got == 1
                r["status"] = "warn" if early else "fail"
                r["why"] = "이 단계에서 벌써 켜짐 — 앞 조건만으로도 섬" if early else "학습 때와 다름"
        except Exception as ex:                        # 한 단계가 실패해도 나머지를 돌린다 — 사유를 남긴다
            r.update(got="오류", status="fail", why=f"시뮬레이션 오류 — {ex}")
        out.append(r)
    n = Counter(r.get("status") for r in out)
    return {"steps": out, "counts": dict(n),
            "status": next((s for s in ("fail", "warn", "gap", "pass") if n.get(s)), "skip")}


# ───────────────────────────── 화면용 목록 ─────────────────────────────

def overview(res, learned):
    st = structure(res, learned)
    eqs = st["equips"]
    out = []
    no = 0
    for S in st["systems"]:
        row = {"code": S["code"], "name": S["name"], "sheets": len(S["sheets"]), "equips": []}
        for name in S["equips"]:
            e = eqs[name]
            no += 1
            acts = sorted({x.get("action") for x in e["items"]}, key=lambda a: ACT_ORDER.get(a, 5))
            row["equips"].append({"no": no, "equip": name, "sheets": e["sheets"], "actions": acts, "deps": e["deps"],
                                  "reached": sum(1 for x in e["items"] if x.get("reached")), "commands": len(e["items"]),
                                  "interlocks": sum(len(x.get("interlocks") or []) for x in e["items"])})
        if row["equips"]:
            out.append(row)
    return out


# ───────────────────────────── Word ─────────────────────────────

LATCH_KO = {True: "자기유지(래치)", False: "누르는 동안만", None: "펄스"}


def _latch(x):
    if x.get("latch") is True:
        return "자기유지(래치)"
    if x.get("latch") is False:
        return "누르는 동안만 (자기유지 없음)"
    if x.get("pulse_s") is not None:
        return f"펄스 (약 {x['pulse_s']}초)"
    return "—"


def docx_bytes(res, learned, status=None, review=None, dangling=None, progress=None, budget=120.0) -> bytes:
    """전체 System Description (Word) — 계통 · 장비 · 명령마다 기동 조건 · 인터록 · 정지 · 유지 · 시운전 단계."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    status = status or {}
    st = structure(res, learned)
    eqs = st["equips"]
    sheets = res.get("sheets") or []
    summ = (learned or {}).get("summary") or {}
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    base = doc.styles["Normal"]
    base.font.name = "Malgun Gothic"
    base.font.size = Pt(9.5)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    for nm in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            hs = doc.styles[nm]
            hs.font.name = "Malgun Gothic"
            hs.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        except Exception:
            pass

    def table(head, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = h
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(8.5)
        for row in rows:
            cs = t.add_row().cells
            for i, v in enumerate(row):
                cs[i].text = str(v if v is not None else "")
                for p in cs[i].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        return t

    name = status.get("name") or "LOGIC"
    proj = status.get("project") or ""
    title = doc.add_heading("System Description", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{proj + ' — ' if proj else ''}{name}").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"작성 {time.strftime('%Y-%m-%d')} · LOGIC 분석 도구가 도면 로직을 읽고 시뮬레이션으로 확인해 만든 문서").italic = True

    doc.add_heading("1. 이 문서에 대하여", 1)
    doc.add_paragraph(
        "이 문서는 Control Logic 도면(PDF)을 도구가 읽어 게이트 · 타이머 · 메모리 · 매크로와 시트 사이 FROM/TO 연결을 해석하고, "
        "명령(START · STOP · OPEN · CLOSE · RUN · TRIP)마다 같은 로직 시뮬레이터로 신호를 켜고 꺼 보며 배운 내용을 정리한 것입니다. "
        "조건 · 인터록 · 자기유지는 사람이 쓴 설명이 아니라 시뮬레이션 결과입니다 — 도면을 읽지 못한 자리(판정 못 한 블록 · 끊긴 선)는 "
        "그렇다고 적었습니다.  시운전 단계는 화면 'LOGIC → 📘 시운전' 에서 단계마다 시뮬레이션으로 다시 돌려 볼 수 있습니다.")
    rows = [("도면 쪽 수 / 로직 시트", f"{res.get('pdf_pages') or status.get('pages') or ''} 쪽 / {len(sheets)} 장"),
            ("계통 (System)", f"{sum(1 for S in st['systems'] if S['equips'])} 개"),
            ("장비 (Equipment)", f"{len(eqs)} 개"),
            ("명령", f"{summ.get('commands', 0)} 개 — 시뮬레이션으로 섬 {summ.get('reached', 0)} · 못 세움 {summ.get('unreached', 0)}"),
            ("명령 성격", f"자기유지 {summ.get('latched', 0)} · 누르는 동안만 {summ.get('momentary', 0)} · 펄스 {summ.get('pulse', 0)}"),
            ("인터록이 있는 명령", f"{summ.get('with_interlock', 0)} 개"),
            ("반대 명령 짝 (START/STOP · OPEN/CLOSE)", f"서로 막음 {summ.get('exclusive', 0)} · 함께 나갈 수 있음 {len(summ.get('both_on') or [])} (확인 대상)"),
            ("신호 없이 켜지는 명령", f"{len(summ.get('always_on') or [])} 개 (확인 대상)")]
    if status.get("idrate") is not None:
        rows.append(("도면 읽기 (식별률)", f"{round(status['idrate'] * 100, 1) if status['idrate'] <= 1 else status['idrate']} %"))
    if dangling:
        rows.append(("연결 안 된 선 끝", dangling))
    table(["항목", "내용"], rows, [6, 11])

    # 2. 기능 블록 (범례)
    doc.add_heading("2. 로직 기능 (범례)", 1)
    doc.add_paragraph("도면이 쓰는 기능 블록과 그 동작 (범례 기준 — 시뮬레이터가 이 정의로 돈다).")
    try:
        from . import functions
        fb = functions.book()
        used = Counter(b.get("logic") for s in sheets for b in (s.get("blocks") or {}).values())
        frows = []
        for k, n in used.most_common():
            f = fb.get(k)
            if f and n:
                frows.append((f.get("ko") or k, n, f.get("does") or ""))
        if frows:
            table(["기능", "도면 개수", "동작"], frows, [3.5, 2, 11.5])
    except Exception:
        pass
    mac = Counter()
    for s in sheets:
        for b in (s.get("blocks") or {}).values():
            m = b.get("macro") or {}
            if m.get("sheet"):
                mac[(m.get("sheet"), m.get("title") or "")] += 1
    if mac:
        doc.add_paragraph("매크로(Typical) — 장비 구동 블록은 Typical 시트의 로직으로 펼쳐 같은 시간으로 돈다.")
        table(["Typical 시트", "제목", "쓰인 수"], [(k[0], k[1], n) for k, n in mac.most_common()], [3, 11, 3])

    # 3. 공통 신호
    if summ.get("common_need") or summ.get("common_block"):
        doc.add_heading("3. 여러 명령이 함께 쓰는 신호", 1)
        if summ.get("common_need"):
            doc.add_paragraph("여러 명령이 함께 기다리는 허가 신호:")
            table(["시트", "신호", "명령 수"], [(c["sheet"], c["text"], c["count"]) for c in summ["common_need"]], [3, 12, 2])
        if summ.get("common_block"):
            doc.add_paragraph("여러 명령을 막는 신호 (공통 트립 · 인터록):")
            table(["시트", "신호", "명령 수"], [(c["sheet"], c["text"], c["count"]) for c in summ["common_block"]], [3, 12, 2])

    # 4. 계통별 설명
    doc.add_heading("4. 계통별 설명", 1)
    sysno = 0
    for S in st["systems"]:
        if not S["equips"]:
            continue
        sysno += 1
        doc.add_heading(f"4.{sysno} {S['code']} — {S['name']}", 2)
        gs = [t for t in S["sheets"] if "_G" in (t["sheet"] or "")]
        ps = [t for t in S["sheets"] if "_P" in (t["sheet"] or "")]
        doc.add_paragraph(f"로직 시트 {len(gs)} 장 · 루프 시트 {len(ps)} 장 · 장비 {len(S['equips'])} 개.  "
                          "장비는 선행 관계 순서(다른 장비의 명령을 조건으로 받는 장비가 뒤)로 적었습니다.")
        for name in S["equips"]:
            e = eqs[name]
            doc.add_heading(name, 3)
            line = f"시트 {', '.join(e['sheets'])}"
            if e["deps"]:
                line += f" · 선행 장비: {', '.join(e['deps'])}"
            doc.add_paragraph(line)
            for x in sorted(e["items"], key=lambda x: (ACT_ORDER.get(x.get("action"), 5), x.get("page") or 0)):
                act = x.get("action")
                p = doc.add_paragraph()
                r = p.add_run(f"{act}  ")
                r.bold = True
                p.add_run(f"{x.get('desc') or ''} ({x.get('sheet')}-{x.get('row')})")
                if not x.get("reached"):
                    b = x.get("blocked_at") or {}
                    doc.add_paragraph(f"시뮬레이션으로 세우지 못함 — {x.get('why') or ''}"
                                      + (f" (막힌 곳: {b.get('label')})" if b.get("label") else ""), style="List Bullet")
                    continue
                real = [l for l in x.get("labels") or [] if not l.get("assumed")]
                need1 = [l for l in real if l.get("value") != 0]
                need0 = [l for l in real if l.get("value") == 0]
                if x.get("always_on"):
                    doc.add_paragraph("아무 신호 없이도 켜져 있음 — 안전 쪽 설계(건강 신호가 와야 풀림)인지 확인.", style="List Bullet")
                if need1:
                    doc.add_paragraph("기동 조건 (모두): " + " · ".join(f"{_lab(l)[0]} [{_lab(l)[1]}]" for l in need1), style="List Bullet")
                if need0:
                    doc.add_paragraph("없어야 하는 신호: " + " · ".join(f"{_lab(l)[0]} [{_lab(l)[1]}]" for l in need0), style="List Bullet")
                for i, a in enumerate(x.get("alts") or []):
                    ls = [l for l in a.get("labels") or [] if not l.get("assumed")]
                    doc.add_paragraph(f"다른 길 {i + 1}: " + " · ".join(f"{_lab(l)[0]}" for l in ls), style="List Bullet")
                doc.add_paragraph("출력 성격: " + _latch(x), style="List Bullet")
                if x.get("interlocks"):
                    doc.add_paragraph("인터록 (있으면 안 나감): " + " · ".join(f"{_lab(l)[0]} [{_lab(l)[1]}]" for l in x["interlocks"]),
                                      style="List Bullet")
                if x.get("stoppers"):
                    doc.add_paragraph("정지 (유지된 명령을 끔): " + " · ".join(f"{_lab(l)[0]}" for l in x["stoppers"]), style="List Bullet")
                if x.get("drops"):
                    doc.add_paragraph("이 허가를 잃으면 꺼짐: " + " · ".join(f"{_lab(l)[0]}" for l in x["drops"]), style="List Bullet")
                if x.get("effects"):
                    doc.add_paragraph("명령이 서면 함께 켜지는 출력: " + " · ".join(
                        f"{f.get('text')} ({f.get('sheet')}-{f.get('row')})" for f in x["effects"][:8]), style="List Bullet")
                if x.get("holds"):
                    doc.add_paragraph(f"도면 읽기 빈틈 {x['holds']} 곳을 '만족' 으로 보고 세웠음 — 현장 확인.", style="List Bullet")
            pairs = [p for p in (learned or {}).get("pairs") or [] if p.get("equip") == name]
            for pr in pairs:
                doc.add_paragraph(f"{pr['a']} · {pr['b']}: " + ("서로 막음 (동시 지령 불가)" if pr.get("exclusive")
                                                                 else "함께 나갈 수 있음 — 설계 확인" if pr.get("both") else "—"),
                                  style="List Bullet")

    # 5. 시운전 절차
    doc.add_heading("5. 시운전 절차 (단계별 동작)", 1)
    doc.add_paragraph(
        "계통 순서 · 장비 선행 순서대로 진행합니다.  장비마다 ① 신호 없음 확인 → ② 허가 조건을 하나씩 넣으며 명령이 마지막에만 "
        "나오는지 확인 → ③ 누름단추를 놓았을 때 유지/꺼짐 확인 → ④ 정지 신호 → ⑤ 인터록마다 명령이 막히는지 → ⑥ 반대 명령 동시 금지 "
        "순서입니다.  기대값은 시뮬레이션이 배운 값이고, 화면 'LOGIC → 📘 시운전' 에서 단계마다 '▶ 시뮬레이션' 으로 다시 돌려 봅니다.")
    doc.add_heading("5.0 공통 준비", 2)
    for t in ("전원 · DCS 통신 정상, 제어 모드(REMOTE/AUTO) 확인",
              "입력 신호(I/O) 점검 — 로직 시트 입력 행의 FROM(FM) 이 가리키는 신호가 실제로 들어오는지",
              "시트 사이 FROM/TO 정합성 — 'LOGIC → 정합성 검토' 의 설계 오류 0 확인",
              "타이머 설정값 확인 — 도면 값이 없는 타이머는 시뮬레이션 기본 5초로 돌았음",
              "인터록 · 트립 시험 신호를 줄 수 있는 방법(강제 · 시뮬레이터) 준비"):
        doc.add_paragraph(t, style="List Number")
    sysno = 0
    eno = 0
    opp = sequence.vocab().get("opposites") or []
    resolve = resolver(res, learned, (res.get("sheets") or []) + list(res.get("library_sheets") or []), {}, budget)
    for S in st["systems"]:
        if not S["equips"]:
            continue
        sysno += 1
        doc.add_heading(f"5.{sysno} {S['code']} — {S['name']}", 2)
        for name in S["equips"]:
            eno += 1
            e = eqs[name]
            if progress:
                progress(eno, len(eqs))
            doc.add_heading(f"[{eno}] {name}", 3)
            if e["deps"]:
                doc.add_paragraph("먼저: " + ", ".join(e["deps"]))
            stp = steps_of(e["items"], opp, resolve)
            table(["#", "단계", "기대 결과"],
                  [(s["no"], s["title"], _expect_ko(s)) for s in stp], [1, 12, 4])
    # 6. 확인 대상
    doc.add_heading("6. 설계 확인 대상", 1)
    both = summ.get("both_on") or []
    if both:
        doc.add_paragraph("반대 명령이 함께 나갈 수 있는 장비 (서로 막는 신호가 로직 안에 없음 — 현장 신호 조합에 기댐):")
        table(["장비", "짝", "시트"], [(p["equip"], f"{p['a']} · {p['b']}", f"{(p.get('a_at') or [''])[0]}") for p in both[:200]], [8, 4, 5])
    ao = summ.get("always_on") or []
    if ao:
        doc.add_paragraph("아무 신호 없이 켜지는 명령:")
        table(["장비", "명령", "시트-행"], [(a["equip"], a["action"], f"{a['sheet']}-{a['row']}") for a in ao], [9, 3, 5])
    un = [x for x in (learned or {}).get("items") or [] if not x.get("reached")]
    if un:
        doc.add_paragraph("시뮬레이션으로 세우지 못한 명령 (도면 읽기 빈틈 · 현장 되먹임 필요):")
        table(["장비", "명령", "시트-행", "사유"],
              [(x.get("equip"), x.get("action"), f"{x.get('sheet')}-{x.get('row')}", x.get("why") or "") for x in un], [6, 2, 3.5, 5.5])
    if review:
        doc.add_paragraph(f"정합성 검토 (LOGIC → 📋 정합성 검토): 설계 오류 {review.get('error', 0)} · 설계 확인 {review.get('check', 0)} · "
                          f"읽기 빈틈 {review.get('gap', 0)}")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _expect_ko(s):
    if s["mode"] == "none":
        return "시험 불가 — " + (s.get("note") or "")
    if s["mode"] == "pair":
        return "하나만 나감"
    a = s.get("action") or "명령"
    v = s["expect"]
    if s["kind"] == "block":
        return f"{a} 안 나감 (막힘)"
    if s["mode"] == "any" and v == 1:
        return f"{a} 켜짐 (1)"
    return f"{a} {'켜짐 (1)' if v == 1 else '꺼짐 (0)'}"
