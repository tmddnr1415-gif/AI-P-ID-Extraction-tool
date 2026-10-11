"""LOGIC 점검 결과를 묶음 단위로 펴고 Excel 로 낸다 (L7).  판정 0 — 장마다의 `findings` · `rows` 를 옮겨 적을 뿐이다."""
from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# 점검 코드 → 사람 말.  모르는 코드는 코드 그대로 보인다 (숨기지 않는다).
LABELS = {
    "FT_MISMATCH": "from-to 불일치",
    "FT_DESC_DIFF": "from-to 설명 다름",
    "C1_DANGLING": "선 끝이 아무 데도 닿지 않음",
    "C1_ARROW_NOWHERE": "화살촉이 가리키는 블록 없음",
    "C3_GATE_NO_INPUT": "입력 없는 게이트",
    "C3_BLOCK_UNCONNECTED": "선이 닿지 않은 블록",
    "C7_CONNECTOR_NO_SOURCE": "연결 기호의 상대 없음",
    "C7_CONNECTOR_NO_SOURCE_OUT": "연결 기호의 상대 없음 (장 밖일 수 있음)",
    "BLOCK_UNKNOWN": "판정 못 한 블록",
}
ORDER = list(LABELS)


def findings(sheets):
    items, counts = [], {}
    for s in sheets:
        for f in s.get("findings", []):
            counts[f["code"]] = counts.get(f["code"], 0) + 1
            items.append({"page": s["page"], "sheet_no": s.get("sheet_no"), "code": f["code"],
                          "label": LABELS.get(f["code"], f["code"]), "text": f.get("text", ""),
                          "row": f.get("row"), "side": f.get("side"), "block": f.get("block"), "at": f.get("at")})
    rank = {c: i for i, c in enumerate(ORDER)}
    items.sort(key=lambda x: (rank.get(x["code"], len(ORDER)), x["page"], x.get("row") or 0))
    codes = [{"code": c, "label": LABELS.get(c, c), "n": counts[c]}
             for c in sorted(counts, key=lambda c: (rank.get(c, len(ORDER)), c))]
    return {"items": items, "codes": codes, "total": len(items)}


def _sheet(wb, title, header):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE6F7")
    ws.freeze_panes = "A2"
    return ws


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def workbook(res, name="") -> bytes:
    sheets = res.get("sheets", [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.append(["LOGIC 검토", name])
    ws.append(["PDF 쪽수", res.get("pdf_pages"), "Logic 표 시트", len(sheets), "분석 시간(초)", res.get("seconds")])
    ws.append([])
    ws.append(["점검 종류", "건수"])
    fd = findings(sheets)
    for c in fd["codes"]:
        ws.append([c["label"], c["n"]])
    ws.append(["합계", fd["total"]])
    ws.append([])
    ws.append(["시트", "쪽", "제목", "블록", "선", "입력 행", "출력 행", "from-to 불일치", "설명 다름", "점검 합계"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for s in sheets:
        fs = s.get("findings", [])
        ws.append([s.get("sheet_no"), s["page"], s.get("title"), len(s.get("blocks", {})), len(s.get("nets", [])),
                   sum(1 for r in s.get("rows", []) if r["side"] == "L"),
                   sum(1 for r in s.get("rows", []) if r["side"] == "R"),
                   sum(1 for f in fs if f["code"] == "FT_MISMATCH"),
                   sum(1 for f in fs if f["code"] == "FT_DESC_DIFF"), len(fs)])
    _widths(ws, [14, 8, 40, 8, 8, 10, 10, 14, 10, 10])

    ws = _sheet(wb, "from-to 불일치", ["시트", "쪽", "쪽(L입력·R출력)", "행", "상대", "사유"])
    for s in sheets:
        for r in s.get("rows", []):
            ft = r.get("ft")
            if ft and ft.get("ok") is False:
                ws.append([s.get("sheet_no"), s["page"], r["side"], r["n"], ft.get("peer"), " · ".join(ft.get("why") or [])])
    _widths(ws, [10, 6, 14, 6, 20, 90])

    ws = _sheet(wb, "설명 다름", ["시트", "쪽", "행", "상대", "이 쪽 설명", "상대 설명"])
    for s in sheets:
        for r in s.get("rows", []):
            ft = r.get("ft")
            if ft and ft.get("warn") and ft.get("ok") is not False:
                peer = ft.get("peer") or ""
                other = ""
                for w in ft["warn"]:
                    if "↔" in w:
                        other = w.split("↔", 1)[1].strip().strip("'")
                ws.append([s.get("sheet_no"), s["page"], r["n"], peer, r.get("desc"), other])
    _widths(ws, [10, 6, 6, 20, 50, 50])

    ws = _sheet(wb, "연결 점검", ["시트", "쪽", "점검", "코드", "행", "블록", "자리 x", "자리 y", "내용"])
    for it in fd["items"]:
        if it["code"].startswith("FT_"):
            continue
        at = it.get("at") or [None, None]
        ws.append([it["sheet_no"], it["page"], it["label"], it["code"], it.get("row"), it.get("block"), at[0], at[1], it["text"]])
    _widths(ws, [10, 6, 28, 26, 6, 8, 9, 9, 70])

    ws = _sheet(wb, "입력 행", ["시트", "쪽", "행", "FM", "설명", "KKS", "from-to 판정", "상대", "비고"])
    for s in sheets:
        for r in s.get("rows", []):
            if r["side"] != "L":
                continue
            ft = r.get("ft")
            verdict = "" if not ft else ("불일치" if ft.get("ok") is False else ("설명 다름" if ft.get("warn") else "일치"))
            note = " · ".join((ft or {}).get("why") or (ft or {}).get("warn") or [])
            ws.append([s.get("sheet_no"), s["page"], r["n"], r.get("fm"), r.get("desc"), r.get("kks"), verdict, (ft or {}).get("peer"), note])
    _widths(ws, [10, 6, 6, 12, 50, 18, 12, 20, 70])

    ws = _sheet(wb, "출력 행", ["시트", "쪽", "행", "KKS", "설명", "TO", "from-to 판정", "상대", "비고"])
    for s in sheets:
        for r in s.get("rows", []):
            if r["side"] != "R":
                continue
            ft = r.get("ft")
            verdict = "" if not ft else ("불일치" if ft.get("ok") is False else "일치")
            note = " · ".join((ft or {}).get("why") or [])
            ws.append([s.get("sheet_no"), s["page"], r["n"], r.get("kks"), r.get("desc"), r.get("to"), verdict, (ft or {}).get("peer"), note])
    _widths(ws, [10, 6, 6, 18, 50, 18, 12, 24, 70])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
