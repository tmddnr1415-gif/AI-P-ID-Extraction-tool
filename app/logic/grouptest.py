"""L36 — 블록 그룹(장비) 원 버튼 시험.

분석이 끝나면 동작 학습(behavior.json)이 명령마다 '무엇이 켜고 · 막고 · 유지하고 · 끄는가' 를 배운다.  이 모듈은 그 배운
것을 **시험 항목**으로 바꿔 지금 결과(사람이 고친 값까지 얹은)로 다시 돌린다 — 장비 하나를 누르면 그 장비의 모든 명령을
한 번에.  판정 0줄: 결과는 언제나 같은 묶음 시뮬레이터가 낸다.

시험 항목 (명령마다):
  on      배운 조합을 주면 명령이 선다                       — 안 서면 실패 (학습 때와 달라졌거나 지금 도면으로 안 섬)
  alt     다른 길(다른 최소 조합)로도 선다
  block   막는 신호를 더하면 안 선다 (인터록 · 트립)
  hold    누름단추를 놓았을 때 — 자기유지면 남고, 아니면 꺼진다 (학습 때와 같은가)
  stop    자기유지된 명령을 끄는 신호가 정말 끈다
  always  아무 신호 없이 서 있는 명령 — 확인 대상
  tied    도면이 상수로 묶어 둔 길로만 서는 명령 — 확인 대상
장비마다 (반대 명령 짝):
  pair    START·STOP / OPEN·CLOSE 조건을 함께 주어도 둘이 함께 나가지 않는다   — 함께 나가면 확인
  cancel  자기유지된 START(OPEN) 에 STOP(CLOSE) 조건을 주면 꺼진다             — 남으면 확인
상태: pass(통과) · fail(실패 — 배운 것과 지금 결과가 다름) · warn(확인 필요 — 설계 원칙과 다름 · 도면 확인) ·
      gap(시험 불가 — 지금 도면으로 그 명령을 세울 길이 없음 · 읽기 빈틈) · skip(시험할 것 없음)
"""
from __future__ import annotations

import threading
import time

from . import behavior, sequence, sim


def _F(forces):
    return {int(p): dict(v) for p, v in (forces or {}).items()}


def _plus(F, page, key, v):
    G = {p: dict(d) for p, d in F.items()}
    G.setdefault(int(page), {})[key] = v
    return G


def _run(sheets, F, state, page, net, cache, timers):
    r = sim.simulate_bundle(sheets, F, state, page, timers=timers, cache=cache)
    fr = (r["focus"] or {}).get("frames") or []
    vals = [f["values"].get(net) for f in fr]
    return vals, r.get("state") or {}


def _chk(kind, x, title, expect, got, status, why="", forces=None, at=None, key=None):
    return {"kind": kind, "action": x.get("action"), "title": title, "expect": expect, "got": got, "status": status,
            "why": why, "forces": {str(p): v for p, v in (forces or {}).items()},
            "at": at or {"page": x["page"], "row": x["row"], "sheet": x.get("sheet")}, "key": key}


def _say(v):
    return "1" if v == 1 else "0" if v == 0 else "모름"


def test_command(sheets, x, cache, timers=None):
    """명령 하나 — 배운 것을 지금 결과로 다시 돌린다."""
    out = []
    act = x.get("action")
    if not x.get("reached"):
        # 학습이 어느 조합으로도 세우지 못한 명령 — 시험이 틀린 것이 아니라 시험할 길이 없다 (도면 읽기 빈틈 · 되먹임)
        b = x.get("blocked_at") or {}
        why = (x.get("why") or "안 섬") + (f" · 막힌 곳 {b.get('label')}" if b.get("label") else "")
        out.append(_chk("on", x, f"{act} 켜짐", "1", "안 섬", "gap", why))
        return out
    net = x.get("net")
    page = x["page"]
    if net is None:                                     # 예전 학습 (v2 이전) — 선이 없다
        out.append(_chk("on", x, f"{act} 켜짐", "1", "모름", "skip", "다시 학습 필요"))
        return out
    F = _F(x.get("forces"))
    vals, st = _run(sheets, F, None, page, net, cache, timers)
    on = 1 in vals
    real = [l["text"] for l in x.get("labels") or [] if not l.get("assumed")]
    out.append(_chk("on", x, f"{act} 켜짐 ← " + (" + ".join(real[:4]) or "신호 없음"),
                    "1", _say(1 if on else (vals[-1] if vals else None)), "pass" if on else "fail",
                    "" if on else "학습 때와 다름 — 지금은 안 섬", F))
    if not on:
        return out
    if x.get("always_on"):
        v0, _ = _run(sheets, {}, None, page, net, cache, timers)
        out.append(_chk("always", x, f"{act} 신호 없이 켜짐?", "0", _say(v0[-1] if v0 else None),
                        "warn" if v0 and v0[-1] == 1 else "pass",
                        "안전 쪽 설계인지 확인"))
    if x.get("tied_off"):
        out.append(_chk("tied", x, f"{act} 상수로 막힘", "다른 길", "상수 길뿐", "warn",
                        x["tied_off"][0]))
    for i, a in enumerate(x.get("alts") or []):
        G = _F(a.get("forces"))
        v, _ = _run(sheets, G, None, page, net, cache, timers)
        ok = 1 in v
        out.append(_chk("alt", x, f"다른 길 {i + 1} — " + " + ".join(l["text"] for l in a.get("labels") or [] if not l.get("assumed"))[:160],
                        "1", _say(1 if ok else (v[-1] if v else None)), "pass" if ok else "fail",
                        "" if ok else "학습 때와 다름", G))
    for l in x.get("interlocks") or []:
        G = _plus(F, l["page"], l["key"], 1)
        v, _ = _run(sheets, G, None, page, net, cache, timers)
        ok = 1 not in v
        out.append(_chk("block", x, f"{l.get('sheet') or ''} {l['text']} = 1 → {act} 막힘",
                        "0", _say(1 if not ok else 0), "pass" if ok else "fail",
                        "" if ok else "막지 못함 (학습 때와 다름)", G,
                        at={"page": l["page"], "sheet": l.get("sheet"), "key": l["key"]}, key=l["key"]))
    latch = x.get("latch")
    if latch is not None and vals and vals[-1] == 1:
        rest = {p: {k: v for k, v in d.items() if not (x.get("latch_released") == "all" or k.startswith("block:"))}
                for p, d in F.items()}
        rest = {p: d for p, d in rest.items() if d}
        v2, st2 = _run(sheets, rest, st, page, net, cache, timers)
        now = v2[-1] if v2 else None
        what = "단추 놓음" if x.get("latch_released") == "buttons" else "신호 거둠"
        if latch:
            out.append(_chk("hold", x, f"{what} → {act} 유지", "1", _say(now),
                            "pass" if now == 1 else "fail", "" if now == 1 else "학습 때와 다름", rest))
            for l in x.get("stoppers") or []:
                G = _plus(rest, l["page"], l["key"], 1)
                v3, _ = _run(sheets, G, st2, page, net, cache, timers)
                ok = bool(v3) and v3[-1] != 1
                out.append(_chk("stop", x, f"{l.get('sheet') or ''} {l['text']} = 1 → {act} 꺼짐",
                                "0", _say(v3[-1] if v3 else None), "pass" if ok else "fail",
                                "" if ok else "꺼지지 않음 (학습 때와 다름)", G,
                                at={"page": l["page"], "sheet": l.get("sheet"), "key": l["key"]}, key=l["key"]))
        else:
            out.append(_chk("hold", x, f"{what} → {act} 꺼짐", "0", _say(now),
                            "pass" if now != 1 else "fail", "" if now != 1 else "학습 때와 다름 — 유지됨", rest))
    elif x.get("pulse_s") is not None:
        ends = bool(vals) and vals[-1] != 1
        out.append(_chk("hold", x, f"{act} 펄스 {x['pulse_s']}초", "0", _say(vals[-1] if vals else None),
                        "pass" if ends else "fail", "" if ends else "학습 때와 다름", F))
    return out


def test_group(sheets, items, cache, timers=None, opposites=None):
    """장비 하나 — 명령마다 시험하고 반대 명령 짝을 함께 돌려 본다."""
    t0 = time.time()
    checks = []
    for x in items:
        try:
            checks += test_command(sheets, x, cache, timers)
        except Exception as ex:                          # 한 명령이 실패해도 나머지는 시험한다 — 사유를 남긴다
            checks.append(_chk("on", x, f"{x.get('action')} 시험", "", "오류", "fail", f"시뮬레이션 오류 — {ex}"))
    by = {x["action"]: x for x in items if x.get("reached") and x.get("net") is not None}
    for a, b in opposites if opposites is not None else (sequence.vocab().get("opposites") or []):
        if a not in by or b not in by:
            continue
        A, B = by[a], by[b]
        FA, FB = _F(A["forces"]), _F(B["forces"])
        G = {p: dict(d) for p, d in FA.items()}
        for p, d in FB.items():
            G.setdefault(p, {}).update(d)
        va, _ = _run(sheets, G, None, A["page"], A["net"], cache, timers)
        vb, _ = _run(sheets, G, None, B["page"], B["net"], cache, timers)
        both = bool(va) and bool(vb) and va[-1] == 1 and vb[-1] == 1
        checks.append(_chk("pair", A, f"{a} · {b} 동시 금지", "하나만", "둘 다 1" if both else "하나만",
                           "warn" if both else "pass",
                           "서로 막는 신호 없음 — 확인"
                           if both else "", G))
        if A.get("latch"):
            vals, st = _run(sheets, FA, None, A["page"], A["net"], cache, timers)
            if vals and vals[-1] == 1:
                rest = {p: {k: v for k, v in d.items() if not k.startswith("block:")} for p, d in FA.items()}
                H = {p: dict(d) for p, d in rest.items() if d}
                for p, d in FB.items():
                    H.setdefault(p, {}).update(d)
                v2, _ = _run(sheets, H, st, A["page"], A["net"], cache, timers)
                kept = bool(v2) and v2[-1] == 1
                checks.append(_chk("cancel", A, f"{b} 조건 → {a} 꺼짐", "0", _say(v2[-1] if v2 else None),
                                   "warn" if kept else "pass",
                                   f"{a} 가 남음 — 확인" if kept else "", H))
    n = {s: sum(1 for c in checks if c["status"] == s) for s in ("pass", "fail", "warn", "gap", "skip")}
    worst = next((s for s in ("fail", "warn", "gap", "pass") if n[s]), "skip")
    return {"status": worst, "counts": n, "checks": checks, "seconds": round(time.time() - t0, 2), "at": time.time()}


def groups(learned):
    """배운 것 → 장비(블록 그룹) 목록.  계통은 그 장비 첫 명령 시트의 제목."""
    eq = {}
    for x in (learned or {}).get("items") or []:
        e = eq.setdefault(x["equip"], {"equip": x["equip"], "actions": [], "sheets": []})
        e["actions"].append({k: x.get(k) for k in ("action", "page", "row", "sheet", "desc", "reached", "latch",
                                                     "pulse_s", "always_on", "signals")})
        e["actions"][-1]["interlocks"] = len(x.get("interlocks") or [])
        e["actions"][-1]["alts"] = len(x.get("alts") or [])
        e["actions"][-1]["effects"] = len(x.get("effects") or [])
        if x.get("sheet") and x["sheet"] not in e["sheets"]:
            e["sheets"].append(x["sheet"])
    pairs = {}
    for p in (learned or {}).get("pairs") or []:
        pairs.setdefault(p["equip"], []).append(p)
    for name, e in eq.items():
        e["pairs"] = pairs.get(name, [])
    return sorted(eq.values(), key=lambda e: e["equip"])


def items_of(learned, equip):
    return [x for x in (learned or {}).get("items") or [] if x.get("equip") == equip]


# ------------------------------------------------------------------ 저장 (분석 폴더 grouptest.json)

FILE = "grouptest.json"


def load(d):
    import json
    try:
        return json.loads((d / FILE).read_text(encoding="utf-8"))
    except Exception:
        return {"groups": {}, "history": []}


_SAVE = threading.Lock()


def save(d, equip, res, author, key):
    with _SAVE:                                          # 두 사람이 동시에 시험해도 앞 결과를 잃지 않게
        return _save(d, equip, res, author, key)


def _save(d, equip, res, author, key):
    import json
    from . import store
    cur = load(d)
    res = dict(res, author=author or "", key=key)
    cur.setdefault("groups", {})[equip] = res
    cur.setdefault("history", []).append({"equip": equip, "status": res["status"], "counts": res["counts"],
                                          "author": author or "", "at": res["at"]})
    cur["history"] = cur["history"][-2000:]
    store._atomic(d / FILE, json.dumps(cur, ensure_ascii=False))
    return res


def clear(d):
    try:
        (d / FILE).unlink()
    except OSError:
        pass


# ------------------------------------------------------------------ Excel (표 그대로)

STATUS_KO = {"pass": "통과", "fail": "실패", "warn": "확인 필요", "gap": "시험 불가", "skip": "시험 안 함", None: "미시험"}
KIND_KO = {"on": "켜기", "alt": "다른 길", "block": "막는 신호", "hold": "유지/놓기", "stop": "끄는 신호",
           "always": "늘 켜짐", "tied": "상수 묶임", "pair": "반대 명령 동시", "cancel": "반대 명령이 끔"}


def workbook(learned, tested, name):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "그룹 요약"
    fill = {"pass": "D9F2E1", "fail": "F8D7D7", "warn": "FFF1C2", "gap": "E5E7EB"}
    head = ["장비(블록 그룹)", "시트", "동작", "결과", "통과", "실패", "확인", "시험 불가", "시험한 사람", "시험 시각"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    tg = (tested or {}).get("groups") or {}
    for g in groups(learned):
        r = tg.get(g["equip"]) or {}
        n = r.get("counts") or {}
        ws.append([g["equip"], " ".join(g["sheets"][:4]), " ".join(a["action"] for a in g["actions"]),
                   STATUS_KO.get(r.get("status")), n.get("pass"), n.get("fail"), n.get("warn"), n.get("gap"), r.get("author"),
                   time.strftime("%Y-%m-%d %H:%M", time.localtime(r["at"])) if r.get("at") else ""])
        if r.get("status") in fill:
            ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=fill[r["status"]])
    ws2 = wb.create_sheet("시험 항목")
    ws2.append(["장비", "동작", "항목", "시험 내용", "기대", "결과", "판정", "사유", "시트", "행"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for eq, r in sorted(tg.items()):
        for c in r.get("checks") or []:
            at = c.get("at") or {}
            ws2.append([eq, c.get("action"), KIND_KO.get(c["kind"], c["kind"]), c["title"], c["expect"], c["got"],
                        STATUS_KO.get(c["status"]), c.get("why") or "", at.get("sheet") or "", at.get("row") or ""])
            if c["status"] in fill:
                ws2.cell(ws2.max_row, 7).fill = PatternFill("solid", fgColor=fill[c["status"]])
    for w, cols in ((ws, (30, 18, 20, 12, 7, 7, 7, 9, 12, 17)), (ws2, (28, 8, 12, 70, 10, 10, 10, 50, 10, 6))):
        for i, v in enumerate(cols):
            w.column_dimensions[chr(65 + i)].width = v
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
