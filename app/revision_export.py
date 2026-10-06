"""hotfix41 — 개정 변경 내역 Excel (추가 · 수정 · 삭제 · 장 변경 · 요약).

사용자 목적(hotfix40): *최신 식별 결과를 이전과 비교해 변경이 있으면 눈으로 대조한다.*  화면이 보여
주는 것을 **파일로도** 가져갈 수 있어야 회의·회신에 쓸 수 있다.  여기서 판정하지 않는다 —
`revision_state`(행 상태 · 바뀐 칸) · `deleted_candidate` · 장부의 `sheets` 를 그대로 옮겨 적는다.
발주처 양식(`excel_out`)과는 다른 파일이다: 그 양식은 발주처 열만 담고 개정은 REMARK 한 줄이다.
"""
from __future__ import annotations

import io
from typing import Iterable

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HEAD_FILL = PatternFill("solid", fgColor="EFE8D8")
STATE_FILL = {"ADDED": "FFF1D6", "MODIFIED": "E3F1E6", "DELETED_CANDIDATE": "F8E1DF", "DELETED": "F2C9C6"}

ROW_COLS = [("P&ID No.", "drawing_no"), ("쪽", "page_no"), ("탭", "tab"), ("Type", "type"),
            ("Tag No.", "tag_no"), ("Line No.", "line_no"), ("Q'ty", "qty"), ("SCOPE", "scope"),
            ("Description", "description"), ("안정 ID", "stable_id"), ("짝 근거", "basis")]


def _val(row: dict, pages: dict, name: str):
    if name == "drawing_no":
        return pages.get(row.get("page_no")) or row.get("drawing_no") or ""
    if name in ("page_no", "tab"):
        return row.get(name)
    if name == "stable_id":
        return (row.get("rev") or {}).get("id") or ""
    if name == "basis":
        return (row.get("rev") or {}).get("basis") or ""
    return (row.get("values") or {}).get(name, "")


def _sheet(wb, title: str, header: Iterable[str]):
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    ws.freeze_panes = "A2"
    return ws


def _fit(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def build(job: dict, rows: list, rev_states: dict, deleted: list, pages: dict,
          sheets: dict, matched_by: dict) -> bytes:
    """워크북 바이트.  `rows` 는 `db.merged_rows` 의 행(`rev` 를 얹은 것) · `pages` 는 쪽 → 도면번호."""
    revision = job.get("revision") or ""
    against = job.get("compared_with") or ""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "요약"
    added = [r for r in rows if (r.get("rev") or {}).get("state") == "ADDED"]
    modified = [r for r in rows if (r.get("rev") or {}).get("state") == "MODIFIED"]
    unchanged = sum(1 for r in rows if (r.get("rev") or {}).get("state") == "UNCHANGED")
    cand = [d for d in deleted if not d.get("confirmed")]
    conf = [d for d in deleted if d.get("confirmed")]
    for line in [("문서", job.get("pdf_name") or ""), ("리비전", revision), ("비교 대상", against),
                 ("", ""), ("추가", len(added)), ("수정", len(modified)), ("변경 없음", unchanged),
                 ("삭제 후보", len(cand)), ("삭제 확정", len(conf)), ("", ""),
                 ("짝 — 태그", (matched_by or {}).get("TAG", 0)),
                 ("짝 — 기하", (matched_by or {}).get("GEOMETRY", 0)), ("", ""),
                 ("도면번호 바뀐 장", len((sheets or {}).get("renumbered") or [])),
                 ("이번에만 있는 장", len((sheets or {}).get("only_now") or [])),
                 ("직전에만 있는 장", len((sheets or {}).get("only_before") or [])), ("", ""),
                 ("읽는 법", "판정은 분석 때 저장된 그대로입니다 (행 상태 · 바뀐 태그 · 삭제 후보).  "
                            "수정은 태그가 달라진 행뿐이고(자리 이동 · 수량 · SCOPE · Description 차이는 "
                            "참고로만 적습니다), 삭제 후보는 사람이 확정하기 전에는 발주처 양식에 나가지 않습니다.")]:
        ws.append(list(line))
    ws["A1"].font = ws["A2"].font = ws["A3"].font = Font(bold=True)
    _fit(ws, {"A": 18, "B": 80})

    def rows_sheet(title, items, extra=()):
        w = _sheet(wb, title, ["NO"] + [h for h, _ in ROW_COLS] + list(extra))
        for i, r in enumerate(items, 1):
            line = [i] + [_val(r, pages, n) for _, n in ROW_COLS]
            if extra:
                ch = (r.get("rev") or {}).get("changed") or []
                line.append(" · ".join(f"{c.get('field')}: {c.get('was') or '(빈칸)'} → {c.get('now') or '(빈칸)'}"
                                       for c in ch if isinstance(c, dict)))
                # hotfix43 — 상태를 정하지 않은 값 차이는 따로 적는다 (참고)
                fd = (r.get("rev") or {}).get("field_diffs") or []
                line.append(" · ".join(f"{c.get('field')}: {c.get('was') or '(빈칸)'} → {c.get('now') or '(빈칸)'}"
                                       for c in fd if isinstance(c, dict)))
                line.append((r.get("rev") or {}).get("moved_pt"))
            w.append(line)
            fill = PatternFill("solid", fgColor=STATE_FILL[(r.get("rev") or {}).get("state")])
            w.cell(row=i + 1, column=1).fill = fill
        _fit(w, {"A": 5, "B": 26, "C": 5, "D": 10, "E": 10, "F": 16, "G": 12, "H": 6, "I": 18,
                 "J": 44, "K": 16, "L": 9, "M": 36, "N": 50, "O": 9})
        return w

    rows_sheet("추가", sorted(added, key=lambda r: (r.get("page_no") or 0, _val(r, pages, "tag_no"))))
    rows_sheet("수정", sorted(modified, key=lambda r: (r.get("page_no") or 0, _val(r, pages, "tag_no"))),
               extra=("바뀐 태그 (전 → 후)", "그 밖에 값이 다른 칸 (참고 · 개정 판정에 안 씀)", "자리 이동 pt"))

    w = _sheet(wb, "삭제", ["NO", "상태", "안정 ID", "P&ID No.", "쪽", "탭", "Type", "Tag No.",
                           "Description", "근거", "직전 Rev 좌표", "이번 분석의 다른 도면"])
    for i, d in enumerate(sorted(deleted, key=lambda d: (d.get("drawing_no") or "", d.get("id") or "")), 1):
        if d.get("tag_no"):
            why = (f"태그 {d['tag_no']} 가 이번 분석에서 다른 도면에 섰습니다 — 옮김일 수 있습니다"
                   if d.get("tag_elsewhere") else f"태그 {d['tag_no']} 가 이번 분석 어디에도 없습니다")
        elif d.get("nearest_distance") is None:
            why = "그 도면에 같은 TYPE 이 하나도 없습니다"
        else:
            why = f"가장 가까운 같은 TYPE 이 {d['nearest_distance']}pt (반경 {d.get('radius')}pt)"
        w.append([i, "삭제 확정" if d.get("confirmed") else "삭제 후보", d.get("id", ""),
                  d.get("drawing_no", ""), d.get("page_no"), d.get("tab", ""), d.get("type", ""),
                  d.get("tag_no", ""), d.get("description", ""), why,
                  ", ".join(str(v) for v in (d.get("anchor") or [])),
                  ", ".join(d.get("tag_elsewhere") or [])])
        w.cell(row=i + 1, column=1).fill = PatternFill(
            "solid", fgColor=STATE_FILL["DELETED" if d.get("confirmed") else "DELETED_CANDIDATE"])
    _fit(w, {"A": 5, "B": 10, "C": 16, "D": 26, "E": 5, "F": 10, "G": 8, "H": 16, "I": 44, "J": 52, "K": 16, "L": 26})

    w = _sheet(wb, "장", ["구분", "직전 도면번호", "이번 도면번호", "공유 태그", "직전 태그 수", "이번 태그 수"])
    for e in (sheets or {}).get("renumbered") or []:
        w.append(["도면번호 바뀜", e.get("before", ""), e.get("now", ""), e.get("shared"),
                  e.get("before_tags"), e.get("now_tags")])
    for d in (sheets or {}).get("only_before") or []:
        w.append(["직전에만 있음", d, "", "", "", ""])
    for d in (sheets or {}).get("only_now") or []:
        w.append(["이번에만 있음", "", d, "", "", ""])
    _fit(w, {"A": 14, "B": 26, "C": 26, "D": 10, "E": 12, "F": 12})
    for s in wb.worksheets:
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
