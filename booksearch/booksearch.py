# -*- coding: utf-8 -*-
"""
내용 기반 도서 검색기 (Book Content Search) — 하이브리드 검색판

- 폴더를 선택하면 하위 폴더까지 모든 도서 파일(PDF, EPUB, DOCX, TXT, MD, HWPX)을 읽어
  본문을 문장 단위로 묶은 조각(청크)으로 나누고 '의미 벡터(임베딩)'로 인덱싱합니다.
- 모든 처리는 PC 안에서만 이루어집니다(외부 서버로 문서 전송 없음).
  최초 1회만 임베딩 모델을 내려받습니다.

검색 알고리즘
 1) 청크 분할: 문장 경계에서 자르고 ~450자로 묶음(앞 조각의 마지막 문장 1개 겹침).
    e5 입력 한도(512토큰) 안에 들어가 뒷부분이 잘려 버려지지 않음.
 2) 의미 검색: multilingual-e5 임베딩의 코사인 유사도(전체 청크 행렬곱 1회).
 3) 어휘 검색: BM25. 한국어는 조사를 떼어낸 어절 + 2글자 조각을 '부분 문자열'로 세므로
    형태소 분석기 없이도 '보안에', '정보보안', '보안성' 같은 변형을 잡음.
 4) 두 순위를 청크 단위 Reciprocal Rank Fusion(RRF)으로 결합 — 점수 척도가 달라도 안정적.
 5) 책 점수 = 그 책 상위 청크 RRF 점수의 감쇠 합(1, 1/2, 1/3, …).
    한 번 스쳐 언급한 책보다 여러 곳에서 깊게 다루는 책이 위로 올라옴.
 6) (선택) 재순위: BOOKSEARCH_RERANKER 에 cross-encoder 모델을 지정하면 상위 후보 청크를
    다시 채점해 RRF의 세 번째 순위로 더함. 정확도는 가장 높지만 느리고 모델이 큼.
"""
import os
import sys
import re
import json
import html
import time
import hashlib
import zipfile
import threading
import subprocess

import numpy as np

# ---------------------------------------------------------------- 설정
# 사내망 등에서 다운로드가 막히면 모델 폴더를 미리 받아두고 경로를 지정:
#   set BOOKSEARCH_MODEL=C:\models\multilingual-e5-small
# 정확도를 더 원하면 intfloat/multilingual-e5-base (느리지만 더 정확)
MODEL_NAME = os.environ.get("BOOKSEARCH_MODEL", "intfloat/multilingual-e5-small")
# 선택: 예) set BOOKSEARCH_RERANKER=BAAI/bge-reranker-v2-m3
RERANKER_NAME = os.environ.get("BOOKSEARCH_RERANKER", "").strip()
INDEX_DIR = ".booksearch"            # 선택한 폴더 안에 생성되는 인덱스 폴더
INDEX_VERSION = 2                    # 청크 방식이 바뀌면 올림 → 자동 재인덱싱
SUPPORTED = {".pdf", ".epub", ".docx", ".txt", ".md", ".hwpx"}
CHUNK_CHARS = 450                    # 한 조각 목표 길이(글자)
CHUNK_MIN = 40                       # 이보다 짧은 조각은 버림(쪽번호·머리글 등)
MAX_CHUNKS_PER_BOOK = 4000           # 아주 큰 문서만 고르게 샘플링
BM25_K1, BM25_B = 1.2, 0.75
RRF_K = 60
BOOK_TOP_M = 5                       # 책 점수에 쓰는 상위 청크 수
DECAY = 1.0 / np.arange(1, BOOK_TOP_M + 1)
SNIPPETS_PER_BOOK = 3
RERANK_CANDIDATES = 100
CONFIG_PATH = os.path.join(os.path.expanduser("~"), ".booksearch_config.json")

STOPWORDS = {"관련", "관련된", "관련한", "대한", "대해", "대해서", "관한", "관해", "관하여", "위한",
             "책", "도서", "내용", "주제", "있는", "어떤", "무엇", "찾아줘", "알려줘", "추천", "추천해줘"}
JOSA = sorted(["으로부터", "에서부터", "으로서", "으로써", "에게서", "이라는", "이라고", "라는", "라고",
               "에서", "에게", "으로", "까지", "부터", "처럼", "보다", "하고", "이나", "과", "와",
               "은", "는", "이", "가", "을", "를", "에", "의", "도", "만", "로"], key=len, reverse=True)


# ---------------------------------------------------------------- 텍스트 추출
def _strip_markup(s: str) -> str:
    s = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", s, flags=re.S | re.I)
    s = re.sub(r"</(p|div|h\d|li|br|hp:p)\s*>|<br\s*/?>", "\n\n", s, flags=re.I)   # 문단 경계 보존
    s = re.sub(r"<[^>]+>", " ", s)
    return html.unescape(s)


def extract_pdf(path):
    import fitz  # PyMuPDF
    with fitz.open(path) as doc:
        return "\n\n".join(page.get_text() for page in doc)


def extract_epub(path):
    parts = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if n.lower().endswith((".xhtml", ".html", ".htm")):
                parts.append(_strip_markup(z.read(n).decode("utf-8", errors="ignore")))
    return "\n\n".join(parts)


def extract_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n\n".join(parts)


def extract_hwpx(path):
    parts = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if re.match(r"Contents/section\d+\.xml", n):
                parts.append(_strip_markup(z.read(n).decode("utf-8", errors="ignore")))
    return "\n\n".join(parts)


def extract_txt(path):
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8", "cp949", "euc-kr"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="ignore")


EXTRACTORS = {
    ".pdf": extract_pdf, ".epub": extract_epub, ".docx": extract_docx,
    ".hwpx": extract_hwpx, ".txt": extract_txt, ".md": extract_txt,
}


def extract_text(path):
    return EXTRACTORS[os.path.splitext(path)[1].lower()](path)


# ---------------------------------------------------------------- 청크 분할
SENT_END = re.compile(r"(?<=[.!?。？！…])\s+")


def split_sentences(text):
    out = []
    for para in re.split(r"\n\s*\n", text):
        para = re.sub(r"\s+", " ", para).strip()
        for s in SENT_END.split(para):
            while len(s) > CHUNK_CHARS:            # 문장부호 없이 긴 덩어리는 공백에서 끊음
                cut = s.rfind(" ", 0, CHUNK_CHARS)
                cut = cut if cut > CHUNK_CHARS // 2 else CHUNK_CHARS
                out.append(s[:cut])
                s = s[cut:].lstrip()
            if s:
                out.append(s)
    return out


def chunk_text(text):
    chunks, cur, size = [], [], 0
    for s in split_sentences(text):
        if cur and size + len(s) > CHUNK_CHARS:
            chunks.append(" ".join(cur))
            cur = [cur[-1]] if len(cur[-1]) < CHUNK_CHARS // 3 else []   # 문맥 유지용 1문장 겹침
            size = sum(len(x) + 1 for x in cur)
        cur.append(s)
        size += len(s) + 1
    if cur:
        chunks.append(" ".join(cur))
    chunks = [c for c in chunks if len(c) >= CHUNK_MIN]
    if len(chunks) > MAX_CHUNKS_PER_BOOK:
        idx = np.linspace(0, len(chunks) - 1, MAX_CHUNKS_PER_BOOK).astype(int)
        chunks = [chunks[i] for i in idx]
    return chunks


# ---------------------------------------------------------------- 질의어 분석 (어휘 검색용)
def strip_josa(w):
    for j in JOSA:
        if w.endswith(j) and len(w) - len(j) >= 2:
            return w[:-len(j)]
    return w


def query_terms(query):
    """[(검색어, 가중치)] — 어절은 1.0, 한글 3글자 이상 어절의 2글자 조각은 0.3."""
    terms = {}
    for w in re.findall(r"[가-힣]+|[A-Za-z][A-Za-z0-9+#]*|\d+", query):
        if w in STOPWORDS:
            continue
        hangul = "가" <= w[0] <= "힣"
        if hangul:
            w = strip_josa(w)
            if w in STOPWORDS:
                continue
        w = w.lower()
        if len(w) < 2:
            continue
        terms[w] = 1.0
        if hangul and len(w) >= 3:
            for i in range(len(w) - 1):
                terms.setdefault(w[i:i + 2], 0.3)
    return list(terms.items())


# ---------------------------------------------------------------- 모델
class Embedder:
    def __init__(self, name=MODEL_NAME):
        from sentence_transformers import SentenceTransformer
        self.name = name
        self.model = SentenceTransformer(name)
        self.e5 = "e5" in name.lower()
        self.dim = self.model.get_sentence_embedding_dimension()

    def passages(self, texts, batch_size=32, progress=None):
        if self.e5:
            texts = ["passage: " + t for t in texts]
        out = []
        for i in range(0, len(texts), batch_size):
            out.append(self.model.encode(texts[i:i + batch_size], normalize_embeddings=True))
            if progress:
                progress(min(i + batch_size, len(texts)), len(texts))
        return np.vstack(out).astype(np.float32) if out else np.zeros((0, self.dim), np.float32)

    def query(self, q):
        if self.e5:
            q = "query: " + q
        return self.model.encode([q], normalize_embeddings=True)[0].astype(np.float32)


class Reranker:
    def __init__(self, name):
        from sentence_transformers import CrossEncoder
        self.model = CrossEncoder(name, max_length=512)

    def score(self, query, texts):
        return np.asarray(self.model.predict([(query, t) for t in texts], batch_size=16), np.float32)


def rrf(scores, mask=None):
    order = np.argsort(-scores, kind="stable")
    ranks = np.empty(len(scores), np.float32)
    ranks[order] = np.arange(len(scores))
    r = 1.0 / (RRF_K + ranks + 1.0)
    if mask is not None:
        r[~mask] = 0.0
    return r


# ---------------------------------------------------------------- 인덱스
class BookIndex:
    def __init__(self, folder):
        self.folder = folder
        self.dir = os.path.join(folder, INDEX_DIR)
        os.makedirs(self.dir, exist_ok=True)
        self.meta_path = os.path.join(self.dir, "meta.json")
        self.meta = self._empty_meta()
        if os.path.exists(self.meta_path):
            try:
                with open(self.meta_path, encoding="utf-8") as f:
                    m = json.load(f)
                if m.get("model") == MODEL_NAME and m.get("version") == INDEX_VERSION:
                    self.meta = m                       # 모델·방식이 바뀌면 전체 재인덱싱
            except (OSError, ValueError):
                pass
        self.books = []          # [(rel, title, start, end)]
        self.texts = []
        self.emb = None
        self._blob, self._starts, self._lens = "", None, None

    @staticmethod
    def _empty_meta():
        return {"model": MODEL_NAME, "version": INDEX_VERSION, "files": {}}

    def _key(self, rel):
        return hashlib.sha1(rel.replace("\\", "/").encode("utf-8")).hexdigest()[:16]

    def _save_meta(self, files):
        self.meta = {**self._empty_meta(), "files": files, "built": time.time()}
        tmp = self.meta_path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.meta, f, ensure_ascii=False, indent=1)
        os.replace(tmp, self.meta_path)

    def scan(self):
        found = []
        for root, dirs, files in os.walk(self.folder):
            dirs[:] = [d for d in dirs if d != INDEX_DIR and not d.startswith(".")]
            for fn in files:
                if os.path.splitext(fn)[1].lower() in SUPPORTED and not fn.startswith("~$"):
                    found.append(os.path.relpath(os.path.join(root, fn), self.folder))
        return sorted(found)

    def build(self, embedder, log=lambda m: None, progress=lambda a, b: None):
        files = self.scan()
        old = self.meta["files"]
        new_meta = {}
        for n, rel in enumerate(files, 1):
            full = os.path.join(self.folder, rel)
            st = os.stat(full)
            sig = {"mtime": st.st_mtime, "size": st.st_size}
            key = self._key(rel)
            npy, js = os.path.join(self.dir, key + ".npy"), os.path.join(self.dir, key + ".json")
            prev = old.get(rel)
            if prev and prev["mtime"] == sig["mtime"] and prev["size"] == sig["size"] \
                    and os.path.exists(npy) and os.path.exists(js):
                new_meta[rel] = prev
                log(f"[{n}/{len(files)}] 변경 없음: {rel}")
                continue
            log(f"[{n}/{len(files)}] 읽는 중: {rel}")
            try:
                chunks = chunk_text(extract_text(full))
            except Exception as e:
                log(f"   ⚠ 읽기 실패 ({e.__class__.__name__}: {e})")
                continue
            if not chunks:
                log("   ⚠ 텍스트 없음 (스캔 이미지 PDF일 수 있음)")
                continue
            log(f"[{n}/{len(files)}] {len(chunks)}개 구절 분석 중: {rel}")
            emb = embedder.passages(chunks, progress=progress)
            np.save(npy, emb.astype(np.float16))
            with open(js, "w", encoding="utf-8") as f:
                json.dump(chunks, f, ensure_ascii=False)
            new_meta[rel] = {**sig, "title": os.path.splitext(os.path.basename(rel))[0],
                             "chunks": len(chunks)}
            # 중간 저장: 도중에 꺼져도 다음 실행 때 이미 끝난 책은 건너뜀
            self._save_meta({**{r: v for r, v in old.items() if r not in new_meta}, **new_meta})
        self._save_meta(new_meta)
        # 목록에 없는 인덱스 파일 정리(삭제된 책, 모델 변경 전 파일 포함)
        keep = {self._key(r) for r in new_meta}
        for fn in os.listdir(self.dir):
            stem, ext = os.path.splitext(fn)
            if ext in (".npy", ".json") and fn != "meta.json" and stem not in keep:
                try:
                    os.remove(os.path.join(self.dir, fn))
                except OSError:
                    pass
        self.load()

    def load(self):
        books, texts, embs = [], [], []
        for rel, info in sorted(self.meta["files"].items()):
            key = self._key(rel)
            try:
                e = np.load(os.path.join(self.dir, key + ".npy")).astype(np.float32)
                with open(os.path.join(self.dir, key + ".json"), encoding="utf-8") as f:
                    t = json.load(f)
            except Exception:
                continue
            if len(t) != len(e):
                continue
            start = len(texts)
            texts.extend(t)
            embs.append(e)
            books.append((rel, info.get("title", rel), start, len(texts)))
        # 어휘 검색용: 모든 청크를 한 문자열로 이어 두고 위치→청크 번호를 이분 탐색
        low = [t.lower() for t in texts]
        self._starts = np.cumsum([0] + [len(t) + 1 for t in low[:-1]]) if low else np.zeros(0, int)
        self._blob = "\x00".join(low)
        self._lens = np.array([len(t) for t in low], np.float32)
        self.books, self.texts = books, texts
        self.emb = np.vstack(embs) if embs else None
        return len(self.books)

    def bm25(self, terms):
        n = len(self.texts)
        score = np.zeros(n, np.float32)
        if not terms or n == 0:
            return score
        norm = BM25_K1 * (1 - BM25_B + BM25_B * self._lens / max(self._lens.mean(), 1.0))
        for term, w in terms:
            pos = [m.start() for m in re.finditer(re.escape(term), self._blob)]
            if not pos:
                continue
            tf = np.bincount(np.searchsorted(self._starts, pos, side="right") - 1,
                             minlength=n).astype(np.float32)
            df = int((tf > 0).sum())
            idf = np.log(1.0 + (n - df + 0.5) / (df + 0.5))
            score += w * idf * tf * (BM25_K1 + 1) / (tf + norm)
        return score

    def search(self, embedder, query, top_n=30, reranker=None):
        if self.emb is None or not self.books:
            return []
        dense = self.emb @ embedder.query(query)
        lex = self.bm25(query_terms(query))
        hit = lex > 0
        fused = rrf(dense)
        if hit.any():
            fused += rrf(lex, hit)
        if reranker is not None:
            c = min(RERANK_CANDIDATES, len(fused))
            cand = np.argpartition(-fused, c - 1)[:c]
            rr = reranker.score(query, [self.texts[i] for i in cand])
            fused[cand] += rrf(rr)

        results = []
        for rel, title, s, e in self.books:
            seg = fused[s:e]
            m = min(BOOK_TOP_M, e - s)
            top = np.argpartition(-seg, m - 1)[:m]
            top = top[np.argsort(-seg[top])]
            results.append({"rel": rel, "title": title,
                            "score": float(seg[top] @ DECAY[:m]),
                            "hits": int(hit[s:e].sum()),
                            "snippets": [self.texts[s + int(i)] for i in top[:SNIPPETS_PER_BOOK]]})
        results.sort(key=lambda r: -r["score"])
        return results[:top_n]


# ---------------------------------------------------------------- GUI
def open_file(path):
    if sys.platform.startswith("win"):
        os.startfile(path)
    elif sys.platform == "darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


class App:
    def __init__(self, root):
        import tkinter as tk
        from tkinter import ttk
        self.root = root
        root.title("내용 기반 도서 검색기")
        root.geometry("1000x700")
        self.embedder = None
        self.reranker = None
        self.index = None
        self.results = []
        self.terms = []
        self.working = False

        top = ttk.Frame(root, padding=8)
        top.pack(fill="x")
        self.btn_folder = ttk.Button(top, text="📁 폴더 선택", command=self.choose)
        self.btn_folder.pack(side="left")
        self.folder_var = tk.StringVar(value="(폴더를 선택하세요)")
        ttk.Label(top, textvariable=self.folder_var).pack(side="left", padx=8)
        self.btn_index = ttk.Button(top, text="🔄 인덱싱(새 파일 반영)", command=self.reindex, state="disabled")
        self.btn_index.pack(side="right")

        sf = ttk.Frame(root, padding=(8, 0))
        sf.pack(fill="x")
        self.q_var = tk.StringVar()
        ent = ttk.Entry(sf, textvariable=self.q_var, font=("맑은 고딕", 12))
        ent.pack(side="left", fill="x", expand=True)
        ent.bind("<Return>", lambda e: self.search())
        self.btn_search = ttk.Button(sf, text="🔍 검색", command=self.search, state="disabled")
        self.btn_search.pack(side="left", padx=6)

        pane = ttk.PanedWindow(root, orient="vertical")
        pane.pack(fill="both", expand=True, padx=8, pady=8)

        cols = ("rank", "title", "score", "hits")
        self.tree = ttk.Treeview(pane, columns=cols, show="headings", height=12)
        for c, w, t in [("rank", 50, "순위"), ("title", 560, "도서"),
                        ("score", 90, "관련도"), ("hits", 110, "키워드 등장")]:
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor="w" if c == "title" else "center")
        self.tree.bind("<<TreeviewSelect>>", self.show_snippet)
        self.tree.bind("<Double-1>", self.open_selected)
        pane.add(self.tree, weight=3)

        lf = ttk.LabelFrame(pane, text="근거 구절 (더블클릭하면 파일 열기)", padding=6)
        self.snip = tk.Text(lf, wrap="word", height=10, font=("맑은 고딕", 10))
        self.snip.tag_configure("hl", background="#fff2a8")
        self.snip.tag_configure("head", font=("맑은 고딕", 10, "bold"))
        self.snip.pack(fill="both", expand=True)
        pane.add(lf, weight=2)

        bottom = ttk.Frame(root, padding=(8, 0, 8, 8))
        bottom.pack(fill="x")
        self.pbar = ttk.Progressbar(bottom, mode="determinate")
        self.pbar.pack(fill="x")
        self.status = tk.StringVar(value="준비")
        ttk.Label(bottom, textvariable=self.status).pack(anchor="w")

        cfg = self._load_cfg()
        if cfg.get("folder") and os.path.isdir(cfg["folder"]):
            self.set_folder(cfg["folder"], auto=True)

    # ---- helpers
    def _load_cfg(self):
        try:
            with open(CONFIG_PATH, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}

    def _save_cfg(self, folder):
        try:
            with open(CONFIG_PATH, "w", encoding="utf-8") as f:
                json.dump({"folder": folder}, f, ensure_ascii=False)
        except Exception:
            pass

    def ui(self, fn, *a):
        self.root.after(0, fn, *a)

    def set_status(self, msg):
        self.ui(self.status.set, msg)

    def set_progress(self, done, total):
        self.ui(self.pbar.configure, {"maximum": max(total, 1), "value": done})

    def _set_buttons(self, on):
        st = "disabled" if on else "normal"
        self.btn_folder.configure(state=st)
        self.btn_index.configure(state=st if self.index else "disabled")
        self.btn_search.configure(state=st if self.index and self.index.books else "disabled")

    def busy(self, on):
        """작업 중에는 Enter 검색·폴더 변경까지 모두 막음(인덱스를 동시에 건드리지 않도록)."""
        if on:
            if self.working:
                return False
            self.working = True
            self._set_buttons(True)
            return True
        def done():
            self.working = False
            self._set_buttons(False)
        self.ui(done)
        return True

    def ensure_model(self):
        if self.embedder is None:
            self.set_status(f"모델 불러오는 중… ({MODEL_NAME}) — 최초 실행 시 다운로드로 몇 분 걸릴 수 있습니다")
            self.embedder = Embedder()
        if RERANKER_NAME and self.reranker is None:
            self.set_status(f"재순위 모델 불러오는 중… ({RERANKER_NAME})")
            self.reranker = Reranker(RERANKER_NAME)

    def run_bg(self, fn):
        if not self.busy(True):
            return
        from tkinter import messagebox

        def work():
            try:
                fn()
            except Exception as e:
                self.ui(messagebox.showerror, "오류", f"{e.__class__.__name__}: {e}")
                self.set_status("오류 발생")
            finally:
                self.busy(False)
        threading.Thread(target=work, daemon=True).start()

    # ---- actions
    def choose(self):
        from tkinter import filedialog
        if self.working:
            return
        d = filedialog.askdirectory(title="도서 폴더 선택")
        if d:
            self.set_folder(d)

    def set_folder(self, d, auto=False):
        self.folder_var.set(d)
        self._save_cfg(d)
        self.index = BookIndex(d)
        self.results = []
        self.tree.delete(*self.tree.get_children())
        n_files = len(self.index.scan())
        if self.index.meta["files"]:
            n = self.index.load()
            self.status.set(f"인덱스 {n}권 불러옴 (폴더 내 파일 {n_files}개). 새 파일이 있으면 [인덱싱]을 누르세요.")
        elif auto:
            self.status.set(f"폴더 내 파일 {n_files}개 — [인덱싱]을 눌러 주세요.")
        self._set_buttons(False)
        if not self.index.meta["files"] and not auto:
            self.reindex()

    def reindex(self):
        if not self.index:
            return
        index = self.index

        def job():
            self.ensure_model()
            t0 = time.time()
            index.build(self.embedder, log=self.set_status, progress=self.set_progress)
            self.set_status(f"인덱싱 완료: {len(index.books)}권, "
                            f"{len(index.texts):,}개 구절 ({time.time() - t0:.0f}초)")
        self.run_bg(job)

    def search(self):
        q = self.q_var.get().strip()
        if not q or not self.index or not self.index.books:
            return

        def job():
            self.ensure_model()
            t0 = time.time()
            res = self.index.search(self.embedder, q, reranker=self.reranker)
            self.ui(self.fill, res, q, time.time() - t0)
        self.run_bg(job)

    def fill(self, res, q, secs):
        self.results = res
        self.terms = [t for t, w in query_terms(q) if w >= 1.0]
        self.tree.delete(*self.tree.get_children())
        if not res:
            self.status.set("결과 없음")
            return
        best = res[0]["score"] or 1.0
        for i, r in enumerate(res, 1):
            self.tree.insert("", "end", iid=str(i - 1),
                             values=(i, r["rel"], f"{int(round(100 * r['score'] / best))}%",
                                     f"{r['hits']}곳"))
        self.tree.selection_set("0")
        self.status.set(f"'{q}' 검색 결과 {len(res)}권 ({secs:.2f}초) — 관련도는 1위 대비 비율입니다")

    def show_snippet(self, _=None):
        sel = self.tree.selection()
        if not sel:
            return
        r = self.results[int(sel[0])]
        self.snip.delete("1.0", "end")
        self.snip.insert("end", f"[{r['title']}]\n", "head")
        for k, s in enumerate(r["snippets"], 1):
            self.snip.insert("end", f"\n{k}. …{s}…\n")
        for term in self.terms:
            idx = "1.0"
            while True:
                idx = self.snip.search(term, idx, stopindex="end", nocase=True)
                if not idx:
                    break
                end = f"{idx}+{len(term)}c"
                self.snip.tag_add("hl", idx, end)
                idx = end

    def open_selected(self, _=None):
        sel = self.tree.selection()
        if sel:
            open_file(os.path.join(self.index.folder, self.results[int(sel[0])]["rel"]))


if __name__ == "__main__":
    import tkinter as tk
    from tkinter import ttk
    root = tk.Tk()
    try:
        ttk.Style().theme_use("vista" if sys.platform.startswith("win") else "clam")
    except tk.TclError:
        pass
    App(root)
    root.mainloop()
